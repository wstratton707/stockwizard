"""QuantWizard's Excel equity report.

Reproduces the reference workbook (assets/AAPL_5Y_Analysis (6).xlsx - our
AAPL export rebuilt in Excel by Claude) for any ticker. Its styling, widths,
merges, frozen panes, dropdowns and conditional formats come from
assets/report_template.xlsx, which tools/build_report_template.py derives from
the reference with every Apple-specific value removed. This module writes the
content: labels, live formulas, the numbers they read, source notes, and the
tabs whose length depends on the data (Risk, Catalysts_News, Methodology,
Sources and the data tabs).

Two things the template cannot carry are restored after saving: the reference
charts' XML (so the Summary charts match it exactly) and Excel 365's
dynamic-array flag on the LET/SEQUENCE formulas (openpyxl drops it, which would
leave them showing as legacy {array} formulas).

Rules from the reference's Build_Spec that shape the code:
  R1 every derived number is a live formula; R2 blue-on-yellow inputs, green
     cross-sheet links, black formulas; R3 every sourced figure carries a note
     and a Sources row; R4 missing data reads N/A, never 0; R5 period labels
     come from the actual dates; R6 FY and TTM side by side, labelled;
  R7 market cap = 10-Q cover shares x the report price; R8/R9 the DCF's
     switches and live grids; R10 no copied Apple text - narrative is built
     from this company's filings and data.
"""
import copy
import io
import math
import os
import re
import zipfile
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

_HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(_HERE, "assets", "report_template.xlsx")
CHART_DIR = os.path.join(_HERE, "assets", "report_charts")

# Number formats, exactly as the reference writes them.
F_USD = '"$"#,##0.00;\\("$"#,##0.00\\);"–"'
F_BN = '#,##0.0;\\(#,##0.0\\);"–"'
F_PCT = '0.0%;\\(0.0%\\);"–"'
F_X = '0.0"x";\\(0.0"x"\\);"–"'
F_DATE = "dd\\-mmm\\-yyyy"
F_MONTH = "mmm\\-yyyy"

NAVY, BLUE, BAND = "FF1F3864", "FF2E75B6", "FFD9E1F2"
GREEN_TAB, GREY_TAB = "FF70AD47", "FFA6A6A6"
AUTHOR = "QuantWizard"

FRONT = ["Summary", "Financials", "DCF", "Scenarios", "Multiples", "Risk",
         "Catalysts_News", "Methodology", "Sources"]


# ── small helpers ─────────────────────────────────────────────────────────────
def _num(x):
    return isinstance(x, (int, float, np.integer, np.floating)) and not (
        isinstance(x, float) and math.isnan(x))


def _bn(x):
    return float(x) / 1e9 if _num(x) else None


def _na(x):
    return x if x is not None else "N/A"


def _d(x, fmt="%d-%b-%Y"):
    try:
        return pd.Timestamp(x).strftime(fmt)
    except Exception:
        return str(x or "")


def _part_of_month(ts):
    ts = pd.Timestamp(ts)
    return "early" if ts.day <= 10 else ("mid" if ts.day <= 20 else "late")


def _copy_style(src, dst):
    dst._style = copy.copy(src._style)


def _put(ws, addr, value, note=None, link=None, fmt=None, like=None):
    c = ws[addr]
    if like is not None:
        _copy_style(like, c)
    c.value = value
    if fmt:
        c.number_format = fmt
    if note:
        cm = Comment(note, AUTHOR)
        cm.width, cm.height = 320, 110
        c.comment = cm
    if link:
        c.hyperlink = link if "://" in link else None
        if "://" not in link:
            c.hyperlink = f"#{link}"
            c.hyperlink.location = link
    return c


def _array(ws, addr, formula):
    ws[addr].value = ArrayFormula(addr, formula)


def _reset(ws, min_row, max_row, min_col, max_col):
    """Blank a block, styles included, so a shorter table leaves nothing behind."""
    base = ws.parent._cell_styles[0] if ws.parent._cell_styles else None
    for r in range(min_row, max_row + 1):
        for cidx in range(min_col, max_col + 1):
            c = ws.cell(row=r, column=cidx)
            if type(c).__name__ == "MergedCell":
                continue
            c.value = None
            c.comment = None
            c.hyperlink = None
            if base is not None:
                c._style = copy.copy(base)


def _clear_heights(ws, from_row):
    """Drop row heights the template inherited from the reference's content, so
    rebuilt rows auto-fit (Excel sizes wrapped text on rows with no height)."""
    for r in list(ws.row_dimensions.keys()):
        if r >= from_row:
            ws.row_dimensions[r].height = None


def _snapshot(ws, cells):
    """{addr: StyleArray} taken before a sheet is cleared and rebuilt."""
    return {a: copy.copy(ws[a]._style) for a in cells}


def _style(ws, addr, snap):
    ws[addr]._style = copy.copy(snap)


def _back_link(ws, addr):
    _put(ws, addr, "← Back to Summary", link="'Summary'!A1")


def _clean_name(name, ticker):
    n = re.sub(r"\s+(Class [A-Z]\s+)?(Common Stock|Ordinary Shares|Common Shares|"
               r"Capital Stock|American Depositary Shares)\b.*$", "", name or "", flags=re.I)
    return n.strip() or ticker


def _short_name(name):
    s = re.sub(r"\s*&\s*Co\.?$", "", name.strip())
    for _ in range(2):
        s = re.sub(r",?\s+(Inc\.?|Incorporated|Corporation|Corp\.?|Company|Co\.?|Ltd\.?|"
                   r"plc|PLC|N\.V\.|S\.A\.|Holdings?|Group|Limited|L\.P\.)$", "", s,
                   flags=re.I).strip(" ,")
    return s or name


_EXCH_SHORT = {"nasdaq": "NASDAQ", "new york stock exchange": "NYSE", "nyse": "NYSE",
               "nyse american": "NYSE American", "nyse arca": "NYSE Arca", "cboe": "Cboe"}


def _exch_short(code):
    try:
        from analysis import exchange_name
        n = exchange_name(code) or ""
    except Exception:
        n = code or ""
    return _EXCH_SHORT.get(n.lower(), n.upper() if len(n) <= 6 else n)


# ── the report's inputs, derived once ─────────────────────────────────────────
class _Ctx:
    def __init__(self, ticker, df, fin, fund, dcf, company_details, mc_summary,
                 mc_sim_df, news_rows, peer_fund, peer_group, peer_df, segments,
                 valuation_data, filings, period_label, price_source, generated_at):
        self.t = ticker.upper()
        self.df = df.reset_index(drop=True)
        self.fin = fin or {}
        self.f = fund or {}
        self.dcf = dcf or {}
        self.cd = company_details or {}
        self.mc, self.mc_df = mc_summary or {}, mc_sim_df
        self.news = news_rows or []
        self.peers, self.peer_group, self.peer_df = peer_fund or [], peer_group, peer_df
        self.segments, self.vdata = segments or {}, valuation_data
        self.filings = filings or []
        self.period_label = period_label or ""
        self.src = price_source
        self.now = generated_at or datetime.now()

        self.name = _clean_name(self.cd.get("Name") or self.t, self.t)
        self.short = _short_name(self.name)
        self.exch = _exch_short(self.cd.get("Exchange"))
        inc = self.fin.get("income_statement")
        self.inc = inc if inc is not None else pd.DataFrame()
        bal = self.fin.get("balance_sheet")
        self.bal = bal if bal is not None else pd.DataFrame()
        cf = self.fin.get("cash_flow_statement")
        self.cf = cf if cf is not None else pd.DataFrame()
        self.ttm = self.fin.get("ttm") or {}

        # fiscal calendar
        per = list(self.inc["Period"]) if "Period" in self.inc.columns else []
        self.fy_ends = [pd.Timestamp(p) for p in per]                 # newest first
        self.fy_end = self.fy_ends[0] if self.fy_ends else None
        self.fy = self.fy_end.year if self.fy_end is not None else None
        self.fys = f"FY{str(self.fy)[2:]}" if self.fy else "FY"
        self.flows_end = pd.Timestamp(self.ttm["flows_end"]) if self.ttm.get("flows_end") else None
        self.has_ttm = bool(self.flows_end is not None and self.fy_end is not None
                            and self.flows_end > self.fy_end)
        self.bal_end = (pd.Timestamp(self.ttm["balance_end"]) if self.ttm.get("balance_end")
                        else self.fy_end)
        self.quarters = self.ttm.get("quarters") or []

        # price window
        dates = pd.to_datetime(self.df["Date"])
        self.first_date, self.last_date = dates.iloc[0], dates.iloc[-1]
        self.n = len(self.df)
        self.last_row = self.n + 1
        self.years = (self.last_date - self.first_date).days / 365.25
        self.price = float(self.df["Close"].iloc[-1])
        vol = self.df["Volume"].iloc[-1] if "Volume" in self.df.columns else None
        self.last_bar_no_volume = not _num(vol) or (_num(vol) and float(vol) == 0)
        self.has_bench = ("SPY_Return" in self.df.columns
                          and self.df["SPY_Return"].notna().sum() > 20)

        # filings by period, for source notes
        self.tenk = next((f for f in self.filings if f.get("form") == "10-K"), None)
        self.tenq = next((f for f in self.filings if f.get("form") == "10-Q"), None)
        self.by_period = {}
        for f in self.filings:
            if f.get("form") in ("10-Q", "10-K") and f.get("period"):
                self.by_period.setdefault(f["period"], f)
        cik = None
        try:
            from data import _sec_cik_for
            cik = _sec_cik_for(self.t, log=lambda *a, **k: None)
        except Exception:
            pass
        self.cik = cik
        self.edgar_index = (f"https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany"
                            f"&CIK={cik}&type=10-&dateb=&owner=include&count=40"
                            if cik else "https://www.sec.gov/edgar/search/")

    # -- labels ---------------------------------------------------------------
    def fy_desc(self):
        """'the last Saturday of September' / 'June 30' / 'late September'."""
        ends = self.fy_ends[:5]
        if not ends:
            return "at its fiscal year-end"
        wk = {e.weekday() for e in ends}
        last_week = all((e + pd.offsets.MonthEnd(0) - e).days < 7 for e in ends)
        if len(wk) == 1 and len({e.day for e in ends}) > 1 and last_week:
            return f"on the last {ends[0].strftime('%A')} of {ends[0].strftime('%B')}"
        if all(e.is_month_end for e in ends):
            return f"on {ends[0].strftime('%B')} {ends[0].day}"
        return f"in {_part_of_month(ends[0])} {ends[0].strftime('%B')}"

    def filing_for(self, period_end):
        key = _d(period_end, "%Y-%m-%d")
        f = self.by_period.get(key)
        if f is None:
            for p, v in self.by_period.items():
                if abs((pd.Timestamp(p) - pd.Timestamp(period_end)).days) <= 7:
                    return v
        return f

    def src_note(self, what, period_end=None, form=None):
        f = self.filing_for(period_end) if period_end is not None else None
        f = f or (self.tenq if form == "10-Q" else self.tenk if form == "10-K" else None)
        if f:
            return (f"Source: {self.name} Form {f.get('form')} for the period ended "
                    f"{_d(f.get('period'))} (filed {_d(f.get('filed'))}), {what}, "
                    f"SEC EDGAR XBRL, {f.get('url')}")
        return f"Source: SEC EDGAR XBRL company facts, {what}, {self.edgar_index}"


# ── data tabs ─────────────────────────────────────────────────────────────────
_PD_COLS = ["Date", "Open", "High", "Low", "Close", "Volume", "Daily_Return",
            "Cumulative_Index", "MA50", "MA200", "Volatility_20d", "Drawdown_60d",
            "Pct_From_52W_High", "RSI14", "MACD_Hist"]


def _price_data(wb, c):
    ws = wb["Price_Data"]
    hdr = {cidx: copy.copy(ws.cell(row=1, column=cidx)._style) for cidx in range(1, 18)}
    body = {cidx: copy.copy(ws.cell(row=3, column=cidx)._style) for cidx in range(1, 18)}
    first = {cidx: copy.copy(ws.cell(row=2, column=cidx)._style) for cidx in range(1, 18)}
    cols = list(_PD_COLS)
    extra = c.has_bench
    for j, name in enumerate(cols + ["Running_Peak", "Drawdown_From_Peak"]
                             + (["SPY_Return"] if extra else []), 1):
        cell = ws.cell(row=1, column=j, value=name)
        cell._style = copy.copy(hdr.get(j, hdr[17]))
    df = c.df
    for i in range(c.n):
        r = i + 2
        row = df.iloc[i]
        for j, name in enumerate(cols, 1):
            v = row.get(name) if name in df.columns else None
            if name == "Date":
                v = pd.Timestamp(v).to_pydatetime()
            elif name == "Volume":
                v = int(v) if _num(v) and float(v) > 0 else "n/a"
            elif _num(v):
                v = float(v)
            else:
                v = None
            cell = ws.cell(row=r, column=j, value=v)
            cell._style = copy.copy((first if r == 2 else body)[j])
        p = ws.cell(row=r, column=16, value=f"=MAX($E$2:E{r})")
        p._style = copy.copy((first if r == 2 else body)[16])
        q = ws.cell(row=r, column=17, value=f"=E{r}/P{r}-1")
        q._style = copy.copy((first if r == 2 else body)[17])
        if extra:
            sv = row.get("SPY_Return")
            s = ws.cell(row=r, column=18, value=float(sv) if _num(sv) else None)
            s._style = copy.copy(body[7])
    last = c.last_row
    ws.conditional_formatting.add(f"G2:G{last}", ColorScaleRule(
        start_type="num", start_value=-0.05, start_color="FFFFAAAA",
        mid_type="num", mid_value=0, mid_color="FFFFFFFF",
        end_type="num", end_value=0.05, end_color="FFAAFFAA"))
    ws.conditional_formatting.add(f"N2:N{last}", CellIsRule(
        operator="greaterThan", formula=["70"], font=Font(color="FFC00000")))
    ws.conditional_formatting.add(f"N2:N{last}", CellIsRule(
        operator="lessThan", formula=["30"], font=Font(color="FF006100")))
    if extra:
        ws.column_dimensions["R"].width = 14


def _raw_fundamentals(wb, c, f_fy):
    ws = wb["Raw_Fundamentals"]
    v, m, r, l, g = (f_fy.get(k, {}) for k in ("valuation", "margins", "returns",
                                                "leverage", "growth"))
    q, fc = f_fy.get("quality", {}), f_fy.get("fcf", {})

    def pct(x):
        return x / 100 if _num(x) else None
    rows = {
        1: ("Fundamentals & Valuation   ·   source: SEC EDGAR   ·   FY ending "
            f"{_d(c.fy_end, '%d %b %Y')}", None),
        3: ("Valuation", None), 4: ("P/E", v.get("pe")), 5: ("P/S", v.get("ps")),
        6: ("P/B", v.get("pb")), 7: ("EV / EBITDA", f_fy.get("ev_ebitda")),
        8: ("Earnings Yield", pct(v.get("earnings_yield"))),
        9: ("FCF Yield", pct(fc.get("fcf_yield"))),
        10: ("Implied growth — earnings basis, flat 9% discount", f_fy.get("implied_growth")),
        11: ("  (cross-check only; the headline reverse DCF is on the DCF tab)", None),
        13: ("Profitability & Returns", None), 14: ("Gross Margin", pct(m.get("gross"))),
        15: ("Operating Margin", pct(m.get("operating"))), 16: ("Net Margin", pct(m.get("net"))),
        17: ("Return on Equity", pct(r.get("roe"))), 18: ("Return on Assets", pct(r.get("roa"))),
        19: ("Free Cash Flow", fc.get("fcf")),
        21: ("Growth", None), 22: ("Revenue YoY", pct(g.get("revenue_yoy"))),
        23: ("EPS YoY", pct(g.get("eps_yoy"))), 24: ("Revenue CAGR", pct(g.get("revenue_cagr"))),
        25: ("EPS CAGR", pct(g.get("eps_cagr"))),
        27: ("Balance Sheet & Quality", None), 28: ("Current Ratio", l.get("current_ratio")),
        29: ("Debt / Equity", l.get("debt_to_equity")),
        30: ("Piotroski F-Score", f"{q['f_score']} / 9" if q.get("f_score") is not None else None),
        31: ("Altman Z-Score", f"{q['z_score']} ({q['z_zone']})" if q.get("z_score") is not None else None),
    }
    for row, (label, val) in rows.items():
        ws.cell(row=row, column=1, value=label)
        if row not in (1, 3, 11, 13, 21, 27):
            ws.cell(row=row, column=2, value=_na(val))
    tr = f_fy.get("trend") or {}
    per = list(tr.get("periods") or [])[-10:]
    pad = 10 - len(per)
    ws.cell(row=33, column=1, value="Fiscal Period")
    for label, key, rr in (("Revenue ($B)", "revenue", 34), ("Net Income ($B)", "net_income", 35),
                           ("Free Cash Flow ($B)", "fcf", 36)):
        ws.cell(row=rr, column=1, value=label)
        vals = list(tr.get(key) or [])[-10:]
        for k in range(10):
            cell = ws.cell(row=rr, column=2 + k)
            idx = k - pad
            cell.value = (_bn(vals[idx]) if (0 <= idx < len(vals) and _num(vals[idx]))
                          else "N/A")
    for k in range(10):
        idx = k - pad
        ws.cell(row=33, column=2 + k, value=per[idx] if 0 <= idx < len(per) else "N/A")
    return per, pad


def _raw_table(wb, name, headers, rows):
    ws = wb[name]
    hdr = copy.copy(ws["A1"]._style)
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell._style = copy.copy(hdr)
    for i, rec in enumerate(rows, 2):
        for j, v in enumerate(rec, 1):
            ws.cell(row=i, column=j, value=v)


def _monte_carlo(wb, c):
    ws = wb["Monte_Carlo"]
    ws["A1"] = "Monte Carlo Simulation Summary"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color=NAVY)
    ws["A2"], ws["B2"] = "Field", "Value"
    for cell in (ws["A2"], ws["B2"]):
        cell.font = Font(name="Calibri", bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)

    def val(v):
        if isinstance(v, str) and v.strip().endswith("%"):
            try:
                return float(v.strip().rstrip("%")) / 100
            except ValueError:
                return v
        return v
    r = 3
    for k, v in (c.mc or {}).items():
        ws.cell(row=r, column=1, value=k)
        cell = ws.cell(row=r, column=2, value=val(v))
        if isinstance(cell.value, float) and isinstance(v, str) and v.endswith("%"):
            cell.number_format = "0.0%"
        elif isinstance(cell.value, float) and k not in ("Forecast Horizon (days)", "Simulations"):
            cell.number_format = "#,##0.00"
        r += 1
    mc_df = c.mc_df
    if mc_df is not None and len(mc_df):
        start = max(r + 2, 19)
        ws.cell(row=start, column=1,
                value=f"Percentile price paths across all {mc_df.shape[1]:,} simulations").font = \
            Font(name="Calibri", italic=True, color="FF595959")
        hdrs = ["Day", "P5 (Bear)", "P25 (Low)", "P50 (Median)", "P75 (Bull)", "P95 (Best)"]
        for j, h in enumerate(hdrs, 1):
            cell = ws.cell(row=start + 1, column=j, value=h)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
        arr = np.asarray(mc_df, dtype=float)
        pct = np.percentile(arr, [5, 25, 50, 75, 95], axis=1).T
        for d in range(arr.shape[0]):
            ws.cell(row=start + 2 + d, column=1, value=d)
            for j in range(5):
                cell = ws.cell(row=start + 2 + d, column=2 + j, value=round(float(pct[d, j]), 2))
                cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16


# ── Financials ────────────────────────────────────────────────────────────────
def _balance(frame_row, get):
    """(cash, current securities, non-current securities, CP, term debt, equity,
    assets, total debt) in $B from one balance sheet."""
    cash, sti, lts = get("cash"), get("short_term_investments"), get("lt_securities")
    cp = get("commercial_paper")
    total = get("debt_total_incl_current")
    if total is None:
        dct = get("debt_current_total")
        if dct is None:
            parts = [x for x in (get("debt_current"), cp, get("short_term_borrowings")) if x is not None]
            dct = sum(parts) if parts else None
        ltd = get("long_term_debt")
        total = (sum(x for x in (ltd, dct) if x is not None)
                 if (ltd is not None or dct is not None) else None)
    term = (total - (cp or 0.0)) if total is not None else None
    return {"cash": _bn(cash), "sti": _bn(sti), "lts": _bn(lts), "cp": _bn(cp),
            "term": _bn(term), "equity": _bn(get("equity")), "assets": _bn(get("assets")),
            "total": _bn(total)}


def _financials(wb, c, per, pad):
    ws = wb["Financials"]
    ws["A1"] = f"{c.name} ({c.t}) — Historical Financials"
    ws["A2"] = (f"$ in billions except per-share data · {c.short}'s fiscal year ends "
                f"{c.fy_desc()}")
    _back_link(ws, "L2")
    labels = []
    for k in range(10):
        idx = k - pad
        lab = f"FY{str(per[idx])[:4]}" if 0 <= idx < len(per) else "N/A"
        labels.append(lab)
        ws.cell(row=5, column=2 + k, value=lab)
    ws["L5"], ws["M5"] = "9-yr CAGR", "5-yr CAGR"
    # A missing year (a tag the filer did not use, a loss) reads N/A in every
    # cell that depends on it, never #VALUE! or #NUM!.
    raw = wb["Raw_Fundamentals"]
    have = {rr: [_num(raw.cell(row=rr, column=2 + k).value) for k in range(10)] for rr in (34, 35, 36)}
    pos = {rr: [_num(raw.cell(row=rr, column=2 + k).value) and raw.cell(row=rr, column=2 + k).value > 0
                for k in range(10)] for rr in (34, 35, 36)}
    need = {7: [(34, 0), (34, -1)], 9: [(35, 0), (34, 0)], 11: [(36, 0), (34, 0)], 12: [(36, 0), (35, 0)]}
    for rr, deps in need.items():
        for k in range(10):
            ok = all(0 <= k + off < 10 and have[src][k + off] for src, off in deps)
            if not ok and not (rr == 7 and k == 0):
                ws.cell(row=rr, column=2 + k).value = "N/A"
    for rr, src in ((6, 34), (8, 35), (10, 36)):
        if not (pos[src][0] and pos[src][9]):
            ws[f"L{rr}"] = "N/A"
        if not (pos[src][4] and pos[src][9]):
            ws[f"M{rr}"] = "N/A"
    if pad:
        # Fewer than ten fiscal years filed: CAGRs over what exists, never #VALUE!.
        span = 9 - pad
        first = get_column_letter(2 + pad)
        ws["L5"] = f"{span}-yr CAGR"
        for rr in (6, 8, 10):
            ws[f"L{rr}"] = f'=IFERROR((K{rr}/{first}{rr})^(1/{span})-1,"N/A")'
            ws[f"M{rr}"] = f'=IFERROR((K{rr}/F{rr})^(1/5)-1,"N/A")'
        for k in range(pad):
            col = get_column_letter(2 + k)
            for rr in (7, 9, 11, 12):
                ws[f"{col}{rr}"] = None
        for k in range(pad, pad + 1):
            ws[f"{get_column_letter(2 + k)}7"] = None
    # latest four quarters
    ws["A14"] = f"Latest four quarters — {c.short} 10-Q / 10-K filings ($B except per share)"
    q = c.quarters[-4:]
    notes = []
    for k in range(4):
        col = get_column_letter(2 + k)
        if k < len(q):
            r = q[k]
            src = c.src_note(f"quarter ended {_d(r['end'])}", r["end"])
            ws[f"{col}15"] = r["label"]
            ws[f"{col}16"] = pd.Timestamp(r["end"]).to_pydatetime()
            _put(ws, f"{col}17", _bn(r["revenue"]), note=src)
            _put(ws, f"{col}18", _na(r.get("rev_yoy")), note=src)
            _put(ws, f"{col}19", _na(r.get("eps")), note=src)
            _put(ws, f"{col}20", _na(r.get("eps_yoy")), note=src)
        else:
            for rr in range(15, 21):
                ws[f"{col}{rr}"] = "N/A"
    ws["F15"] = "TTM"
    if q:
        ws["F16"] = pd.Timestamp(q[-1]["end"]).to_pydatetime()
    derived = [r["label"] for r in q if not r.get("eps_exact")]
    if derived:
        notes.append(f"{', '.join(derived)} EPS is the full year less the first nine months "
                     f"(the 10-K reports the year, not the quarter).")
    typ = c.typical_tax
    for r in q:
        ep = r.get("etr_prev")
        if typ is not None and ep is not None and abs(ep - typ) > 0.05 and r.get("eps_yoy") is not None:
            notes.append(f"{r['label']} EPS growth compares with {r.get('prev_label')}, which was "
                         f"taxed at {ep:.0%} against a usual {typ:.0%} — a one-off tax item "
                         f"in the base quarter.")
    if q and q[-1].get("gross_profit") and q[-1].get("revenue"):
        notes.append(f"{q[-1]['label']} gross margin: {q[-1]['gross_profit'] / q[-1]['revenue']:.1%}.")
    ws["A21"] = " ".join(notes) or "Growth compares each quarter with the same quarter a year earlier, as filed."
    # balance sheet & other items
    ws["A23"] = "Balance sheet & other items ($B)"
    fyb = _balance(None, lambda k: (c.bal.iloc[0].get(k) if len(c.bal) and k in c.bal.columns
                                    and _num(c.bal.iloc[0].get(k)) else None))
    tb = (c.ttm.get("balance") or {})
    latest = _balance(None, lambda k: tb.get(k) if _num(tb.get(k)) else None) if tb else {}
    same = c.bal_end is None or c.fy_end is None or c.bal_end <= c.fy_end
    ws["B24"] = f"{c.fys} · {_d(c.fy_end)} (10-K)"
    ws["C24"] = ("TTM · " if c.has_ttm else "") + (f"{_d(c.bal_end)} (10-Q)" if not same else "n/a")
    fy_src = c.src_note("balance sheet", c.fy_end, "10-K")
    q_src = c.src_note("balance sheet", c.bal_end, "10-Q")
    for rr, key, label in ((25, "cash", "Cash & cash equivalents"),
                           (26, "sti", "Current marketable securities"),
                           (27, "lts", "Non-current marketable securities"),
                           (29, "cp", "Commercial paper"),
                           (30, "term", "Term debt (notes, carrying amount)"),
                           (33, "equity", "Shareholders' equity"),
                           (34, "assets", "Total assets")):
        ws[f"A{rr}"] = label
        bv = fyb.get(key)
        _put(ws, f"B{rr}", bv if bv is not None else (0.0 if key in ("sti", "lts", "cp") else "N/A"),
             note=fy_src if bv is not None else None)
        lv = latest.get(key) if not same else None
        _put(ws, f"C{rr}", lv if lv is not None else ((0.0 if key in ("sti", "lts", "cp") else "n/a")
                                                       if not same else "n/a"),
             note=q_src if lv is not None else None, like=ws[f"B{rr}"])
    ws["A28"], ws["A31"] = "Total cash & marketable securities", "Total debt"
    ws["A32"] = "Net cash / (net debt) — all securities"
    ws["B28"], ws["B31"], ws["B32"] = "=SUM(B25:B27)", "=B29+B30", "=B28-B31"
    for a in ("C28", "C31", "C32"):
        _copy_style(ws["B" + a[1:]], ws[a])
    ws["C28"], ws["C31"], ws["C32"] = "=SUM(C25:C27)", "=C29+C30", "=C28-C31"
    if same:
        ws["C28"] = ws["C31"] = ws["C32"] = "n/a"
    # flows: last fiscal year in B, latest twelve months in C
    cf0 = c.cf.iloc[0] if len(c.cf) else {}
    inc0 = c.inc.iloc[0] if len(c.inc) else {}
    ttm_cf, ttm_inc = c.ttm.get("cash_flow") or {}, c.ttm.get("income") or {}
    fy_cfo, fy_sbc, fy_da = (_bn(cf0.get("net_cash_flow_from_operating_activities")),
                             _bn(cf0.get("sbc")), _bn(inc0.get("depreciation_amortization")))
    ws["A35"], ws["A36"] = "Cash from operations", "Capital expenditures"
    ws["A37"], ws["A38"] = "Share-based compensation", "FCF after SBC"
    ws["A39"], ws["A40"] = "Depreciation & amortisation", "Operating income"
    ws["A41"] = ("Shares outstanding (B) — "
                 + (f"{_d(c.f.get('shares_date'))} 10-Q cover" if c.f.get("market_cap_basis") == "filing"
                    else "market cap ÷ price"))
    _put(ws, "B35", _na(fy_cfo), note=c.src_note("cash flow statement", c.fy_end, "10-K"))
    ws["B36"], ws["B38"] = "=B35-K10", "=K10-B37"
    _put(ws, "B37", _na(fy_sbc), note=c.src_note("cash flow statement", c.fy_end, "10-K"))
    _put(ws, "B39", _na(fy_da), note=c.src_note("income / cash flow statement", c.fy_end, "10-K"))
    ws["B40"] = "=Raw_Fundamentals!B15*K6"
    shares = c.f.get("shares_outstanding") or (c.f.get("market_cap") / c.price
                                               if c.f.get("market_cap") and c.price else None)
    _put(ws, "B41", _na(_bn(shares)),
         note=(c.src_note(f"cover page, shares outstanding as of {_d(c.f.get('shares_date'))}",
                          None, "10-Q") if c.f.get("market_cap_basis") == "filing"
               else "Source: vendor market capitalisation ÷ report price (no current cover-page count)"))
    if c.has_ttm:
        tsrc = c.src_note(f"four quarters to {_d(c.flows_end)} (10-Q/10-K cash flow)", c.flows_end, "10-Q")
        for rr, v in ((35, ttm_cf.get("net_cash_flow_from_operating_activities")),
                      (36, ttm_cf.get("capex")), (37, ttm_cf.get("sbc")),
                      (39, ttm_inc.get("depreciation_amortization")),
                      (40, ttm_inc.get("operating_income_loss"))):
            _put(ws, f"C{rr}", _na(_bn(v)), note=tsrc if v is not None else None, like=ws["B37"])
        _put(ws, "C38", "=C35-C36-C37", like=ws["B38"])
    from openpyxl.styles import Alignment
    for rr in range(25, 41):
        cell = ws[f"C{rr}"]
        cell.alignment = Alignment(horizontal="right")
    ws["A43"] = ("Annual figures link to the SEC EDGAR pull on Raw_Fundamentals; quarterly, balance-"
                 "sheet and TTM figures are from the XBRL in each 10-Q / 10-K (hover for the filing). "
                 "Blue = hard-coded source input · Green = link to another tab · Black = formula.")
    return labels


# ── DCF ───────────────────────────────────────────────────────────────────────
def _dcf(wb, c):
    ws = wb["DCF"]
    d = c.dcf
    ws["A1"] = f"{c.name} ({c.t}) — Discounted Cash Flow Valuation"
    ws["A2"] = ("Two-stage free-cash-flow DCF · $ in billions except per share · "
                "blue/yellow cells are editable inputs")
    _back_link(ws, "G2")
    ws["B5"] = f"=Price_Data!E{c.last_row}"
    ws["C5"] = (f"Report price: close of {_d(c.last_date)}"
                + (f" ({c.src})" if c.src else "")
                + (". The final bar has no volume — confirm the official close." if c.last_bar_no_volume else "."))
    ws["C6"] = ((f"Shares outstanding per 10-Q cover ({_d(c.f.get('shares_date'))}). "
                 "Diluted count would be slightly higher.")
                if c.f.get("market_cap_basis") == "filing"
                else "Market cap ÷ price (no current 10-Q cover count).")
    wb_ = d.get("wacc_basis") or {}
    ws["B7"] = float(d.get("wacc") or 0.09)
    if wb_ and not wb_.get("fallback") and wb_.get("beta") is not None:
        ws["C7"] = (f"CAPM: {wb_['risk_free']:.2%} 10-yr Treasury + {wb_['beta_adjusted']:.2f} "
                    f"adjusted beta × {wb_['erp']:.1%} ERP = {wb_['cost_of_equity']:.1%} cost of "
                    f"equity, blended with {wb_['cost_of_debt'] * (1 - wb_['tax_rate']):.1%} after-tax "
                    f"debt at a {wb_['debt_weight']:.0%} weight. Cross-check below.")
    else:
        ws["C7"] = ("Default cost of capital — beta could not be measured (no benchmark data). "
                    "Cross-check: implied equity risk premium at beta 1.0 shown below.")
    ws["B8"] = float(d.get("terminal_growth") or 0.025)
    ws["C8"] = ("Perpetual growth after year 10. Must be below WACC; ~nominal GDP growth "
                "or less.")
    ws["A9"] = "Stage-1 FCF growth (year 1)"
    g_base = d.get("base_growth")
    cagr_ok = c.fcf_cagr_9 is not None and g_base is not None and abs(c.fcf_cagr_9 - g_base) < 1e-9
    if cagr_ok:
        ws["B9"] = "=Financials!L10"
        ws["C9"] = (f"Defaults to the {c.span}-yr historical FCF CAGR; fades linearly to terminal "
                    f"growth by year 10. Overwrite to test.")
    else:
        _put(ws, "B9", float(g_base) if g_base is not None else 0.05)
        ws["B9"].font = Font(name="Calibri", size=10, color="FF0000FF")
        ws["C9"] = ("Historical FCF CAGR bounded to a sane range (or revenue CAGR where FCF was "
                    "negative); fades linearly to terminal growth by year 10. Overwrite to test.")
    opts = [f"{c.fys} reported"] + (["TTM"] if c.has_ttm else []) + ["3-yr average"]
    basis = d.get("base_fcf_basis")
    ws["B10"] = {"ttm": "TTM", "fy_avg3": "3-yr average"}.get(basis, f"{c.fys} reported")
    ws.data_validations.dataValidation = [dv for dv in ws.data_validations.dataValidation
                                          if "B10" not in str(dv.sqref)]
    dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=False)
    dv.add("B10")
    ws.add_data_validation(dv)
    ws["C10"] = ("Choose: " + " / ".join(
        [f"{c.fys} reported"] + ([f"TTM (4 qtrs to {_d(c.flows_end, '%b-%y')})"] if c.has_ttm else [])
        + [f"3-yr average (FY{c.fy - 2 - 2000:02d}–{c.fys})" if c.fy else "3-yr average"]) + ".")
    fy_sbc = c.fy_sbc
    ws["C11"] = ("Yes = treat SBC as a real cost (more conservative"
                 + (f"; {c.fys} SBC ${fy_sbc:,.1f}B)." if fy_sbc else ")."))
    lts = c.lts_latest
    ws["C12"] = ((f"{c.short}'s ${lts:,.1f}B of non-current marketable securities are liquid "
                  f"holdings; many companies count them in their own net cash.")
                 if lts else "No non-current marketable securities reported.")
    ws["A14"] = "Risk-free rate (3-month T-bill)"
    ws["B14"] = float(c.rf)
    # base FCF and net debt: the fiscal year, the latest twelve months, or three years
    fy_col = "B" if c.bal_same else "C"
    ws["B17"] = ('=IF(B10="3-yr average",AVERAGE(Financials!I10:K10),'
                 + ('IF(B10="TTM",Financials!C35-Financials!C36,Financials!K10))' if c.has_ttm
                    else 'Financials!K10)')
                 + '-IF(B11="Yes",' + ('IF(B10="TTM",Financials!C37,Financials!B37)' if c.has_ttm
                                       else "Financials!B37") + ",0)")
    ws["B18"] = (f"=Financials!{fy_col}31-Financials!{fy_col}25-Financials!{fy_col}26"
                 f'-IF(B12="Yes",Financials!{fy_col}27,0)')
    ws["C18"] = (f"Negative = net cash (added to equity value). Balance sheet at "
                 f"{_d(c.bal_end if not c.bal_same else c.fy_end)} "
                 f"({'10-Q' if not c.bal_same else '10-K'}).")
    ws["A20"] = "Implied equity risk premium at beta 1.0"
    ws["C20"] = "WACC − risk-free. 4–6% is a typical range."
    ws["A22"] = "10-year free cash flow projection ($B)"
    ws["J23"] = "Year-1 growth"
    for i in range(10):
        r = 24 + i
        ws[f"A{r}"] = i + 1
        ws[f"B{r}"] = f'="FY"&({c.fy}+A{r})'
    k_tmpl = ws["K24"].value.text if hasattr(ws["K24"].value, "text") else str(ws["K24"].value)
    k_style = copy.copy(ws["K25"]._style)
    j_style = copy.copy(ws["J25"]._style)
    for i in range(161):                          # -20.0% .. +60.0%
        r = 24 + i
        ws[f"J{r}"] = -0.20 if i == 0 else f"=J{r - 1}+0.005"
        ws[f"J{r}"]._style = copy.copy(j_style)
        _array(ws, f"K{r}", k_tmpl.replace("J24", f"J{r}"))
        ws[f"K{r}"]._style = copy.copy(k_style)
    for addr in ("B49", "B50"):
        f0 = ws[addr].value.text if hasattr(ws[addr].value, "text") else str(ws[addr].value)
        f0 = f0.replace("K24:K154", "K24:K184").replace("J24:J154", "J24:J184")
        _array(ws, addr, '=IFERROR(' + f0.lstrip("=") + ',"n/a")')
    ws["B51"] = '=IFERROR(B17*(1+B50)^10,"n/a")'
    ws["B54"] = '=IFERROR(B50-B52,"n/a")'
    ws["C49"] = ("Solved so the DCF equals today's price, holding WACC, terminal growth and base "
                 "FCF constant." if c.dcf.get("market_implied_growth") is not None else
                 "Today's price is outside the -20% to +60% range the model can solve; no single "
                 "year-1 rate reproduces it.")
    ws["C37"] = "Gordon growth on year-11 FCF."
    ws["A37"] = "Terminal value at year 10 ($B)"
    ws["C38"] = "Discounted a full 10 years."
    ws["C45"] = "Above ~75% means the valuation rests mostly on the perpetuity assumption."
    ws["C46"] = "Should be 0.00."
    ws["A49"] = "Implied year-1 FCF growth (fading to terminal)"
    ws["A50"] = "Equivalent 10-year FCF CAGR"
    ws["C50"] = ("The single annual rate that produces the same year-10 FCF — the number to "
                 "compare with history.")
    ws["A51"] = f"Implied FY{c.fy + 10} free cash flow ($B)"
    ws["C51"] = (f"Compare with {c.fys} FCF of ~${c.fy_fcf:,.0f}B." if c.fy_fcf else "")
    ws["A52"] = f"Historical {c.span}-yr FCF CAGR ({c.first_fy_short}–{c.fys})"
    ws["A53"] = "Latest 3 quarters avg. revenue growth y/y"
    ws["C53"] = (f"{c.quarters[-3]['label']}–{c.quarters[-1]['label']} as filed."
                 if len(c.quarters) >= 3 else "")
    ws["A56"] = ("Notes: FCF is cash from operations less capital expenditures, i.e. after interest "
                 "and tax (not a strict unlevered FCF). Balance-sheet items are as of "
                 f"{_d(c.bal_end if not c.bal_same else c.fy_end)}. Valuation date ≈ fiscal "
                 f"year-end, so year 1 = FY{c.fy + 1}. Same model as the QuantWizard site.")


# ── Scenarios ────────────────────────────────────────────────────────────────
def _scenarios(wb, c):
    ws = wb["Scenarios"]
    s = c.dcf.get("scenarios") or {}
    bear, base, bull = s.get("bear") or {}, s.get("base") or {}, s.get("bull") or {}
    ws["A1"] = f"{c.name} ({c.t}) — Scenarios & Sensitivity"
    ws["A2"] = "Every value on this tab is live: it recalculates from the DCF tab inputs · $ per share"
    _back_link(ws, "G2")
    ws["A6"] = "Year-1 FCF growth (fades to terminal)"
    for col, sc in (("B", bear), ("D", bull)):
        ws[f"{col}6"] = float(sc.get("growth", 0.0))
        ws[f"{col}7"] = float(sc.get("wacc", c.dcf.get("wacc", 0.09)))
        ws[f"{col}8"] = float(sc.get("terminal_growth", c.dcf.get("terminal_growth", 0.025)))
    ws["B9"], ws["C9"], ws["D9"] = (float(bear.get("probability") or 0.25),
                                    float(base.get("probability") or 0.5),
                                    float(bull.get("probability") or 0.25))
    g, tg = c.dcf.get("base_growth") or 0.0, c.dcf.get("terminal_growth") or 0.025
    rg = (c.f.get("growth") or {}).get("revenue_yoy")
    ws["B14"] = (f"Growth normalises: FCF rises {bear.get('growth', 0):.1%} in year one, fading to "
                 f"{bear.get('terminal_growth', tg):.1%}; a {bear.get('wacc', 0):.1%} discount rate "
                 f"prices in more risk.")
    ws["C14"] = (f"FCF compounds near its {c.span}-year historical rate (~{g:.0%}), fading to "
                 f"{tg:.1%} long-run growth.")
    ws["D14"] = (f"Recent momentum holds: year-one FCF growth of {bull.get('growth', 0):.0%}"
                 + (f" (revenue grew {rg:+.0f}% over the latest twelve months)" if _num(rg) else "")
                 + f", a {bull.get('wacc', 0):.1%} discount rate and "
                   f"{bull.get('terminal_growth', tg):.1%} terminal growth.")
    ws["A34"] = ('="Yellow = base case. Green = fair value at or above today\'s price of "&TEXT(DCF!B5,"$0.00")'
                 '&". The base case assumes "&TEXT(DCF!B9,"0.0%")&" year-1 growth; "&IFERROR("the market '
                 'price implies "&TEXT(DCF!B49,"0.0%")&" (≈"&TEXT(DCF!B50,"0.0%")&" 10-yr CAGR).","the '
                 'market price is outside the range the model can solve.")')
    ws["A18"] = ("Sensitivity 1 — fair value per share: WACC (rows) × terminal growth (columns), "
                 "base-case growth")
    ws["A26"] = "Sensitivity 2 — fair value per share: year-1 FCF growth (rows) × WACC (columns)"
    s0, e0 = c.wk52_rows
    ws["A42"], ws["B42"] = "52-week low", f"=MIN(Price_Data!D{s0}:D{e0})"
    ws["A43"], ws["B43"] = "52-week high", f"=MAX(Price_Data!C{s0}:C{e0})"
    # The note under the grids says green means at or above the price; make it true.
    for rng in ("C20:G24", "C28:G32"):
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"{rng.split(':')[0]}>=DCF!$B$5"],
            fill=PatternFill("solid", bgColor="FFC6EFCE"), font=Font(color="FF006100")))


# ── Multiples ────────────────────────────────────────────────────────────────
def _multiples(wb, c):
    ws = wb["Multiples"]
    ws["A1"] = f"{c.name} ({c.t}) — Multiples & Quality Metrics"
    ws["A2"] = (f"FY{c.fy} = last audited fiscal year (ended {_d(c.fy_end)})"
                + (f" · TTM = four quarters to {_d(c.flows_end)}" if c.has_ttm else "")
                + " · $ in billions except per share")
    _back_link(ws, "D2")
    ws["C5"] = f"Report price, {_d(c.last_date)}"
    ws["C6"] = (f"10-Q cover, {_d(c.f.get('shares_date'))}" if c.f.get("market_cap_basis") == "filing"
                else "Market cap ÷ price")
    ws["C7"] = "Price × shares — one price, one share count, one moment"
    dps, dps_end = c.ttm.get("dps_quarter"), c.ttm.get("dps_quarter_end")
    if dps is not None:
        _put(ws, "B10", float(dps), note=c.src_note("dividends declared per share", dps_end))
        ws["C10"] = f"Declared for the quarter to {_d(dps_end, '%b-%Y')}"
    else:
        _put(ws, "B10", 0.0)
        ws["C10"] = "No dividend declared in the latest two quarters"
    ws["C11"] = "Quarterly dividend × 4 ÷ price"
    ws["B14"] = f"FY{c.fy}"
    ws["C14"] = f"TTM ({_d(c.flows_end, '%b-%y')})" if c.has_ttm else "TTM"
    ws["D15"] = "FY: market cap ÷ net income · TTM: price ÷ diluted EPS (sum of last 4 quarters)"
    if c.has_ttm:
        ws["C18"] = "=B9/(Financials!C40+Financials!C39)"
        ws["C19"] = "=B7/(Financials!C35-Financials!C36)"
        ws["C20"] = "=(Financials!C35-Financials!C36)/B7"
        ws["C21"] = "=Financials!C38/B7"
        for a in ("C18", "C19"):
            ws[a].number_format = F_X
        for a in ("C20", "C21"):
            ws[a].number_format = F_PCT
    if not c.bal_same:
        ws["C24"] = "=B7/Financials!C33"
        ws["C24"].number_format = F_X
    ws["D18"] = "EBITDA = operating income + D&A"
    ws["D23"] = f"Equity 'spread' over T-bills (risk-free {c.rf:.2%})"
    ws["D24"] = "Book equity; distorted where buybacks have shrunk equity"
    ws["A26"] = f"Profitability, cash quality & balance sheet (FY{c.fy} unless stated)"
    lq = c.quarters[-1] if c.quarters else None
    if lq and lq.get("gross_profit") and lq.get("revenue"):
        _put(ws, "C27", lq["gross_profit"] / lq["revenue"],
             note=c.src_note(f"gross margin, quarter ended {_d(lq['end'])}", lq["end"]))
        ws["D27"] = f"Latest quarter ({lq['label']})"
    else:
        ws["C27"] = "n/a"
    tr = c.f_fy.get("trend") or {}
    fcf_s, ni_s = list(tr.get("fcf") or []), list(tr.get("net_income") or [])
    conv = [(f / n) if (_num(f) and _num(n) and n) else None for f, n in zip(fcf_s[-2:], ni_s[-2:])]
    if len(conv) == 2 and all(x is not None for x in conv):
        a, b = conv[1], conv[0]
        ws["D31"] = (f"{'Below' if a < 1 else 'Above'} 1.0x in {c.fys} after "
                     f"{'>' if b > 1 else '<'}1.0x in FY{str(c.fy - 1)[2:]}")
    # peers - the reference held a one-line note here in a merged cell
    row = 38
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row > row:
            ws.unmerge_cells(str(mr))
    ws[f"A{row}"] = "Peer comparison"
    _clear_heights(ws, row + 1)
    if c.peers and len(c.peers) > 1:
        # Compact here, in the sheet's own columns (A label, B/C values, D the
        # wide definition column); the full table is on the Peers tab.
        from analysis import peer_median
        heads = [("A", "Ticker"), ("B", "P/E (TTM)"), ("C", "EV / EBITDA"), ("D", "Revenue growth · FCF yield")]
        for col, h in heads:
            _put(ws, f"{col}{row + 1}", h, like=ws[f"{col}14"])
        med = peer_median(c.peers, skip=c.t)
        for i, rec in enumerate(list(c.peers) + [dict(med, ticker="Peer median")]):
            r = row + 2 + i
            _put(ws, f"A{r}", rec.get("ticker"), like=ws["A15"])
            _put(ws, f"B{r}", _na(rec.get("pe")), like=ws["B15"], fmt=F_X)
            _put(ws, f"C{r}", _na(rec.get("ev_ebitda")), like=ws["B15"], fmt=F_X)
            rg, fy_ = rec.get("rev_growth"), rec.get("fcf_yield")
            _put(ws, f"D{r}", (f"{rg:+.1f}% · {fy_:.1f}%" if _num(rg) and _num(fy_)
                               else f"{rg:+.1f}% · n/a" if _num(rg) else "n/a"), like=ws["D15"])
            if rec.get("ticker") in (c.t, "Peer median"):
                for col in "ABCD":
                    ws[f"{col}{r}"].font = Font(name="Calibri", size=11, bold=True)
        note = ws[f"A{row + 3 + len(c.peers)}"]
        note.value = (f"Peers: {c.peer_group or 'selected by industry and size'} — each on its own latest "
                      f"filings and market cap. Full table on the Peers tab.")
        note.font = Font(name="Calibri", size=9, italic=True, color="FF595959")
    else:
        ws[f"A{row + 1}"] = "No peer set was available for this company."


# ── Risk ─────────────────────────────────────────────────────────────────────
def _risk(wb, c):
    ws = wb["Risk"]
    L = c.last_row
    # prototype styles, from the template rows the reference used
    snap = _snapshot(ws, ["A4", "A5", "B5", "C6", "B7", "B8", "B10", "B12", "B15", "B18", "B21",
                          "B22", "B24", "B25", "B26", "F4", "F5", "G5", "H5", "F6", "G6", "H6",
                          "F31", "H31", "F35", "A36", "A37", "B37", "C37", "D37", "A38", "B38",
                          "C38", "D38", "A42", "A43", "B43", "B46", "B47", "B49", "C47", "A50",
                          "B29", "B31", "B34", "B11"])
    # Rows 4-34 are fixed and keep the template's styles (notes column
    # included); the monthly table (F:H) and everything below row 35 vary with
    # the window and are rebuilt.
    for rr in range(5, 35):
        for col in "ABC":
            ws[f"{col}{rr}"].value = None
            ws[f"{col}{rr}"].comment = None
    _reset(ws, 6, max(ws.max_row, 60), 6, 9)
    _reset(ws, 35, max(ws.max_row, 60), 1, 5)
    ws["A1"] = f"{c.name} ({c.t}) — Price Performance & Risk"
    basis = ("total return — Yahoo Finance closes are dividend-adjusted" if c.src == "Yahoo Finance"
             else "price-only — dividends excluded" if c.src == "Polygon.io"
             else "adjusted for splits")
    ws["A2"] = f"All statistics computed live from Price_Data (daily closes, {basis}) · window shown below"
    _back_link(ws, "H2")

    def lab(r, text, sty="A5"):
        _put(ws, f"A{r}", text)
        _style(ws, f"A{r}", snap[sty])

    def val(r, formula, sty="B5", fmt=None):
        _put(ws, f"B{r}", formula)
        _style(ws, f"B{r}", snap[sty])
        if fmt:
            ws[f"B{r}"].number_format = fmt

    def note(r, text):
        if text:
            _put(ws, f"C{r}", text)
            _style(ws, f"C{r}", snap["C6"])

    lab(4, "Performance window", "A4")
    lab(5, "First observation"); val(5, f"=MIN(Price_Data!A2:A{L})", "B5")
    lab(6, "Last observation"); val(6, f"=MAX(Price_Data!A2:A{L})", "B5")
    note(6, "Final bar had no volume — may be an intraday snapshot" if c.last_bar_no_volume else None)
    lab(7, "Length of window (years)"); val(7, "=(B6-B5)/365.25", "B7")
    note(7, f"Report period: {c.period_label}" if c.period_label else None)
    lab(8, "Starting close ($)"); val(8, "=Price_Data!E2", "B8")
    lab(9, "Latest close ($)"); val(9, f"=Price_Data!E{L}", "B8")
    lab(10, "Price return over window"); val(10, "=B9/B8-1", "B10")
    dy = (c.f.get("capital_return") or {}).get("dividend_yield")
    note(10, ("Includes dividends (adjusted closes)" if c.src == "Yahoo Finance"
              else f"Excludes ~{dy / 100:.1%}/yr dividend yield" if _num(dy) and dy > 0 else None))
    lab(11, "Annualised return (CAGR)"); val(11, "=(1+B10)^(1/B7)-1", "B10")
    lab(12, "Trading days"); val(12, f"=COUNT(Price_Data!E2:E{L})", "B12")
    lab(14, "Risk statistics (window above)", "A4")
    lab(15, "Annualised volatility"); val(15, f"=_xlfn.STDEV.S(Price_Data!G3:G{L})*SQRT(252)", "B15")
    note(15, "Sample st. dev. of daily returns × √252")
    lab(16, "Arithmetic annualised return"); val(16, f"=AVERAGE(Price_Data!G3:G{L})*252", "B15")
    note(16, "Mean daily return × 252 (used in Sharpe)")
    lab(17, "Risk-free rate"); val(17, "=DCF!B14", "B15", fmt="0.00%")
    note(17, "3-month T-bill (FRED DGS3MO)")
    lab(18, "Sharpe ratio"); val(18, "=(B16-B17)/B15", "B18")
    lab(19, "Downside deviation (annualised)")
    val(19, f"=SQRT(SUMPRODUCT((Price_Data!G3:G{L}<0)*Price_Data!G3:G{L}^2)"
            f"/COUNT(Price_Data!G3:G{L}))*SQRT(252)", "B15")
    note(19, "Target return 0%")
    lab(20, "Sortino ratio"); val(20, "=(B16-B17)/B19", "B18")
    lab(21, "Maximum drawdown (peak to trough)"); val(21, f"=MIN(Price_Data!Q2:Q{L})", "B21")
    lab(22, "  Trough date"); val(22, f"=INDEX(Price_Data!A2:A{L},MATCH(B21,Price_Data!Q2:Q{L},0))", "B22")
    lab(23, "Current drawdown from peak"); val(23, f"=Price_Data!Q{L}", "B21")
    lab(24, "20-day volatility (latest)"); val(24, f"=Price_Data!K{L}", "B24")
    lab(25, "Beta vs S&P 500")
    if c.has_bench:
        val(25, f"=SLOPE(Price_Data!G3:G{L},Price_Data!R3:R{L})", "B18")
        note(25, "Regression of daily returns on SPY over the window (raw, not Blume-adjusted)")
    else:
        val(25, "N/A", "B25")
        note(25, "Benchmark series was not delivered by the data feed — beta and correlation "
                 "cannot be computed")
    lab(26, "RSI (14)"); val(26, f"=Price_Data!N{L}", "B26")
    lab(28, "Trend (latest)", "A4")
    s0, e0 = c.wk52_rows
    lab(29, "52-week high ($)"); val(29, f"=MAX(Price_Data!C{s0}:C{e0})", "B8"); note(29, "Intraday")
    lab(30, "52-week low ($)"); val(30, f"=MIN(Price_Data!D{s0}:D{e0})", "B8"); note(30, "Intraday")
    lab(31, "% below 52-week high"); val(31, "=B9/B29-1", "B10")
    lab(32, "50-day moving average ($)"); val(32, f"=Price_Data!I{L}", "B8")
    lab(33, "200-day moving average ($)"); val(33, f"=Price_Data!J{L}", "B8")
    lab(34, "Price vs 200-day average"); val(34, "=B9/B33-1", "B10")

    # monthly returns, F:H
    for col, sty in (("F", "F4"), ("G", "F4"), ("H", "F4")):
        _style(ws, f"{col}4", snap[sty])
    _put(ws, "F4", "Monthly price returns")
    for col, text in (("F", "Month"), ("G", "Month-end close"), ("H", "Return")):
        _put(ws, f"{col}5", text)
        _style(ws, f"{col}5", snap[f"{col}5"])
    # First row = EOMONTH(first date, 1), as the F6 formula computes it.
    months = pd.date_range(c.first_date + pd.offsets.MonthEnd(0) + pd.offsets.MonthEnd(1),
                           c.last_date + pd.offsets.MonthEnd(0), freq="ME")
    r = 6
    for i, _m in enumerate(months):
        _put(ws, f"F{r}", "=EOMONTH(B5,1)" if i == 0 else f"=EOMONTH(F{r - 1},1)")
        _style(ws, f"F{r}", snap["F6"])
        _put(ws, f"G{r}", f"=_xlfn.XLOOKUP(F{r}+0.99,Price_Data!$A$2:$A${L},Price_Data!$E$2:$E${L},,-1)")
        _style(ws, f"G{r}", snap["G6"])
        _put(ws, f"H{r}", (f"=G{r}/_xlfn.XLOOKUP(EOMONTH(F{r},-1)+0.99,Price_Data!$A$2:$A${L},"
                           f"Price_Data!$E$2:$E${L},,1)-1") if i == 0 else f"=G{r}/G{r - 1}-1")
        _style(ws, f"H{r}", snap["H6"])
        r += 1
    last_m = r - 1
    ws.conditional_formatting.add(f"H6:H{last_m}", ColorScaleRule(
        start_type="min", start_color="FFF8696B", mid_type="num", mid_value=0,
        mid_color="FFFFFFFF", end_type="max", end_color="FF63BE7B"))
    r += 1
    for text, formula in (("Best month", f"=MAX(H6:H{last_m})"), ("Worst month", f"=MIN(H6:H{last_m})"),
                          ("Positive months", f'=COUNTIF(H6:H{last_m},">0")&" of "&COUNT(H6:H{last_m})')):
        _put(ws, f"F{r}", text); _style(ws, f"F{r}", snap["F31"])
        _put(ws, f"H{r}", formula); _style(ws, f"H{r}", snap["H31"])
        r += 1
    r += 1
    _prior = c.first_date + pd.offsets.MonthEnd(0)
    _put(ws, f"F{r}", (f"First month is partial (starts {_d(c.first_date)} → "
                       f"{_d(months[0], '%b')} uses {_d(_prior, '%d-%b')} close as base)."
                       if len(months) else ""))
    _style(ws, f"F{r}", snap["F35"])

    # calendar years
    top = 36          # columns A:D - the monthly table runs alongside in F:H, whatever its length
    for col in "ABCD":
        _style(ws, f"{col}{top}", snap["A36"])
    lab(top, "Calendar-year returns", "A36")
    for col, text in (("A", "Year"), ("B", "Coverage"), ("C", "Price return"), ("D", "Ann. volatility")):
        _put(ws, f"{col}{top + 1}", text); _style(ws, f"{col}{top + 1}", snap[f"{col}37"])
    years = list(range(c.first_date.year, c.last_date.year + 1))
    rng_a, rng_e = f"Price_Data!$A$2:$A${L}", f"Price_Data!$E$2:$E${L}"
    first_row = top + 2
    for i, y in enumerate(years):
        rr = first_row + i
        _put(ws, f"A{rr}", y); _style(ws, f"A{rr}", snap["A38"])
        if i == 0 and c.first_date > pd.Timestamp(y, 1, 5):
            cov = '="Partial: from "&TEXT(B5,"d-mmm")'
        elif y == c.last_date.year and c.last_date < pd.Timestamp(y, 12, 24):
            cov = '="YTD to "&TEXT(B6,"d-mmm")'
        else:
            cov = "Full year"
        _put(ws, f"B{rr}", cov); _style(ws, f"B{rr}", snap["B38"])
        end_px = (f"_xlfn.XLOOKUP(DATE(A{rr},12,31)+0.99,{rng_a},{rng_e},,-1)"
                  if y != c.last_date.year else "B9")
        start_px = ("Price_Data!E2" if i == 0
                    else f"_xlfn.XLOOKUP(DATE(A{rr - 1},12,31)+0.99,{rng_a},{rng_e},,-1)")
        _put(ws, f"C{rr}", f"={end_px}/{start_px}-1"); _style(ws, f"C{rr}", snap["C38"])
        _array(ws, f"D{rr}", f"=_xlfn.STDEV.S(_xlfn._xlws.FILTER(Price_Data!$G$3:$G${L},"
                             f"YEAR(Price_Data!$A$3:$A${L})=A{rr}))*SQRT(252)")
        _style(ws, f"D{rr}", snap["D38"])
    c.ytd_cell = f"Risk!C{first_row + len(years) - 1}"

    # Monte Carlo
    mc = first_row + len(years) + 1
    for col in "ABCD":
        _style(ws, f"{col}{mc}", snap["A42"])
    lab(mc, "Monte Carlo illustration (QuantWizard simulation)", "A42")
    for i, (text, ref, sty) in enumerate((("Median 1-year price (P50)", "=Monte_Carlo!B7", "B43"),
                                          ("5th percentile (P5)", "=Monte_Carlo!B8", "B43"),
                                          ("95th percentile (P95)", "=Monte_Carlo!B11", "B43"),
                                          ("Probability of gain", "=Monte_Carlo!B12", "B46"),
                                          ("Drift assumption", "=Monte_Carlo!B14", "B47"),
                                          ("Starting price used", "=Monte_Carlo!B3", "B43"),
                                          ("Paths simulated", "=Monte_Carlo!B5", "B49")), 1):
        lab(mc + i, text, "A43")
        val(mc + i, ref, sty)
    basis_txt = (c.mc.get("Drift basis") or "").strip()
    note(mc + 5, basis_txt or None)
    if _num(c.mc.get("Last Price")) and abs(float(c.mc["Last Price"]) - c.price) > 0.005:
        note(mc + 6, "Differs slightly from the report price (separate data pull)")
    _put(ws, f"A{mc + 8}", "A random-walk illustration of the price range implied by historical "
                           "volatility — not a forecast or price target.")
    _style(ws, f"A{mc + 8}", snap["A50"])
    # Blank cells in the lower blocks carry the column's styles, as in the
    # reference: values right-aligned in B, grey notes in C.
    table_rows = set(range(top, first_row + len(years)))
    for rr in range(35, mc + 8):
        if rr in table_rows or rr in (top, mc):
            continue
        if ws[f"B{rr}"].value is None:
            _style(ws, f"B{rr}", snap["B5"])
        if ws[f"C{rr}"].value is None:
            _style(ws, f"C{rr}", snap["C6"])


# ── Catalysts & news ─────────────────────────────────────────────────────────
def _catalysts(c):
    """(timing, event, why, status) rows built from the filing calendar and the
    themes of this company's recent headlines - no hand-written Apple text."""
    rows = []
    q = c.quarters
    if q:
        last_end = pd.Timestamp(q[-1]["end"])
        lag = 35
        f = c.filing_for(last_end)
        if f and f.get("filed"):
            lag = max(20, min(60, (pd.Timestamp(f["filed"]) - last_end).days))
        fy_m = c.fy_end.month if c.fy_end is not None else 12
        from data import fiscal_quarter_label
        for k in (1, 2):
            nxt_end = last_end + pd.DateOffset(months=3 * k)
            qq, fy = fiscal_quarter_label(nxt_end, fy_m)
            when = nxt_end + pd.Timedelta(days=lag)
            timing = f"{_part_of_month(when).capitalize()} {when:%b-%Y}"
            label = f"{qq} FY{str(fy)[2:]}"
            if qq == "Q4":
                rows.append((timing, f"{label} earnings",
                             f"Completes FY{fy}: full-year revenue, margins and cash flow; guidance "
                             "for the new fiscal year.", "Date not yet announced" if k == 1 else "Expected"))
                rows.append((timing, f"FY{fy} Form 10-K",
                             "Audited balance sheet, segment and geographic data, stock-based pay "
                             "and buybacks for the full year.", "Expected"))
            else:
                rows.append((timing, f"{label} earnings",
                             "Refreshes revenue, margins and the trailing-twelve-month figures in "
                             "this report; tests whether recent growth is holding.",
                             "Date not yet announced" if k == 1 else "Expected"))
    if (c.ttm.get("dps_quarter") or 0) > 0:
        rows.append(("With results", "Dividend declaration",
                     f"Last declared ${c.ttm['dps_quarter']:.2f} a share a quarter.", "Expected"))
    themes = {}
    for n in c.news:
        themes[str(n.get("Theme") or "")] = themes.get(str(n.get("Theme") or ""), 0) + 1
    for theme, event, why in (("Legal", "Legal and regulatory matters",
                               "Recent headlines flag legal or regulatory exposure."),
                              ("Regulat", "Regulatory rulings",
                               "Recent headlines flag regulatory exposure."),
                              ("Management", "Management changes",
                               "Recent headlines concern leadership."),
                              ("M&A", "Deal activity", "Recent headlines concern acquisitions or disposals."),
                              ("Product", "Product launches and demand",
                               "Recent headlines concern new products.")):
        if any(theme.lower() in k.lower() for k in themes):
            rows.append(("Ongoing", event, why, "Monitor"))
    return rows[:8]


def _news_sheet(wb, c):
    ws = wb["Catalysts_News"]
    snap = _snapshot(ws, ["A4", "A5", "B5", "C5", "D5", "A6", "B6", "C6", "D6", "A13", "A14",
                          "B14", "C14", "D14", "E14", "F14", "A15", "B15", "C15", "D15", "E15",
                          "F15", "A29"])
    _reset(ws, 4, max(ws.max_row, 40), 1, 7)
    _clear_heights(ws, 4)
    ws["A1"] = "Catalysts & News"
    ws["A2"] = "Upcoming events and a filtered list of recent headlines (QuantWizard feed, cleaned)"
    _back_link(ws, "F2")
    for col in "ABCDEF":
        _style(ws, f"{col}4", snap["A4"])
    _put(ws, "A4", "Upcoming catalysts")
    for col, text in zip("ABCD", ("Timing", "Event", "Why it matters", "Status")):
        _put(ws, f"{col}5", text); _style(ws, f"{col}5", snap[f"{col}5"])
    cats = _catalysts(c)
    r = 6
    for rec in cats:
        for col, v in zip("ABCDEF", tuple(rec) + (None, None)):
            _put(ws, f"{col}{r}", v); _style(ws, f"{col}{r}", snap[f"{col if col in 'ABCD' else 'D'}6"])
        r += 1
    c.catalyst_rows = cats
    r += 1
    for col in "ABCDEF":
        _style(ws, f"{col}{r}", snap["A13"])
    _put(ws, f"A{r}", "Recent headlines (filtered)")
    r += 1
    for col, text in zip("ABCDEF", ("Date", "Headline", "Publisher", "Theme", "Tone (reviewed)", "Note")):
        _put(ws, f"{col}{r}", text); _style(ws, f"{col}{r}", snap[f"{col}14"])
    r += 1
    top = r
    for n in c.news[:13]:
        dt = n.get("Date") or ""
        try:
            dt = pd.Timestamp(str(dt)[:10]).to_pydatetime()
        except Exception:
            pass
        note = "SEC filing" if "SEC" in str(n.get("Publisher") or "") else None
        vals = (dt, n.get("Headline"), n.get("Publisher"), n.get("Theme") or "General",
                n.get("Sentiment") or "Neutral", note)
        for col, v in zip("ABCDEF", vals):
            _put(ws, f"{col}{r}", v); _style(ws, f"{col}{r}", snap[f"{col}15"])
        url = n.get("URL")
        if url and str(url).startswith("http"):
            ws[f"B{r}"].hyperlink = url
        r += 1
    if r > top:
        rng = f"E{top}:E{r - 1}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'NOT(ISERROR(SEARCH("Negative",E{top})))'],
                                                       font=Font(color="FFC00000")))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'NOT(ISERROR(SEARCH("Positive",E{top})))'],
                                                       font=Font(color="FF006100")))
    else:
        _put(ws, f"A{r}", "No recent headlines were available."); r += 1
    r += 1
    _put(ws, f"A{r}", "Filtered from the aggregated feed: quote and option-chain pages, headlines about "
                      "other companies and off-topic press releases are removed; tone is scored on the "
                      "headline. Full list on Raw_News.")
    _style(ws, f"A{r}", snap["A29"])


# ── Methodology & Sources ────────────────────────────────────────────────────
def _methodology(wb, c):
    ws = wb["Methodology"]
    snap = _snapshot(ws, ["A4", "B4", "C4", "A5", "B5", "C5", "A9", "B9", "C9", "A6", "B6", "C6"])
    _reset(ws, 4, max(ws.max_row, 45), 1, 4)
    _clear_heights(ws, 4)
    ws["A1"] = "Methodology"
    ws["A2"] = "How each number in this report is built — and its known limitations"
    _back_link(ws, "C2")
    d, wb_ = c.dcf, (c.dcf.get("wacc_basis") or {})
    rf = c.rf
    wacc_txt = (f"CAPM: {wb_.get('risk_free', 0):.2%} 10-year Treasury + adjusted beta "
                f"{wb_.get('beta_adjusted', 1):.2f} × {wb_.get('erp', 0):.1%} ERP, blended with "
                f"after-tax debt → {d.get('wacc', 0):.1%}."
                if wb_ and not wb_.get("fallback") and wb_.get("beta") is not None
                else f"{d.get('wacc', 0.09):.1%} default.")
    blocks = [
        ("Fundamentals", [
            ("Fiscal periods", f"FY = {c.short}'s fiscal year, ending {c.fy_desc()}. TTM = the four quarters "
                               f"to {_d(c.flows_end)}, rebuilt from the cumulative figures in each "
                               "10-Q / 10-K." if c.has_ttm else
                               f"FY = {c.short}'s fiscal year, ending {c.fy_desc()}.",
             "Quarterly figures are reviewed, not audited; the fourth quarter is the full year less "
             "nine months."),
            ("Free cash flow", "Cash from operations minus capital expenditures.",
             "After interest and tax; not a strict unlevered FCF. Optional SBC deduction on the DCF tab."),
            ("Units", "$ billions unless stated; per-share figures in $; shares in billions.", None)]),
        ("DCF", [
            ("Structure", "10-year two-stage FCF model. Growth fades in a straight line from the "
                          "year-1 rate to terminal growth; Gordon-growth terminal value.",
             "Horizon fixed at 10 years."),
            ("Discounting", "Mid-year convention (toggle on DCF tab); terminal value discounted a full "
                            "10 years.", "Valuation date ≈ fiscal year-end, so year 1 is the current "
                                         "fiscal year."),
            ("WACC", wacc_txt, "Beta is estimated from the report window's daily returns; a model "
                               "input, not a measured cost of capital."),
            ("Net cash", "Cash + current + non-current marketable securities − commercial paper − "
                         "term debt (non-current securities are a toggle).",
             f"Balance sheet at {_d(c.bal_end if not c.bal_same else c.fy_end)}."),
            ("Share count", (f"{_bn(c.f.get('shares_outstanding')):,.3f}B shares outstanding at "
                             f"{_d(c.f.get('shares_date'))} (10-Q cover page)."
                             if c.f.get("market_cap_basis") == "filing"
                             else "Market capitalisation ÷ report price."),
             "Basic shares; a diluted count would be slightly higher."),
            ("Reverse DCF", "Solves for the year-1 growth rate at which DCF value = price, via a "
                            "lookup over a grid of growth rates; also reports the 10-year CAGR and "
                            "implied year-10 FCF.", "Described as a fading path, not a flat rate."),
            ("Scenarios", "Bear / base / bull flex growth, WACC and terminal growth together; "
                          "weights 25/50/25 are inputs.", "Sensitivity grids are live formulas.")]),
        ("Market & risk statistics", [
            ("Price data", f"Daily closes {_d(c.first_date)} to {_d(c.last_date)} "
                           f"({c.years:.1f} years), {c.src or 'data feed'}.",
             ("Yahoo closes are dividend-adjusted (total return)." if c.src == "Yahoo Finance"
              else "Polygon closes are split-adjusted only (price return).")
             + (" The last bar has no volume — confirm the official close." if c.last_bar_no_volume else "")),
            ("Volatility", "Sample standard deviation of daily returns × √252.", None),
            ("Sharpe / Sortino", f"(Annualised arithmetic mean return − {rf:.2%} risk-free) ÷ "
                                 "annualised volatility (downside deviation for Sortino).",
             f"A {c.years:.1f}-year window; a single sample, highly period-dependent."),
            ("Maximum drawdown", "Largest peak-to-trough decline in daily closes over the full window.", None),
            ("Calendar-year rows", "The first and last years may be partial and are labelled so.",
             "Not comparable with full-year figures."),
            ("Beta / correlation", ("Regression of daily returns on SPY over the window."
                                    if c.has_bench else
                                    "Not available — the benchmark series returned no data."),
             ("The WACC uses a Blume-adjusted beta (⅔ raw + ⅓)." if c.has_bench
              else "The WACC falls back to a default rate; this report shows beta as N/A.")),
            ("Monte Carlo", f"{c.mc.get('Simulations', 1000):,} paths, "
                            f"{c.mc.get('Forecast Horizon (days)', 252)} trading days; drift = CAPM "
                            "expected return, volatility from the window.",
             "Illustrative only: a random walk, not a forecast.")]),
        ("General", [
            ("Colour code", "Blue text on yellow = input you can change. Black = formula. Green = link "
                            "from another tab.", None),
            ("Disclaimer", "For information only; not investment advice. Figures from company filings "
                           "and third-party feeds may contain errors — check the Sources tab.", None)]),
    ]
    for col, text in zip("ABC", ("Item", "Method", "Limitations")):
        _put(ws, f"{col}4", text); _style(ws, f"{col}4", snap[f"{col}4"])
    r = 5
    for bi, (title, items) in enumerate(blocks):
        for col in "ABC":
            _style(ws, f"{col}{r}", snap[f"{col}5" if bi == 0 else f"{col}9"])
        _put(ws, f"A{r}", title)
        r += 1
        for item, method, lim in items:
            for col, v in zip("ABC", (item, method, lim)):
                _put(ws, f"{col}{r}", v); _style(ws, f"{col}{r}", snap[f"{col}6"])
            ws.row_dimensions[r].height = 30
            r += 1


def _sources(wb, c):
    ws = wb["Sources"]
    snap = _snapshot(ws, ["A4", "B4", "C4", "D4", "E4", "F4", "A5", "B5", "C5", "D5", "E5", "F5", "A20"])
    _reset(ws, 4, max(ws.max_row, 40), 1, 7)
    _clear_heights(ws, 4)
    ws["A1"] = "Sources & Data Log"
    ws["A2"] = "Every externally sourced figure used in this report, with where it lives and how reliable the source is"
    _back_link(ws, "F2")
    for col, text in zip("ABCDEF", ("Figure", "Used in", "Value", "Source", "Type", "Link")):
        _put(ws, f"{col}4", text); _style(ws, f"{col}4", snap[f"{col}4"])
    rows = []
    for qd in c.quarters[-4:]:
        f = c.filing_for(qd["end"])
        rows.append((f"{qd['label']} revenue / diluted EPS", "Financials quarters",
                     f"${_bn(qd['revenue']):,.1f}B / ${qd['eps']:.2f}" if qd.get("eps") is not None else
                     f"${_bn(qd['revenue']):,.1f}B",
                     f"Form {f['form']} for the period ended {_d(f['period'])}" if f else "SEC XBRL",
                     "SEC filing", f["url"] if f else c.edgar_index))
    if c.tenk:
        rows.append((f"{c.fys} balance sheet, cash flow, SBC, D&A", "Financials column B",
                     f"FY ended {_d(c.fy_end)}", f"Form 10-K filed {_d(c.tenk.get('filed'))}",
                     "SEC filing", c.tenk.get("url")))
    if c.tenq and not c.bal_same:
        rows.append(("Latest balance sheet & TTM cash flow", "Financials column C; DCF net debt",
                     f"At {_d(c.bal_end)}", f"Form 10-Q filed {_d(c.tenq.get('filed'))}",
                     "SEC filing", c.tenq.get("url")))
    if c.f.get("market_cap_basis") == "filing":
        rows.append(("Shares outstanding", "DCF B6; Multiples B6",
                     f"{c.f.get('shares_outstanding'):,.0f} ({_d(c.f.get('shares_date'))})",
                     "10-Q / 10-K cover page", "SEC filing", (c.tenq or c.tenk or {}).get("url")))
    if c.ttm.get("dps_quarter"):
        f = c.filing_for(c.ttm.get("dps_quarter_end"))
        rows.append(("Quarterly dividend", "Multiples B10", f"${c.ttm['dps_quarter']:.2f}",
                     "Dividends declared per share", "SEC filing", (f or {}).get("url")))
    rows.append((f"FY{c.first_fy_short[2:]}–{c.fys} revenue, net income, FCF", "Financials B6:K12",
                 "Various", "SEC EDGAR XBRL company facts (10-K)", "SEC filing", c.edgar_index))
    if (c.segments or {}).get("axes"):
        rows.append(("Revenue by segment and geography", "Business_Mix",
                     f"FY ended {_d(c.segments.get('fy_end'))}", "10-K XBRL instance",
                     "SEC filing", c.segments.get("url")))
    rows.append(("Price history & current price", "Price_Data; Risk; DCF B5",
                 f"${c.price:,.2f} ({_d(c.last_date)})", c.src or "Market data feed", "Feed",
                 f"https://finance.yahoo.com/quote/{c.t}/history/" if c.src != "Polygon.io"
                 else f"https://polygon.io/quote/{c.t}"))
    rows.append(("Risk-free rate", "DCF B14; Risk", f"{c.rf:.2%}", "US Treasury 3-month (FRED DGS3MO)",
                 "Government data", "https://fred.stlouisfed.org/series/DGS3MO"))
    if (c.dcf.get("wacc_basis") or {}).get("risk_free"):
        rows.append(("10-year Treasury (WACC)", "DCF B7", f"{c.dcf['wacc_basis']['risk_free']:.2%}",
                     "US Treasury 10-year (FRED DGS10)", "Government data",
                     "https://fred.stlouisfed.org/series/DGS10"))
    if c.peers:
        rows.append(("Peer multiples", "Multiples peer table", ", ".join(p["ticker"] for p in c.peers),
                     "Each peer's own 10-Q / 10-K XBRL; market caps from the data feed", "SEC filing / feed",
                     None))
    rows.append(("Headlines", "Catalysts_News", "—", "QuantWizard news feed (aggregators + SEC 8-Ks)",
                 "Third-party", None))
    rows.append(("Monte Carlo outputs", "Risk", "—", "QuantWizard simulation", "Model output", None))
    r = 5
    for rec in rows:
        for col, v in zip("ABCDEF", rec):
            val = ("Open source" if col == "F" and v else v)
            _put(ws, f"{col}{r}", val)
            _style(ws, f"{col}{r}", snap["E5" if (col == "F" and not v) else f"{col}5"])
            if col == "F" and v:
                ws[f"F{r}"].hyperlink = v
        ws.row_dimensions[r].height = 30
        r += 1
    ws.conditional_formatting.add(f"E5:E{r - 1}", FormulaRule(
        formula=['NOT(ISERROR(SEARCH("verif",E5)))'], fill=PatternFill("solid", bgColor="FFFFF2CC")))
    r += 1
    _put(ws, f"A{r}", "Source priority: SEC filings > data feeds > third-party sites. Every filing figure "
                      "comes from the XBRL the company filed; the link opens that filing.")
    _style(ws, f"A{r}", snap["A20"])


# ── Summary ──────────────────────────────────────────────────────────────────
def _bull_bear(c):
    """Four bull and four bear points, each built from this company's numbers."""
    bull, bear = [], []
    g = c.f.get("growth") or {}
    q = c.quarters[-3:]
    ys = [r["rev_yoy"] for r in q if r.get("rev_yoy") is not None]
    fy_g = (c.f_fy.get("growth") or {}).get("revenue_yoy")
    if len(ys) == 3:
        lo, hi = min(ys) * 100, max(ys) * 100
        rng = f"{lo:+.0f}%" if round(lo) == round(hi) else f"{lo:+.0f}–{hi:.0f}%"
        if _num(fy_g) and min(ys) * 100 > fy_g + 2:
            bull.append(f"Revenue re-accelerated to {rng} y/y in each of the last three quarters "
                        f"({q[0]['label']}–{q[-1]['label']}), from {fy_g:+.0f}% in {c.fys}.")
        elif min(ys) > 0.08:
            bull.append(f"Revenue grew {rng} y/y in each of the last three quarters.")
        elif max(ys) < 0:
            bear.append(f"Revenue fell in each of the last three quarters ({rng} y/y).")
    axes = (c.segments or {}).get("axes") or {}
    prod = axes.get("Products & services") or axes.get("Reportable segments") or []
    geo = axes.get("Geography") or axes.get("Reportable segments") or []
    growers = [r for r in prod if r.get("growth") is not None and r["share"] >= 0.05]
    if growers:
        best = max(growers, key=lambda r: r["growth"])
        if best["growth"] > 0.05:
            bull.append(f"{best['name']} (~{best['share']:.0%} of revenue) grew {best['growth']:+.0%} "
                        f"in {c.fys} — the fastest-growing line.")
    fallers = [r for r in (geo + prod) if r.get("growth") is not None and r["growth"] < 0
               and r["share"] >= 0.05]
    if fallers:
        worst = min(fallers, key=lambda r: r["growth"])
        bear.append(f"{worst['name']} revenue fell {abs(worst['growth']):.0%} in {c.fys} "
                    f"({worst['share']:.0%} of the total).")
    cr = c.f.get("capital_return") or {}
    sh = [x for x in ((c.f_fy.get("trend") or {}).get("diluted_shares") or []) if x]
    shrink = ((sh[-1] / sh[-4]) ** (1 / 3) - 1) if len(sh) >= 4 and 0.67 < sh[-1] / sh[-4] < 1.5 else None
    tot_cash = (c.cash_total_latest or 0)
    if cr.get("buybacks") and cr["buybacks"] > 0:
        bull.append(f"${tot_cash:,.0f}B of cash and securities"
                    + (", net cash" if (c.dcf.get("net_debt") or 0) < 0 else "")
                    + f", and ${cr['buybacks'] / 1e9:,.0f}B of buybacks over the last twelve months"
                    + (f" shrink the share count ~{abs(shrink):.0%} a year." if shrink and shrink < 0 else "."))
    m_fy, m_now = (c.f_fy.get("margins") or {}).get("gross"), (c.f.get("margins") or {}).get("gross")
    if _num(m_fy) and _num(m_now) and m_now > m_fy + 1:
        bull.append(f"Gross margin widened to {m_now:.1f}% over the last twelve months from "
                    f"{m_fy:.1f}% in {c.fys}.")
    ic, hc = c.dcf.get("market_implied_cagr"), c.fcf_cagr_9
    if _num(ic) and _num(hc) and ic > hc + 0.02:
        bear.insert(0, f"The price already discounts ~{ic:.0%} FCF compounding for 10 years vs "
                       f"~{hc:.0%} historically.")
    elif _num(ic) and _num(hc) and ic < hc:
        bull.insert(0, f"The price discounts only ~{ic:.0%} FCF growth a year for a decade, below "
                       f"the ~{hc:.0%} delivered historically.")
    pe = (c.f.get("valuation") or {}).get("pe")
    if c.vhist and _num(pe) and c.vhist.get("median_pe"):
        med, pct = c.vhist["median_pe"], c.vhist.get("percentile")
        if pe > med * 1.15:
            bear.append(f"At {pe:.0f}x trailing earnings the shares trade "
                        + ("above every year's average P/E of the last decade"
                           if pct == 1 else f"above {pct:.0%} of the last decade's yearly averages")
                        + f" (median {med:.0f}x).")
        elif pe < med * 0.85:
            bull.append(f"At {pe:.0f}x trailing earnings the shares trade below their decade median "
                        f"of {med:.0f}x.")
    if c.peers and _num(pe):
        from analysis import peer_median
        pm = peer_median(c.peers, skip=c.t).get("pe")
        if _num(pm) and pe > pm * 1.25:
            bear.append(f"P/E of {pe:.0f}x against a peer median of {pm:.0f}x "
                        f"({c.peer_group or 'industry peers'}).")
        elif _num(pm) and pe < pm * 0.8:
            bull.append(f"P/E of {pe:.0f}x against a peer median of {pm:.0f}x.")
    offs = g.get("tax_one_offs") or []
    if offs and _num(g.get("eps_yoy")) and _num(g.get("eps_yoy_ex_tax_one_offs")):
        bear.append(f"Reported EPS growth of {g['eps_yoy']:+.0f}% leans on a one-off tax item in the "
                    f"base period; excluding it EPS grew {g['eps_yoy_ex_tax_one_offs']:+.0f}%.")
    fc = c.f.get("fcf") or {}
    if _num(fc.get("sbc")) and _num(fc.get("fcf")) and fc["fcf"] > 0 and fc["sbc"] / fc["fcf"] > 0.08:
        bear.append(f"Stock-based pay of ${fc['sbc'] / 1e9:,.1f}B ({fc['sbc'] / fc['fcf']:.0%} of FCF) "
                    "is added back in reported free cash flow.")
    lev = (c.f.get("leverage") or {}).get("debt_to_equity")
    if _num(lev) and lev > 2:
        bear.append(f"Debt of {lev:.1f}x equity leaves less room for error.")
    return bull[:4], bear[:4]


def _quality_flags(c):
    flags = []
    if not c.has_bench:
        flags.append("Beta and benchmark correlation unavailable: the data feed returned no "
                     "S&P 500 series.")
    if c.last_bar_no_volume:
        flags.append(f"The final price bar ({_d(c.last_date)}) has no volume — an intraday snapshot; "
                     "confirm the official close.")
    flags.append(f"Price history covers {c.years:.1f} years; returns are "
                 + ("total returns (dividend-adjusted closes)." if c.src == "Yahoo Finance"
                    else "price-only." if c.src == "Polygon.io" else "as supplied by the feed."))
    flags.append((f"Balance sheet at {_d(c.bal_end)} (10-Q); TTM sums the four quarters to "
                  f"{_d(c.flows_end)}; annual figures are {c.fys} (audited)."
                  if c.has_ttm else f"Financials are {c.fys} (audited 10-K)."))
    if c.has_ttm:
        flags.append("Quarterly figures are from 10-Q XBRL (reviewed, not audited); the fourth "
                     "quarter is the full year less nine months.")
    flags = flags[:3]
    flags.append("Every external figure and its link is on the Sources tab.")
    return flags


def _summary(wb, c):
    ws = wb["Summary"]
    ws["A1"] = f"{c.name} ({c.exch}: {c.t}) — Equity Research Summary"
    fy_m = c.fy_end.strftime("%B") if c.fy_end is not None else ""
    ws["A2"] = (f"Report date {_d(c.now)} · Price $ per share · Financials $ in billions · "
                f"Fiscal year ends {_part_of_month(c.fy_end) if c.fy_end is not None else ''} {fy_m}")
    ws["A8"], ws["B8"] = "52-week low", "=Risk!B30"
    ws["A9"], ws["B9"] = "52-week high", "=Risk!B29"
    ws["A10"], ws["B10"] = f"{c.last_date.year} YTD price return", f"={c.ytd_cell}"
    ws["E13"] = "Implied year-1 FCF growth"
    ws["E14"] = "Implied 10-year FCF CAGR"
    ws["E15"] = f"Actual {c.span}-year FCF CAGR"
    ws["E16"] = f"Implied FCF in FY{c.fy + 10} ($B)"
    ws["B15"] = f"FY{c.fy}"
    if c.has_ttm:
        ws["C18"] = "=Multiples!C18"
        ws["C19"] = "=Multiples!C20"
        ws["C20"] = "=Multiples!C21"
    ws["A23"] = (f"Revenue, TTM to {_d(c.flows_end, '%b-%Y')} ($B)" if c.has_ttm
                 else f"Revenue, {c.fys} ($B)")
    if not c.has_ttm:
        ws["B23"] = "=Financials!K6"
    ws["A24"] = "Revenue growth, last 3 qtrs avg y/y"
    lq = c.quarters[-1] if c.quarters else None
    ws["A26"] = f"Gross margin, {lq['label']}" if lq else "Gross margin, latest quarter"
    ws["C26"] = f"quarter to {_d(lq['end'], '%b-%y')}" if lq else None
    ws["A27"] = f"Free cash flow {c.fys} ($B)"
    ws["A28"] = f"{c.span}-year FCF CAGR ({c.first_fy_short}–{c.fys})"
    # the analyst view, live off the cells above
    ws["E19"] = (
        f'="At "&TEXT(B5,"$0")&", {c.short} trades at "&IFERROR(TEXT(C16,"0.0x"),"n/m")&'
        '" trailing earnings. "&IF(ISNUMBER(F14),"To justify the price, free cash flow must compound at '
        'about "&TEXT(F14,"0%")&" a year for a decade"&IFERROR(IF(F15>0," — roughly "&TEXT(F14/F15,"0.0")&'
        f'"x its {c.span}-year record of "&TEXT(F15,"0%"),""),"")&" — reaching about $"&TEXT(F16,"0")&'
        f'"B by FY{c.fy + 10}. ","The price sits outside the growth range the model can solve, so no '
        f'single implied rate is quoted. ")&IFERROR(IF(B24>=0,"Recent results are strong (revenue +",'
        '"Recent results are soft (revenue ")&TEXT(B24,"0%")&" y/y over the last three quarters), ","")&'
        'IF(H7<0,"even the bull case ("&TEXT(H6,"$0")&") sits "&TEXT(-H7,"0%")&" below the price.",'
        '"the bull case ("&TEXT(H6,"$0")&") sits "&TEXT(H7,"0%")&" above the price.")&" On cash-flow '
        'fundamentals the shares look "&IF(F10<-0.15,"fully valued",IF(F10>0.15,"undervalued",'
        '"roughly fairly valued"))&IFERROR(IF(F14>F15,"; the case for owning them rests on growth well '
        'above the company\'s record proving durable.","."),".")')
    bull, bear = _bull_bear(c)
    for i in range(4):
        ws[f"A{32 + i}"] = ("• " + bull[i]) if i < len(bull) else None
        ws[f"D{32 + i}"] = ("• " + bear[i]) if i < len(bear) else None
    cats = [f"{r[1]} ({r[0][0].lower() + r[0][1:]}): {r[2]}"
            for r in (c.catalyst_rows or [])][:4]
    flags = _quality_flags(c)
    for i in range(4):
        ws[f"A{38 + i}"] = ("• " + cats[i]) if i < len(cats) else None
        ws[f"D{38 + i}"] = ("• " + flags[i]) if i < len(flags) else None
    links = [("A44", "Financials", "B44", "10-year history + TTM"), ("E44", "Risk", None, None),
             ("A45", "DCF", None, None), ("E45", "Catalysts_News", None, None),
             ("A46", "Scenarios", None, None), ("E46", "Methodology", None, None),
             ("A47", "Multiples", None, None), ("E47", "Sources", None, None)]
    for addr, sheet, daddr, desc in links:
        _put(ws, addr, f"→ {sheet}", link=f"'{sheet}'!A1")
        if daddr:
            ws[daddr] = desc


def _charts(wb, c):
    ws = wb["Summary"]
    sc = wb["Scenarios"]
    bar = BarChart()
    bar.type = "bar"
    bar.add_data(Reference(sc, min_col=2, min_row=38, max_row=44))
    bar.set_categories(Reference(sc, min_col=1, min_row=38, max_row=44))
    bar.anchor = TwoCellAnchor(_from=AnchorMarker(col=9, row=3), to=AnchorMarker(col=16, row=20))
    ws.add_chart(bar)
    pdws = wb["Price_Data"]
    line = LineChart()
    line.add_data(Reference(pdws, min_col=5, min_row=2, max_row=c.last_row))
    line.add_data(Reference(pdws, min_col=10, min_row=2, max_row=c.last_row))
    line.set_categories(Reference(pdws, min_col=1, min_row=2, max_row=c.last_row))
    line.anchor = TwoCellAnchor(_from=AnchorMarker(col=9, row=21), to=AnchorMarker(col=16, row=41))
    ws.add_chart(line)


# ── extras, in the reference's style ─────────────────────────────────────────
def _style_kit(wb):
    s, f = wb["Summary"], wb["Financials"]
    return {"title": copy.copy(s["A1"]._style), "sub": copy.copy(s["A2"]._style),
            "section": copy.copy(s["A4"]._style), "hdr_l": copy.copy(f["A5"]._style),
            "hdr_r": copy.copy(f["B5"]._style), "label": copy.copy(f["A8"]._style),
            "label_i": copy.copy(f["A9"]._style), "num": copy.copy(wb["DCF"]["B17"]._style),
            "pct": copy.copy(wb["DCF"]["B20"]._style), "text": copy.copy(wb["DCF"]["A17"]._style),
            "note": copy.copy(wb["DCF"]["C17"]._style), "back": copy.copy(f["L2"]._style)}


def _extra_sheet(wb, kit, name, title, subtitle, blocks, widths):
    """blocks: [(section title, headers, rows, formats)]."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_TAB
    ncol = max(len(b[1]) for b in blocks)
    for j in range(1, max(ncol, 6) + 1):
        ws.cell(row=1, column=j)._style = copy.copy(kit["title"])
        ws.cell(row=2, column=j)._style = copy.copy(kit["sub"])
    ws["A1"], ws["A2"] = title, subtitle
    ws.row_dimensions[1].height = 30
    back = ws.cell(row=2, column=max(ncol, 6), value="← Back to Summary")
    back._style = copy.copy(kit["back"])
    back.hyperlink = "#'Summary'!A1"
    back.hyperlink.location = "'Summary'!A1"
    r = 4
    for sec, heads, rows, fmts in blocks:
        for j in range(1, max(ncol, 6) + 1):
            ws.cell(row=r, column=j)._style = copy.copy(kit["section"])
        ws.cell(row=r, column=1, value=sec)
        r += 1
        for j, h in enumerate(heads, 1):
            cell = ws.cell(row=r, column=j, value=h)
            cell._style = copy.copy(kit["hdr_l" if j == 1 else "hdr_r"])
        r += 1
        for rec in rows:
            for j, v in enumerate(rec, 1):
                cell = ws.cell(row=r, column=j, value=("N/A" if v is None else v))
                cell._style = copy.copy(kit["text" if j == 1 else "num"])
                if j > 1 and fmts[j - 2] and isinstance(v, (int, float)):
                    cell.number_format = fmts[j - 2]
            r += 1
        r += 1
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A3"
    return ws


def _extras(wb, c):
    kit = _style_kit(wb)
    axes = (c.segments or {}).get("axes") or {}
    if axes:
        blocks = []
        for title, rows in axes.items():
            data = [[r["name"], _bn(r["value"]), _bn(r["prior"]), r["growth"], r["share"]] for r in rows]
            tot = sum(r["value"] for r in rows)
            data.append(["Total", _bn(tot), None, None, 1.0])
            blocks.append((title, ["Line", "Revenue ($B)", "Prior year ($B)", "Growth", "Share"],
                           data, [F_BN, F_BN, F_PCT, F_PCT]))
        _extra_sheet(wb, kit, "Business_Mix", f"{c.name} ({c.t}) — Where the Revenue Comes From",
                     f"Fiscal year ended {_d(c.segments.get('fy_end'))}, from the 10-K filed "
                     f"{_d(c.segments.get('filed'))} · each breakdown adds back to reported revenue",
                     blocks, [40, 14, 16, 11, 10, 12])
    tr = c.f_fy.get("trend") or {}
    per = tr.get("periods") or []
    if per and any(tr.get("buybacks") or []):
        rows = []
        for i, p in enumerate(per):
            fcf = tr["fcf"][i]
            bb, dv = (tr.get("buybacks") or [None] * len(per))[i], (tr.get("dividends") or [None] * len(per))[i]
            sbc, sh = (tr.get("sbc") or [None] * len(per))[i], (tr.get("diluted_shares") or [None] * len(per))[i]
            ret = ((bb or 0) + (dv or 0)) if (bb is not None or dv is not None) else None
            rows.append([f"FY{str(p)[:4]}", _bn(fcf), _bn(bb), _bn(dv), _bn(ret),
                         (ret / fcf) if (ret is not None and _num(fcf) and fcf > 0) else None,
                         _bn(sbc), _bn(sh) if sh else None])
        _extra_sheet(wb, kit, "Capital_Returns", f"{c.name} ({c.t}) — Where the Cash Goes",
                     "Free cash flow against buybacks, dividends and stock-based pay, per fiscal year as filed",
                     [("Capital returned ($B)", ["Fiscal year", "Free cash flow", "Buybacks", "Dividends",
                                                 "Total returned", "Returned / FCF", "Stock-based pay",
                                                 "Diluted shares (B)"],
                       rows, [F_BN, F_BN, F_BN, F_BN, F_PCT, F_BN, "#,##0.000"])],
                     [16, 14, 12, 12, 14, 14, 15, 16])
    if c.peers and len(c.peers) > 1:
        from analysis import peer_median
        med = peer_median(c.peers, skip=c.t)

        def _p(x):
            return x / 100 if _num(x) else None
        rows = [[r.get("ticker"), _bn(r.get("market_cap")), r.get("pe"), r.get("ev_ebitda"), r.get("ps"),
                 _p(r.get("rev_growth")), _p(r.get("gross_margin")), _p(r.get("op_margin")),
                 _p(r.get("net_margin")), _p(r.get("fcf_yield")), _p(r.get("div_yield")), r.get("basis")]
                for r in list(c.peers) + [dict(med, ticker="Peer median", basis="median of peers")]]
        _extra_sheet(wb, kit, "Peers", f"{c.name} ({c.t}) — Valuation vs Peers",
                     f"Peer group: {c.peer_group or 'industry and size'} · each company on its own latest "
                     "filings (TTM where a 10-Q is newer than the 10-K) and its own market cap",
                     [("Peer comparison", ["Ticker", "Market cap ($B)", "P/E", "EV / EBITDA", "P/S",
                                           "Revenue growth", "Gross margin", "Op. margin", "Net margin",
                                           "FCF yield", "Div. yield", "Figures as of"],
                       rows, [F_BN, F_X, F_X, F_X, F_PCT, F_PCT, F_PCT, F_PCT, F_PCT, F_PCT, None])],
                     [14, 14, 9, 11, 9, 13, 12, 11, 11, 10, 10, 20])
    if c.vhist:
        rows = [[str(r["year"]), r["eps"], r["pe_low"], r["pe_avg"], r["pe_high"], r["fcf_yield"]]
                for r in c.vhist["rows"]]
        rows.append(["Today (TTM)", None, None, c.vhist.get("current_pe"), None, None])
        _extra_sheet(wb, kit, "Valuation_History", f"{c.name} ({c.t}) — What the Market Has Paid Before",
                     "Each fiscal year's EPS (on today's share basis) against the stock's average, high and "
                     "low price that calendar year",
                     [("P/E history", ["Fiscal year", "EPS ($)", "P/E low", "P/E average", "P/E high",
                                       "FCF yield"],
                       rows, ['"$"0.00', F_X, F_X, F_X, F_PCT])],
                     [16, 12, 11, 12, 11, 11])


# ── N/A where an input is missing, never an error ────────────────────────────
def _na_guard(wb, c):
    """Ratios whose inputs a filer does not report (a bank has no capex and no
    operating margin) read N/A. Formulas stay exactly as the reference writes
    them whenever their inputs exist."""
    fin, mul, raw = wb["Financials"], wb["Multiples"], wb["Raw_Fundamentals"]

    def num(ws, a):
        return _num(ws[a].value) or (isinstance(ws[a].value, str) and ws[a].value.startswith("="))
    fcf_fy = _num(raw["K36"].value)
    if not fcf_fy:
        for a in ("B36", "B38"):
            fin[a] = "N/A"
        for a in ("B19", "B20", "B21", "B30", "B31", "B34"):
            mul[a] = "N/A"
    if not _num(raw["B15"].value):
        fin["B40"] = "N/A"
        mul["B28"] = "N/A"
    if not (num(fin, "B40") and fin["B40"].value != "N/A" and _num(fin["B39"].value)):
        for a in ("B18", "B36"):
            mul[a] = "N/A"
    if c.has_ttm:
        ttm_fcf = _num(fin["C35"].value) and _num(fin["C36"].value)
        if not ttm_fcf:
            fin["C38"] = "N/A"
            for a in ("C19", "C20", "C21"):
                mul[a] = "N/A"
        if not (_num(fin["C40"].value) and _num(fin["C39"].value)):
            mul["C18"] = "N/A"
    if not _num(fin["B37"].value):
        for a in ("B33", "B34"):
            mul[a] = "N/A"
        fin["B38"] = "N/A"


# ── companies a free-cash-flow DCF does not fit ──────────────────────────────
_NO_DCF_REASONS = {
    "no positive free cash flow to project":
        "it has not reported positive free cash flow in the years available (banks and "
        "brokers, whose cash flow is dominated by lending and trading, rarely do), so there "
        "is nothing to project",
    "no market cap": "no market capitalisation was available to derive a share count",
    "no price": "no current price was available",
    "WACC must exceed terminal growth": "the discount rate is not above terminal growth",
}


def _no_dcf(wb, c):
    """Replace the model with a plain statement instead of a sea of #VALUE!s.
    Inputs, net debt (Multiples reads it) and the price rows stay."""
    reason = _NO_DCF_REASONS.get((c.dcf or {}).get("reason"),
                                 "the inputs a DCF needs were not all available")
    ws = wb["DCF"]
    for r in list(range(16, 18)) + list(range(19, 190)):
        for col in "ABCDEFGHIJK":
            cell = ws[f"{col}{r}"]
            if type(cell).__name__ != "MergedCell":
                cell.value = None
    ws["A16"] = "Derived inputs"
    # the Summary reads these two history rows even without a model
    ws["A52"] = f"Historical {c.span}-yr FCF CAGR ({c.first_fy_short}–{c.fys})"
    ws["B52"] = "=Financials!L10"
    ws["A53"] = "Latest 3 quarters avg. revenue growth y/y"
    ws["B53"] = '=IFERROR(AVERAGE(Financials!C18:E18),"n/a")'
    ws["A22"] = "No discounted-cash-flow model"
    ws["A23"] = (f"No DCF for {c.short}: {reason}. Judge it on the multiples, quality ratios and "
                 "peer comparison instead.")
    ws["A23"].font = Font(name="Calibri", size=10, italic=True, color="FF595959")
    sc = wb["Scenarios"]
    for r in range(5, 36):
        for col in "ABCDEFG":
            cell = sc[f"{col}{r}"]
            if type(cell).__name__ != "MergedCell":
                cell.value = None
    sc["A5"] = f"No scenarios: {reason}."
    for r in range(38, 42):
        sc[f"B{r}"] = "n/a"
    su = wb["Summary"]
    for a in ("F6", "G6", "H6", "F7", "G7", "H7", "F8", "G8", "H8", "F9", "F10", "F13", "F14", "F15", "F16"):
        su[a] = "n/a"
    su["E19"] = (f'="At "&TEXT(B5,"$0")&", {c.short} trades at "&IFERROR(TEXT(C16,"0.0x"),"n/m")&'
                 f'" trailing earnings. No discounted-cash-flow value is shown: {reason}."')


# ── post-save: reference chart XML and the dynamic-array flag ────────────────
_METADATA_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<metadata xmlns="http://schemas.'
    'openxmlformats.org/spreadsheetml/2006/main" xmlns:xda="http://schemas.microsoft.com/office/'
    'spreadsheetml/2017/dynamicarray"><metadataTypes count="1"><metadataType name="XLDAPR" '
    'minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" merge="1" splitFirst="1" '
    'rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1" cellMeta="1"/>'
    '</metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-'
    '496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext>'
    '</extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata>'
    '</metadata>')


def _postprocess(raw, c):
    zin = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    bar_xml = line_xml = None
    try:
        with open(os.path.join(CHART_DIR, "value_bar.xml"), encoding="utf-8") as fh:
            bar_xml = fh.read()
        with open(os.path.join(CHART_DIR, "price_line.xml"), encoding="utf-8") as fh:
            line_xml = fh.read()
    except OSError:
        pass
    L = c.last_row
    win = (f"{c.years:.0f} years" if c.years >= 1.5 else f"{c.years * 12:.0f} months")
    names = zin.namelist()
    for name in names:
        data = zin.read(name)
        if name == "[Content_Types].xml":
            s = data.decode("utf-8")
            if "/xl/metadata.xml" not in s:
                s = s.replace("</Types>", '<Override PartName="/xl/metadata.xml" ContentType="application/'
                                          'vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>')
            data = s.encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            s = data.decode("utf-8")
            if "metadata.xml" not in s:
                s = s.replace("</Relationships>", '<Relationship Id="rIdQWmeta" Type="http://schemas.openxmlformats'
                                                  '.org/officeDocument/2006/relationships/sheetMetadata" '
                                                  'Target="metadata.xml"/></Relationships>')
            data = s.encode("utf-8")
        elif name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            s = data.decode("utf-8")
            s = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f t="array" ref="\1"',
                       lambda m: (f'<c r="{m.group(1)}"{m.group(2)} cm="1"><f t="array" ref="{m.group(1)}"'
                                  if "cm=" not in m.group(2) else m.group(0)), s)
            data = s.encode("utf-8")
        elif name == "xl/charts/chart1.xml" and bar_xml:
            s = bar_xml
            if not (c.dcf or {}).get("ok"):
                # no model: only the price rows (52-week low/high, current) have values
                s = s.replace("$A$38:$A$44", "$A$42:$A$44").replace("$B$38:$B$44", "$B$42:$B$44")
            data = s.encode("utf-8")
        elif name == "xl/charts/chart2.xml" and line_xml:
            s = line_xml.replace("$503", f"${L}").replace("{TITLE}", f"{c.t} price, {win} "
                                 + ("(total return, $)" if c.src == "Yahoo Finance" else "(price-only, $)"))
            data = s.encode("utf-8")
        zout.writestr(name, data)
    if "xl/metadata.xml" not in names:
        zout.writestr("xl/metadata.xml", _METADATA_XML)
    zout.close()
    return out.getvalue()


# ── entry point ──────────────────────────────────────────────────────────────
def build_report(ticker, df, financials=None, fundamentals=None, dcf=None,
                 company_details=None, mc_summary=None, mc_sim_df=None, news_rows=None,
                 peer_fund=None, peer_group=None, peer_df=None, segments=None,
                 valuation_data=None, filings=None, period_label="", price_source=None,
                 generated_at=None):
    """The workbook as BytesIO."""
    from analysis import compute_fundamentals, valuation_history
    from constants import get_risk_free_rate

    c = _Ctx(ticker, df, financials, fundamentals, dcf, company_details, mc_summary,
             mc_sim_df, news_rows, peer_fund, peer_group, peer_df, segments,
             valuation_data, filings, period_label, price_source, generated_at)
    fin_fy = {k: v for k, v in (financials or {}).items() if k != "ttm"}
    c.f_fy = compute_fundamentals(fin_fy, market_cap=(fundamentals or {}).get("market_cap"),
                                  price=c.price) if fin_fy else {}
    c.vhist = valuation_history(valuation_data, fundamentals) if valuation_data else None
    try:
        c.rf = float(get_risk_free_rate())
    except Exception:
        c.rf = 0.04
    tr = (c.f_fy.get("trend") or {})
    fcfs = [x for x in (tr.get("fcf") or [])[-10:]]
    per = list(tr.get("periods") or [])[-10:]
    c.span = max(1, len(per) - 1)
    c.first_fy_short = f"FY{str(per[0])[2:4]}" if per else c.fys
    c.fcf_cagr_9 = ((fcfs[-1] / fcfs[0]) ** (1 / c.span) - 1
                    if len(fcfs) >= 2 and _num(fcfs[0]) and _num(fcfs[-1]) and fcfs[0] > 0
                    and fcfs[-1] > 0 else None)
    c.fy_fcf = _bn(fcfs[-1]) if fcfs and _num(fcfs[-1]) else None
    c.fy_sbc = _bn(c.cf.iloc[0].get("sbc")) if len(c.cf) and "sbc" in c.cf.columns else None
    tb = c.ttm.get("balance") or {}
    c.bal_same = not (c.bal_end is not None and c.fy_end is not None and c.bal_end > c.fy_end and tb)
    src_b = tb if not c.bal_same else (c.bal.iloc[0].to_dict() if len(c.bal) else {})
    c.lts_latest = _bn(src_b.get("lt_securities")) if _num(src_b.get("lt_securities")) else None
    c.cash_total_latest = sum(_bn(src_b.get(k)) or 0 for k in ("cash", "short_term_investments",
                                                                 "lt_securities") if _num(src_b.get(k)))
    c.typical_tax = next((o["typical_rate"] for o in
                          ((fundamentals or {}).get("growth") or {}).get("tax_one_offs") or []), None)
    if c.typical_tax is None and len(c.inc) and "pretax_income" in c.inc.columns:
        rates = []
        for _, rw in c.inc.head(6).iterrows():
            p, t = rw.get("pretax_income"), rw.get("income_tax")
            if _num(p) and p > 0 and _num(t):
                rates.append(t / p)
        rates.sort()
        c.typical_tax = rates[len(rates) // 2] if len(rates) >= 3 else None
    e0 = c.last_row
    c.wk52_rows = (max(2, e0 - 252), e0)
    c.ytd_cell = "Risk!C40"
    c.catalyst_rows = []

    wb = load_workbook(TEMPLATE)
    per10, pad = _raw_fundamentals(wb, c, c.f_fy)
    _price_data(wb, c)
    _monte_carlo(wb, c)
    if c.news:
        _raw_table(wb, "Raw_News", list(c.news[0].keys()), [list(n.values()) for n in c.news])
    if c.peer_df is not None and len(c.peer_df):
        _raw_table(wb, "Raw_Peers", list(c.peer_df.columns),
                   [list(r) for r in c.peer_df.itertuples(index=False)])
    _financials(wb, c, per10, pad)
    _dcf(wb, c)
    _scenarios(wb, c)
    _multiples(wb, c)
    _risk(wb, c)
    _news_sheet(wb, c)
    _methodology(wb, c)
    _sources(wb, c)
    _summary(wb, c)
    _na_guard(wb, c)
    if not (dcf or {}).get("ok"):
        _no_dcf(wb, c)
    _charts(wb, c)
    _extras(wb, c)
    order = FRONT + [n for n in ("Peers", "Business_Mix", "Capital_Returns", "Valuation_History")
                     if n in wb.sheetnames] + ["Raw_Fundamentals", "Price_Data", "Raw_News",
                                               "Raw_Peers", "Monte_Carlo"]
    pos = {n: i for i, n in enumerate(order)}
    wb._sheets.sort(key=lambda s: pos.get(s.title, 999))
    for n in ("Raw_Fundamentals", "Price_Data", "Raw_News", "Raw_Peers", "Monte_Carlo"):
        wb[n].sheet_properties.tabColor = GREY_TAB
    wb.active = 0
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    buf = io.BytesIO()
    wb.save(buf)
    return io.BytesIO(_postprocess(buf.getvalue(), c))
