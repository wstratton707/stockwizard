"""The Excel report reproduces the reference workbook (assets/AAPL_5Y_Analysis
(6).xlsx) for any ticker. Synthetic filings and prices, no network - these pin
structure, live formulas and the absence of anything Apple-specific; the
numbers themselves are checked against the Python model in real Excel by hand
(see the module docstring of excel_report)."""
import io
import re
import zipfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import analysis as A
import data as D
from excel_report import build_report, FRONT


def _fin(years=10):
    ends = [pd.Timestamp(f"{2025 - i}-12-31") for i in range(years)]
    inc = pd.DataFrame([{
        "Period": str(e.date()), "revenues": 100e9 * 1.06 ** (years - i),
        "gross_profit": 45e9 * 1.06 ** (years - i), "operating_income_loss": 25e9 * 1.06 ** (years - i),
        "net_income_loss": 20e9 * 1.06 ** (years - i), "diluted_earnings_per_share": 2.0 * 1.07 ** (years - i),
        "diluted_shares": 10e9 * 0.98 ** (years - i), "depreciation_amortization": 3e9,
        "pretax_income": 24e9 * 1.06 ** (years - i), "income_tax": 4e9 * 1.06 ** (years - i),
    } for i, e in enumerate(ends)])
    bal = pd.DataFrame([{
        "Period": str(e.date()), "assets": 300e9, "current_assets": 120e9, "current_liabilities": 100e9,
        "liabilities": 200e9, "equity": 100e9, "long_term_debt": 60e9, "debt_current": 10e9,
        "cash": 30e9, "short_term_investments": 10e9, "lt_securities": 20e9, "retained_earnings": 50e9,
    } for e in ends])
    cf = pd.DataFrame([{
        "Period": str(e.date()), "net_cash_flow_from_operating_activities": 28e9 * 1.06 ** (years - i),
        "capex": 5e9, "sbc": 2e9, "buybacks": 10e9, "dividends_paid": 4e9,
    } for i, e in enumerate(ends)])
    return {"income_statement": inc, "balance_sheet": bal, "cash_flow_statement": cf, "source": "SEC EDGAR"}


def _prices(n=504):
    days = pd.bdate_range("2024-09-23", periods=n)
    rng = np.random.default_rng(3)
    close = 100 * np.cumprod(1 + rng.normal(0.0005, 0.012, n))
    df = pd.DataFrame({"Date": days, "Open": close, "High": close * 1.01, "Low": close * 0.99,
                       "Close": close, "Volume": 1e6})
    df["Daily_Return"] = df["Close"].pct_change()
    df["Cumulative_Index"] = (1 + df["Daily_Return"].fillna(0)).cumprod() * 100
    df["MA50"] = df["Close"].rolling(50).mean()
    df["MA200"] = df["Close"].rolling(200).mean()
    df["Volatility_20d"] = df["Daily_Return"].rolling(20).std() * np.sqrt(252)
    df["Drawdown_60d"] = df["Close"] / df["Close"].rolling(60, min_periods=1).max() - 1
    df["Pct_From_52W_High"] = df["Close"] / df["High"].rolling(252, min_periods=1).max() - 1
    df["RSI14"], df["MACD_Hist"] = 55.0, 0.1
    df["SPY_Return"] = df["Daily_Return"] * 0.8
    return df


def _build(dcf_ok=True, name="Widget Corp.", ticker="WDGT"):
    fin, df = _fin(), _prices()
    px = float(df["Close"].iloc[-1])
    f = A.compute_fundamentals(fin, market_cap=px * 10e9, price=px)
    d = A.dcf_valuation(f, px, beta=1.0) if dcf_ok else {"ok": False,
                                                         "reason": "no positive free cash flow to project"}
    mc_df, mc = A.run_monte_carlo(df, 200, 60, log=lambda *a, **k: None)
    buf = build_report(ticker, df, financials=fin, fundamentals=f, dcf=d,
                       company_details={"Name": name, "Exchange": "XNYS"},
                       mc_summary=mc, mc_sim_df=mc_df, news_rows=[], filings=[],
                       period_label="2Y", price_source="Yahoo Finance")
    return buf.getvalue(), d


def test_sheets_follow_the_reference_order():
    raw, _ = _build()
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames[:len(FRONT)] == FRONT
    assert {"Raw_Fundamentals", "Price_Data", "Monte_Carlo"} <= set(wb.sheetnames)


def test_the_model_is_live_formulas_with_the_dynamic_array_flag():
    raw, d = _build()
    wb = load_workbook(io.BytesIO(raw))
    dcf, sc, su = wb["DCF"], wb["Scenarios"], wb["Summary"]
    assert dcf["B42"].value == "=B41/B6"
    assert "LET" in getattr(sc["C20"].value, "text", "") and "LET" in getattr(dcf["K24"].value, "text", "")
    assert su["G6"].value == "=Scenarios!C10" and su["B5"].value == "=DCF!B5"
    assert dcf["B7"].value == pytest.approx(d["wacc"])
    z = zipfile.ZipFile(io.BytesIO(raw))
    assert "xl/metadata.xml" in z.namelist()
    sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
    assert any('cm="1"' in z.read(n).decode() for n in sheets)
    line = z.read("xl/charts/chart2.xml").decode()
    assert "Price_Data!$E$2:$E$505" in line and "{TITLE}" not in line and "numCache" not in line


def test_nothing_from_the_reference_company_leaks():
    raw, _ = _build()
    wb = load_workbook(io.BytesIO(raw))
    leak = re.compile(r"Apple|AAPL|iPhone|iPad|Samsung|tariff|Tim Cook|foldable|App Store")
    hits = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows()
            for c in row if isinstance(getattr(c.value, "text", c.value), str)
            and leak.search(getattr(c.value, "text", c.value))]
    assert hits == []
    assert "Widget" in wb["Summary"]["A1"].value and "WDGT" in wb["Summary"]["A1"].value


def test_a_company_without_a_dcf_gets_a_statement_not_errors():
    raw, _ = _build(dcf_ok=False)
    wb = load_workbook(io.BytesIO(raw))
    assert "No DCF for Widget" in wb["DCF"]["A23"].value
    assert wb["Summary"]["F6"].value == "n/a"
    assert wb["DCF"]["B18"].value.startswith("=Financials!")      # net debt still feeds Multiples


def test_the_base_fcf_switch_matches_the_site_model():
    raw, d = _build()
    wb = load_workbook(io.BytesIO(raw))
    dv = [x for x in wb["DCF"].data_validations.dataValidation if "B10" in str(x.sqref)][0]
    assert "reported" in dv.formula1 and "3-yr average" in dv.formula1
    assert d["base_fcf_basis"] == "fy" and wb["DCF"]["B10"].value.endswith("reported")


def test_template_carries_no_reference_values():
    wb = load_workbook("assets/report_template.xlsx")
    leak = re.compile(r"Apple|AAPL|iPhone|Samsung|tariff")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = getattr(c.value, "text", c.value)
                assert not (isinstance(v, str) and leak.search(v)), (ws.title, c.coordinate, v)
                assert not isinstance(v, (int, float)) or ws.title in ("Price_Data",), (ws.title, c.coordinate)


@pytest.mark.parametrize("end,fy_month,label", [
    ("2025-09-27", 9, ("Q4", 2025)), ("2025-12-27", 9, ("Q1", 2026)), ("2025-10-02", 9, ("Q4", 2025)),
    ("2026-06-30", 6, ("Q4", 2026)), ("2025-10-26", 1, ("Q3", 2026)), ("2026-04-26", 1, ("Q1", 2027)),
])
def test_fiscal_quarter_labels(end, fy_month, label):
    assert D.fiscal_quarter_label(end, fy_month) == label
