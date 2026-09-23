"""Regression tests for the AAPL workbook audit (Sep 2026).

Each pins one defect a reviewer found in an exported report. Synthetic inputs,
no network, so they stay fast and do not go stale.
"""
import numpy as np
import pandas as pd
import pytest

import analysis as A
import data as D


# ── The reverse DCF is a fading path, not a flat rate ─────────────────────────

def _fund(fcf):
    return {"ok": True, "market_cap": 1e12, "trend": {"fcf": list(fcf)},
            "net_debt": 0.0, "total_debt": 0.0, "growth": {}}


def test_implied_average_matches_the_fade_it_describes():
    """The solved rate is year one; the headline must be the path's average.
    AAPL read "31.1% a year for 10 years" when the path averaged 16.5%."""
    d = A.dcf_valuation(_fund([80e9, 90e9, 100e9, 110e9]), price=300.0, wacc=0.09)
    g1, tg, n = d["market_implied_growth"], d["terminal_growth"], d["years"]
    assert g1 is not None
    prod = 1.0
    for t in range(1, n + 1):
        prod *= 1 + (g1 + (tg - g1) * (t - 1) / (n - 1))
    assert d["market_implied_cagr"] == pytest.approx(prod ** (1 / n) - 1, rel=1e-12)
    assert d["market_implied_fcf_final"] == pytest.approx(d["base_fcf"] * prod, rel=1e-12)
    # A fade from a high start averages well below the start.
    if g1 > tg:
        assert d["market_implied_cagr"] < g1


# ── Benchmarks join on the market date, whatever each source stamped ──────────

def _bars(hours):
    days = pd.bdate_range("2025-01-02", periods=120)
    c = np.linspace(100, 120, len(days))
    return pd.DataFrame({"Date": days + pd.Timedelta(hours=hours), "Open": c,
                         "High": c + 1, "Low": c - 1, "Close": c, "Volume": 1e6})


def test_benchmark_merge_survives_mixed_timestamps(monkeypatch):
    """Stock at midnight (yfinance), benchmark at 04:00 UTC (Polygon): an exact
    join matched nothing and every SPY cell shipped empty."""
    def _fake(ticker, *a, **k):
        return _bars(4 if ticker == "SPY" else 0)
    monkeypatch.setattr(D, "fetch_ohlcv", _fake)
    df = D.fetch_stock_data("AAPL", "1y", ("SPY",), "", log=lambda *a, **k: None)
    filled = df["SPY_Return"].notna().sum()
    assert filled >= len(df) - 1, f"only {filled} of {len(df)} benchmark rows joined"


# ── Partial years are labelled, and short Sharpes are not quoted ──────────────

def test_partial_first_year_is_labelled(tmp_path):
    from openpyxl import Workbook
    from excel_builder import _build_annual_summary
    days = pd.bdate_range("2024-09-23", "2026-09-22")
    rng = np.random.default_rng(1)
    df = pd.DataFrame({"Date": days,
                       "Close": 100 * np.cumprod(1 + rng.normal(0.0004, 0.01, len(days)))})
    df["Daily_Return"] = df["Close"].pct_change()
    df["Drawdown_60d"] = df["Close"] / df["Close"].rolling(60, min_periods=1).max() - 1
    wb = Workbook()
    _build_annual_summary(wb, df)
    ws = wb["Annual_Summary"]
    first = ws.cell(row=3, column=1).value
    assert "partial" in str(first) and "23 Sep" in str(first), first
    assert ws.cell(row=3, column=5).value == "n/m", "a 3-month Sharpe must not be quoted"
    assert "YTD" in str(ws.cell(row=ws.max_row, column=1).value)


# ── The Valuation sheet is live end to end ────────────────────────────────────

def _live_workbook():
    from openpyxl import load_workbook
    import io as _io
    from excel_builder import build_excel
    periods = ["2025-09-27", "2024-09-28", "2023-09-30", "2022-09-24"]
    fcf = [110e9, 100e9, 90e9, 80e9]
    fin = {
        "income_statement": pd.DataFrame(
            [{"Period": p, "revenues": 400e9, "net_income_loss": 100e9,
              "operating_income_loss": 120e9, "diluted_earnings_per_share": 6.0}
             for p in periods]),
        "balance_sheet": pd.DataFrame(
            [{"Period": p, "equity": 70e9, "long_term_debt": 90e9, "cash": 70e9}
             for p in periods]),
        "cash_flow_statement": pd.DataFrame(
            [{"Period": p, "net_cash_flow_from_operating_activities": x + 10e9,
              "capex": 10e9} for p, x in zip(periods, fcf)]),
        "ttm": {"flows_end": "2025-09-27", "fy_end": "2025-09-27",
                "balance_end": "2025-09-27", "shares_outstanding": 15e9,
                "shares_date": "2025-10-17"},
    }
    fund = A.compute_fundamentals(fin, market_cap=3000e9, price=200.0)
    d = A.dcf_valuation(fund, 200.0, beta=1.1)
    import pytest as _pt
    with _pt.MonkeyPatch.context() as mp:
        mp.setattr(D, "fetch_ohlcv", lambda *a, **k: _bars(0))
        df = D.fetch_stock_data("TEST", "1y", (), "", log=lambda *a, **k: None)
    buf = build_excel("TEST", df, "1Y", company_details={"Name": "Test"},
                      fundamentals=fund, dcf=d, summary_text="x")
    return load_workbook(_io.BytesIO(buf.getvalue())), d


def test_sensitivity_grids_are_formulas_not_snapshots():
    """Editing an input moved every figure on the sheet except the 25 grid
    cells, which were numbers written at generation."""
    wb, d = _live_workbook()
    v = wb["Valuation"]
    grids = [r for r in range(1, v.max_row + 1)
             if str(v.cell(row=r, column=1).value or "").startswith("Sensitivity")]
    assert len(grids) == 2, "WACC x terminal g and WACC x stage-1 g"
    for r in grids:
        for i in range(5):
            for col in range(2, 7):
                val = v.cell(row=r + 2 + i, column=col).value
                assert isinstance(val, str) and val.startswith("=") and "SUMPRODUCT" in val
    assert "sensitivity_growth" in d and len(d["sensitivity_growth"]["grid"]) == 5


def test_wacc_is_built_up_and_feeds_the_input():
    from excel_builder import _VAL_ROWS
    wb, d = _live_workbook()
    v = wb["Valuation"]
    labels = {str(v.cell(row=r, column=1).value): r for r in range(1, 80)}
    r_out = next(r for k, r in labels.items() if k.startswith("WACC used"))
    assert v.cell(row=r_out, column=2).value.startswith("=ROUND(MAX(0.05,MIN(0.2,")
    assert v.cell(row=labels["Discount rate (WACC)"], column=2).value == f"=B{r_out}"
    assert v.cell(row=_VAL_ROWS["gmkt"], column=1).value.endswith("market-implied")


def test_dashboard_links_to_the_live_valuation_cells():
    wb, _ = _live_workbook()
    dsh = wb["Dashboard"]
    vals = {dsh.cell(row=r, column=1).value: dsh.cell(row=r, column=2).value
            for r in range(1, dsh.max_row + 1)}
    assert vals["DCF Fair Value / Share"] == "=Valuation!$B$9"
    assert vals["Upside / Downside"] == "=Valuation!$B$11"
    assert vals["DCF Verdict"] == "=Valuation!$B$12"
