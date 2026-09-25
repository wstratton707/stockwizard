"""The v2 reports: report_inputs (one set of numbers), excel_valuation (the
live workbook in the reference layout) and pptx_report (the reference deck).

Synthetic filings and prices, no network. These pin the rules that decide the
inputs, the structure of the files and the absence of anything carried over
from the reference company; Excel-vs-Python parity is checked against the
reference's own results in test_valuation_model and, by hand, in real Excel
(see excel_valuation's docstring)."""
import io
import re
import zipfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import analysis as A
import report_inputs as RI
import valuation_model as V


def _prices(n=1300, seed=3):
    days = pd.bdate_range("2021-06-01", periods=n)
    rng = np.random.default_rng(seed)
    spy = rng.normal(0.0004, 0.01, n)
    ret = 0.0002 + 1.2 * spy + rng.normal(0, 0.008, n)
    close = 100 * np.cumprod(1 + ret)
    return pd.DataFrame({"Date": days, "Open": close, "High": close * 1.01, "Low": close * 0.99,
                         "Close": close, "Volume": 1_000_000.0, "SPY_Return": spy})


def _quarters(scale=1.0, gains=0.0, oi=True):
    ends = ["2025-09-30", "2025-12-31", "2026-03-31", "2026-06-30"]
    out = []
    for i, e in enumerate(ends):
        rev = 30e9 * scale * (1 + 0.02 * i)
        ni = 6e9 * scale
        out.append({"end": e, "label": ["Q3 FY25", "Q4 FY25", "Q1 FY26", "Q2 FY26"][i], "fiscal_year": 2026,
                    "revenue": rev, "eps": 1.5, "gross_profit": rev * 0.5, "rev_yoy": 0.10, "eps_yoy": 0.12,
                    "eps_exact": True, "net_income": ni, "net_income_prev": ni * 0.9,
                    "operating_income": (rev * 0.25) if oi else None, "operating_income_prev": None,
                    "other_income": gains, "equity_gains": gains, "rnd": rev * 0.1, "rnd_prev": rev * 0.09,
                    "rev_prev": rev / 1.1, "da": rev * 0.04})
    return out


def _fin(years=10, gains=0.0, oi=True, scale=1.0):
    ends = [pd.Timestamp(f"{2025 - i}-12-31") for i in range(years)]
    inc = pd.DataFrame([{
        "Period": str(e.date()), "revenues": 100e9 * scale * 1.06 ** (years - i),
        "gross_profit": 45e9 * scale * 1.06 ** (years - i),
        "operating_income_loss": (25e9 * scale * 1.06 ** (years - i)) if oi else None,
        "net_income_loss": 20e9 * scale * 1.06 ** (years - i), "diluted_earnings_per_share": 2.0 * 1.07 ** (years - i),
        "diluted_shares": 10e9 * 0.98 ** (years - i), "depreciation_amortization": 3e9,
        "pretax_income": 24e9 * 1.06 ** (years - i), "income_tax": 4e9 * 1.06 ** (years - i),
    } for i, e in enumerate(ends)])
    bal = pd.DataFrame([{
        "Period": str(e.date()), "assets": 300e9, "current_assets": 120e9, "current_liabilities": 100e9,
        "liabilities": 200e9, "equity": 100e9, "long_term_debt": 60e9, "debt_current": 10e9,
        "cash": 30e9, "short_term_investments": 10e9, "lt_securities": 20e9, "retained_earnings": 50e9,
        "preferred": 0.0, "nonmkt_securities": 5e9,
    } for e in ends])
    cf = pd.DataFrame([{
        "Period": str(e.date()), "net_cash_flow_from_operating_activities": 28e9 * scale * 1.06 ** (years - i),
        "capex": 5e9 * scale, "sbc": 2e9, "buybacks": 10e9, "dividends_paid": 4e9,
    } for i, e in enumerate(ends)])
    q = _quarters(scale, gains, oi)
    ttm = {"flows_end": "2026-06-30", "fy_end": "2025-12-31", "balance_end": "2026-06-30",
           "income": {"revenues": sum(x["revenue"] for x in q), "gross_profit": sum(x["gross_profit"] for x in q),
                      "operating_income_loss": sum(x["operating_income"] for x in q) if oi else None,
                      "net_income_loss": sum(x["net_income"] for x in q), "depreciation_amortization": 13e9 * scale,
                      "pretax_income": 30e9, "income_tax": 5e9, "eps_diluted": 6.0},
           "cash_flow": {"net_cash_flow_from_operating_activities": 35e9 * scale, "capex": 6e9 * scale,
                         "sbc": 2.5e9, "buybacks": 10e9, "dividends_paid": 4e9},
           "balance": {"assets": 320e9, "equity": 110e9, "long_term_debt": 55e9, "debt_current": 8e9,
                       "cash": 35e9, "short_term_investments": 12e9, "lt_securities": 18e9,
                       "preferred": 2e9, "nonmkt_securities": 6e9},
           "prior": None, "quarters": q, "shares_outstanding": 9.8e9, "shares_date": "2026-07-20",
           "shares_fy": 9.9e9, "dps_quarter": 0.25, "drivers_bal": {"rpo": {"now": 80e9, "ago": 60e9}}}
    return {"income_statement": inc, "balance_sheet": bal, "cash_flow_statement": cf,
            "source": "SEC EDGAR", "ttm": ttm}


def _R(fin=None, df=None, sector="Technology", peers=None, **kw):
    fin = fin or _fin()
    df = df if df is not None else _prices()
    f = A.compute_fundamentals(fin, price=float(df["Close"].iloc[-1]))
    return RI.build("WDGT", df, fin, f, company_details={"Name": "Widget Corp.", "Exchange": "XNAS",
                                                         "Sector": sector},
                    peer_rows=peers, rf=0.04, rf10=0.045, sector=sector, **kw), f


# ── report_inputs ─────────────────────────────────────────────────────────────
def test_model_inputs_come_from_the_filings_block():
    R, _ = _R()
    fb, i = R["fb"], R["inputs"]
    assert R["dcf_ok"]
    assert i.revenue == pytest.approx(sum(q["revenue"] for q in fb["quarters"]))
    assert i.margin0 == pytest.approx(0.25)
    assert i.shares == pytest.approx(9.8)
    assert i.preferred == pytest.approx(2.0)
    assert i.net_cash == pytest.approx((35 + 12 + 18) - (55 + 8))
    assert R["assumptions"]["anchor"] == "TTM"
    assert R["wacc"]["wacc"] == pytest.approx(i.wacc)


def test_beta_is_excel_slope_over_five_years():
    df = _prices()
    R, _ = _R(df=df)
    w = df.tail(RI.DATA_ROWS).reset_index(drop=True)
    y = w["Close"].pct_change().values[1:]
    x = w["SPY_Return"].values[1:]
    assert R["beta_raw"] == pytest.approx(np.polyfit(x, y, 1)[0], rel=1e-9)
    assert len(R["win"]) == RI.DATA_ROWS


def test_clean_eps_only_when_gains_are_material():
    R, _ = _R(_fin(gains=0.0))
    assert not R["fb"]["gains_material"] and R["fb"]["ttm"]["gain_eps"] == 0
    R2, _ = _R(_fin(gains=3e9))                    # 2.4B after tax a quarter vs 6B net income
    assert R2["fb"]["gains_material"]
    q = R2["fb"]["quarters"][0]
    assert q["gain_ni"] == pytest.approx(3 * (1 - RI.GAIN_TAX))
    assert R2["fb"]["ttm"]["eps_clean"] < R2["fb"]["ttm"]["eps"]


def test_scenario_spreads_scale_with_the_company():
    assert V.scaled_spreads(0.17, 0.33) == {"growth": 0.06, "margin": 0.05, "wacc": 0.01, "tg": 0.005}
    small = V.scaled_spreads(0.05, 0.04)
    assert small["margin"] == 0.01 and small["growth"] == 0.02


def test_financial_companies_get_no_dcf():
    R, _ = _R(sector="Banks - Diversified")
    assert not R["dcf_ok"] and "financial" in R["dcf_reason"]
    assert RI.site_dcf(R) == {"ok": False, "reason": R["dcf_reason"]}


def test_loss_makers_anchor_on_profitable_peers_or_get_no_dcf():
    fin = _fin()
    for q in fin["ttm"]["quarters"]:
        q["operating_income"] = -q["revenue"] * 0.1
    R, _ = _R(fin)
    assert not R["dcf_ok"] and "no profitable peers" in R["dcf_reason"]
    peers = [{"ticker": "PEER", "op_margin": 20.0, "pe": 20.0, "market_cap": 1e11}]
    R2, _ = _R(fin, peers=peers)
    assert R2["assumptions"]["anchor"] == "Peer median"
    assert R2["inputs"].margin == pytest.approx(0.20)


def test_site_dcf_speaks_the_old_keys_with_the_new_model():
    R, _ = _R()
    d = RI.site_dcf(R)
    assert d["ok"] and d["model"] == "revenue"
    assert d["fair_value"] == pytest.approx(R["model"]["bridge"]["fair_value"])
    assert set(d["scenarios"]) == {"bear", "base", "bull"}
    assert d["market_implied_growth"] == R["model"]["reverse"]["growth"]
    snap = RI.snapshot(R)
    assert snap["ticker"] == "WDGT" and snap["base"] == pytest.approx(d["fair_value"])


# ── excel_valuation ───────────────────────────────────────────────────────────
def _book(R=None, df=None, fin=None, **kw):
    import excel_valuation as XV
    fin = fin or _fin()
    df = df if df is not None else _prices()
    if R is None:
        R, f = _R(fin, df)
    else:
        f = A.compute_fundamentals(fin, price=float(df["Close"].iloc[-1]))
    buf = XV.build_report("WDGT", df, fin, f, company_details={"Name": "Widget Corp.", "Exchange": "XNAS"},
                          news_rows=[], peer_fund=None, filings=[], price_source="Yahoo Finance", R=R, **kw)
    return buf.getvalue(), R


def test_workbook_follows_the_reference_tabs_and_keeps_its_formulas():
    import excel_valuation as XV
    raw, R = _book()
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == XV.SHEETS
    assert wb["DCF"]["B72"].value == "=B71/B6"
    assert "LET" in getattr(wb["Scenarios"]["B11"].value, "text", "")
    assert wb["Data"]["Q3"].value == "WDGT" and wb["Data"]["Q2"].value == "Widget Corp."
    assert wb["Financials"]["C56"].value == pytest.approx(9.8)
    assert wb["DCF"]["B31"].value == pytest.approx(R["inputs"].tax)
    assert wb["Scenarios"]["F7"].value == pytest.approx(R["assumptions"]["spreads"]["margin"])
    rev = sum(wb["Financials"][f"{c}17"].value for c in "BCDE")
    assert rev == pytest.approx(R["inputs"].revenue)
    z = zipfile.ZipFile(io.BytesIO(raw))
    assert "xl/metadata.xml" in z.namelist()
    assert any('cm="1"' in z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
    line = z.read("xl/charts/chart1.xml").decode()
    assert "Data!$E$2:$E$1256" in line and "{TITLE}" not in line and "numCache" not in line


def test_nothing_from_the_reference_company_leaks():
    raw, _ = _book()
    wb = load_workbook(io.BytesIO(raw))
    leak = re.compile(r"Alphabet|Google|GOOGL|\bTAC\b|antitrust|Wyatt|Stratton|Class A\+B\+C|Gemini")
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = getattr(c.value, "text", c.value)
                if isinstance(v, str) and leak.search(v):
                    hits.append((ws.title, c.coordinate))
        for part in ("oddFooter", "evenFooter", "firstFooter"):
            txt = getattr(ws, part).center.text or ""
            if leak.search(txt) or "{COMPANY}" in txt:
                hits.append((ws.title, part))
    assert hits == []


def test_short_price_history_rewrites_every_data_reference():
    df = _prices(n=600)
    raw, _ = _book(df=df)
    wb = load_workbook(io.BytesIO(raw))
    assert wb["Risk"]["B6"].value == "=MAX(Data!A2:A601)"
    assert "Data!E601" in wb["DCF"]["B5"].value
    assert wb["Data"]["E602"].value is None and wb["Data"]["H700"].value is None
    assert wb["Risk"]["C7"].value == "Report period: 2Y"          # says what the data covers
    assert "$601" in zipfile.ZipFile(io.BytesIO(raw)).read("xl/charts/chart1.xml").decode()


def test_no_dcf_company_reads_n_a_with_the_reason():
    fin = _fin()
    R, _ = _R(fin, sector="Banks - Regional")
    raw, _ = _book(R=R, fin=fin)
    wb = load_workbook(io.BytesIO(raw))
    assert wb["Summary"]["A5"].value == "No DCF"
    assert "financial" in wb["DCF"]["A50"].value
    assert wb["Scenarios"]["B11"].value == "n/a"


def test_valuation_log_rows_are_written_newest_first():
    hist = [{"report_date": "2026-09-20", "price": 100.0, "bear": 80.0, "base": 110.0, "bull": 150.0,
             "pw": 112.0, "upside": 0.12, "verdict": "Near model value", "implied_cagr": 0.08,
             "integrity": "✓ all checks pass"}]
    raw, _ = _book(log_history=hist)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Valuation_Log"]
    assert ws["D6"].value == pytest.approx(110.0)
    assert str(ws["A8"].value).startswith("Row 5")


def test_template_carries_no_reference_values():
    wb = load_workbook("assets/valuation_template.xlsx")
    assert "Website_Spec" not in wb.sheetnames
    leak = re.compile(r"Alphabet|Google|GOOGL|Wyatt|Stratton|20[12]\d")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith("="):
                    assert not leak.search(v), (ws.title, c.coordinate, v)
                assert c.comment is None, (ws.title, c.coordinate)
    dv = [d for d in wb["DCF"].data_validations.dataValidation if d.type == "list" and "Yes" in str(d.formula1)]
    assert all("B18" not in str(d.sqref) and "B19" not in str(d.sqref) for d in dv)


def test_files_are_signed_quantwizard_not_a_person():
    raw, _ = _book()
    core = zipfile.ZipFile(io.BytesIO(raw)).read("docProps/core.xml").decode()
    assert "<cp:lastModifiedBy>QuantWizard</cp:lastModifiedBy>" in core
    for path in ("assets/valuation_template.xlsx", "assets/report_template.xlsx",
                 "static/QuantWizard_Sample_NKE.xlsx"):
        z = zipfile.ZipFile(path)
        for part in z.namelist():
            assert not re.search(rb"Stratton|Wyatt", z.read(part)), (path, part)


# ── pptx_report ───────────────────────────────────────────────────────────────
def test_deck_builds_in_the_reference_design():
    from pptx import Presentation
    import pptx_report as PR
    df = _prices()
    R, f = _R(df=df)
    win = df.tail(504).reset_index(drop=True)
    win["Daily_Return"] = win["Close"].pct_change()
    mc_df, mc = A.run_monte_carlo(win, 200, 60, log=lambda *a, **k: None)
    buf = PR.build_deck("WDGT", win, R, company_details={"Name": "Widget Corp.", "Exchange": "XNAS"},
                        mc_summary=mc, mc_sim_df=mc_df, news_rows=[], fundamentals=f, period_label="2Y",
                        price_source="Yahoo Finance")
    prs = Presentation(buf)
    assert prs.slide_width == PR.W and prs.slide_height == PR.H
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Widget Corp. (WDGT)" in texts[0]
    assert any("Executive summary" in t for t in texts)
    assert "quantwizard.co" in texts[-1] and "quantwizard.app" not in " ".join(texts)
    assert not any("Apple" in t or "AAPL" in t for t in texts)
    assert prs.core_properties.author == "QuantWizard"
    assert prs.core_properties.last_modified_by == "QuantWizard"
