"""QuantWizard's PowerPoint equity review.

Reproduces the reference deck ('AAPL_5Y_Analysis EDITED GOOD VERSION.pptx' - an
earlier export of ours rebuilt by Claude in PowerPoint) for any company: a
navy cover, then consulting-style slides that each carry a section tag, a
sentence title stating the finding, a chart on the left, a ruled panel of
numbers on the right, two bold-lead takeaways and a source line. Every size,
colour, rule weight and position below is read from that deck.

The numbers come from report_inputs.build() - the same context the Excel
workbook and the site use - so the deck, the workbook and the site cannot
disagree. The valuation slides use the revenue-driven model: what the price
implies is year-one revenue growth and a long-run operating margin.

Every sentence is written from the data by a rule with a fallback, and kept
short enough for its box. Four slides beyond the reference (revenue mix, where
the cash goes, valuation range, peers and history) answer the reference's own
"data gaps to close" panel; each appears only when its data exists.
"""
import io
import math
import os
from datetime import datetime

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import matplotlib.ticker as mticker

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

import report_text as RT

_HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(_HERE, "assets", "logo_full.png")

# ── the reference deck's design tokens ────────────────────────────────────────
W, H = 12188825, 6858000
NAVY, TAG, TEXT, GREY, SRC = "1F4E79", "0B6FA4", "1B2533", "5B6475", "6B7280"
RED, GREEN_T = "B42318", "1E7B46"
PANEL, SIDE, RULE = "E8F1F8", "F3F5F8", "D5DAE1"
COVER_BG, COVER_ACCENT, COVER_TEXT, COVER_SRC, COVER_RULE = "14304F", "00B0F0", "A9C4DD", "8FA9C4", "3C5A7A"
BAR, BAR_HI, AXIS = "B8C2CE", "1F4E79", "C6C9CC"
M = 0.56                      # left margin, inches
CW = 12.21                    # content width
FONT = "Calibri"

MPL = {"navy": "#1F4E79", "grey": "#8A94A6", "light": "#B8C2CE", "red": "#B42318",
       "green": "#2E7D32", "orange": "#E8A838", "purple": "#8E44AD", "bg": "#F8F9FB", "grid": "#E3E7EC"}


def _rgb(h):
    return RGBColor.from_string(h)


def _num(x):
    return isinstance(x, (int, float, np.integer, np.floating)) and not (
        isinstance(x, float) and not math.isfinite(x))


def _pct(x, d=1, sign=False):
    if not _num(x):
        return "n/a"
    s = f"{x * 100:+.{d}f}%" if sign else f"{x * 100:.{d}f}%"
    return s.replace("-", "−")


def _usd(x, d=0):
    if not _num(x):
        return "n/a"
    s = f"${abs(x):,.{d}f}"
    return ("−" + s) if x < 0 else s


def _mult(x):
    return f"{x:.1f}x" if _num(x) else "n/a"


# ── primitive shapes ──────────────────────────────────────────────────────────
def _box(slide, x, y, w, h, runs, size=14, color=TEXT, bold=False, align="l", anchor="t",
         spc=None, italic=False, lnspc=None, bullets=False, space_after=0):
    """A text box with no insets and no autofit, as every box in the reference.
    `runs` is a string, a list of (text, {overrides}) runs, or - with
    bullets=True - a list of paragraphs, each a string or a list of runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")
    bp.set("anchor", {"t": "t", "b": "b", "ctr": "ctr", "m": "ctr"}[anchor])
    for child in list(bp):
        bp.remove(child)
    bp.append(bp.makeelement(qn("a:noAutofit"), {}))
    paras = runs if bullets else [runs]
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[align]
        pPr = p._p.get_or_add_pPr()
        if lnspc:
            p.line_spacing = lnspc
        if space_after:
            p.space_after = Pt(space_after)
        if bullets:
            pPr.set("marL", "120650")
            pPr.set("indent", "-120650")
            bc = pPr.makeelement(qn("a:buClr"), {})
            sc = bc.makeelement(qn("a:srgbClr"), {"val": TAG})
            bc.append(sc)
            pPr.append(bc)
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        items = para if isinstance(para, list) else [(para, {})]
        for text, over in items:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.italic = over.get("italic", italic)
            f.color.rgb = _rgb(over.get("color", color))
            if spc or over.get("spc"):
                r._r.get_or_add_rPr().set("spc", str(over.get("spc", spc)))
    return tb


def _line(slide, x1, y1, x2, y2, color=RULE, w=6350):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = _rgb(color)
    ln.line.width = Emu(w)
    # python-pptx gives connectors the theme's line style, whose effect
    # reference draws a soft shadow - the reference's rules are plain hairlines.
    st = ln._element.find(qn("p:style"))
    if st is not None:
        ln._element.remove(st)
    return ln


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    st = s._element.find(qn("p:style"))
    if st is not None:
        s._element.remove(st)
    return s


def _image(slide, buf, x, y, w, h):
    buf.seek(0)
    return slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))


# ── slide furniture ───────────────────────────────────────────────────────────
class Deck:
    def __init__(self, ctx):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.layout = self.prs.slide_layouts[6]
        self.ctx = ctx
        self.page = 0

    def slide(self, tag, title, source):
        s = self.prs.slides.add_slide(self.layout)
        self.page += 1
        _box(s, M, 0.36, CW, 0.22, tag.upper(), size=10, bold=True, color=TAG, spc=120)
        _box(s, M, 0.61, CW, 0.86, title, size=24, bold=True, color=NAVY, anchor="b", lnspc=0.92)
        _line(s, M, 1.58, M + CW, 1.58, NAVY, 12700)
        if source:
            _box(s, M, 6.42, CW, 0.44, source, size=10, color=SRC, anchor="b")
        _line(s, M, 7.08, M + CW, 7.08, RULE, 6350)
        _box(s, M, 7.14, 8.33, 0.22, f"{self.ctx['name']} ({self.ctx['ticker']})  |  Equity review  |  "
                                     f"{self.ctx['month']}", size=9, color=SRC)
        _box(s, 11.10, 7.14, 1.67, 0.22, str(self.page), size=9, bold=True, color=NAVY, align="r")
        return s


def _panel_label(s, x, y, w, text):
    _box(s, x, y, w, 0.25, text, size=12, bold=True, color=TEXT)
    _line(s, x, y + 0.31, x + w, y + 0.31, NAVY, 12700)


def _rows(s, x, y, w, rows, label_w=None, row_h=0.47, label_size=14, value_size=16, extra_w=0):
    """Ruled label / value rows. rows: (label, value, color) or (label, value,
    color, extra) - extra is a third, right-aligned column."""
    label_w = label_w or w * 0.6
    for i, row in enumerate(rows):
        label, value, color = row[:3]
        extra = row[3] if len(row) > 3 else None
        yy = y + i * row_h
        _box(s, x, yy, label_w, row_h - 0.05, label, size=label_size, color=GREY, anchor="ctr")
        vw = w - label_w - extra_w
        _box(s, x + label_w, yy, vw, row_h - 0.05, value, size=value_size, bold=True,
             color=color or TEXT, align="r", anchor="ctr")
        if extra is not None:
            _box(s, x + w - extra_w, yy, extra_w, row_h - 0.05, extra[0], size=14, bold=True,
                 color=extra[1], align="r", anchor="ctr")
        _line(s, x, yy + row_h - 0.03, x + w, yy + row_h - 0.03)
    return y + len(rows) * row_h


def _takeaways(s, items, y=5.64, h=0.78):
    _line(s, M, y - 0.08, M + CW, y - 0.08)
    cols = [(M, 5.97), (6.81, 5.97)] if len(items) > 1 else [(M, CW)]
    for (x, w), (lead, text) in zip(cols, items):
        _box(s, x, y, w, h, [(lead + " ", {"bold": True, "color": NAVY}), (text, {"color": TEXT})],
             size=14, lnspc=0.95)


def _notes(s, x, y, w, h, items, size=13):
    paras = [[(lead + " ", {"bold": True, "color": NAVY}), (text, {"bold": False, "color": TEXT})]
             for lead, text in items]
    tb = _box(s, x, y, w, h, paras[0], size=size, lnspc=0.95)
    tf = tb.text_frame
    for para in paras[1:]:
        p = tf.add_paragraph()
        p.line_spacing = 0.95
        p.space_before = Pt(8)
        for text, over in para:
            r = p.add_run()
            r.text = text
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold, r.font.color.rgb = over["bold"], _rgb(over["color"])
    return tb


# ── charts (images), in the reference's style ────────────────────────────────
def _fig(w, h):
    fig, ax = plt.subplots(figsize=(w, h), dpi=150)
    fig.patch.set_facecolor("white")
    ax.set_facecolor(MPL["bg"])
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    for sp in ("left", "bottom"):
        ax.spines[sp].set_color("#C6C9CC")
    ax.grid(True, color=MPL["grid"], linewidth=0.6)
    ax.tick_params(colors="#5B6475", labelsize=8)
    return fig, ax


def _png(fig):
    buf = io.BytesIO()
    fig.tight_layout(pad=0.4)
    fig.savefig(buf, format="png", dpi=150, facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def _dates(ax):
    loc = mdates.AutoDateLocator(minticks=5, maxticks=9)
    ax.xaxis.set_major_locator(loc)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    for lab in ax.get_xticklabels():
        lab.set_rotation(0)


def chart_cumulative(d, w=8.61, h=3.34):
    fig, ax = _fig(w, h)
    idx = d["Close"] / d["Close"].iloc[0] * 100
    ax.plot(d["Date"], idx, color=MPL["navy"], lw=1.6, label=None)
    ax.fill_between(d["Date"], idx, 100, where=idx >= 100, color="#2E7D32", alpha=0.08, lw=0)
    ax.fill_between(d["Date"], idx, 100, where=idx < 100, color=MPL["red"], alpha=0.10, lw=0)
    ax.axhline(100, color="#9AA3AF", lw=0.8, ls=(0, (4, 3)))
    if "SPY_Return" in d.columns and d["SPY_Return"].notna().sum() > 20:
        spy = (1 + d["SPY_Return"].fillna(0)).cumprod() * 100
        ax.plot(d["Date"], spy, color=MPL["grey"], lw=1.0, alpha=0.9)
        ax.annotate("SPY", (d["Date"].iloc[-1], spy.iloc[-1]), xytext=(4, 0), textcoords="offset points",
                    color=MPL["grey"], fontsize=8, va="center")
    ax.set_ylabel("Index (100 = start)", color="#5B6475", fontsize=8)
    _dates(ax)
    return _png(fig)


def chart_price_ma(d, w=8.61, h=3.34):
    fig, ax = _fig(w, h)
    ax.plot(d["Date"], d["Close"], color=MPL["navy"], lw=1.5, label="Close")
    for n, col, ls in ((20, MPL["orange"], (0, (4, 2))), (50, "#6BA368", (0, (4, 2))),
                       (200, MPL["red"], (0, (4, 2)))):
        ma = d["Close"].rolling(n).mean()
        ax.plot(d["Date"], ma, color=col, lw=1.1, ls=ls, label=f"{n}-day")
    ax.legend(loc="upper left", fontsize=8, frameon=False)
    ax.yaxis.set_major_formatter(mticker.StrMethodFormatter("${x:,.0f}"))
    _dates(ax)
    return _png(fig)


def chart_drawdown(d, w=8.61, h=1.59):
    fig, ax = _fig(w, h)
    dd = d["Close"] / d["Close"].rolling(60, min_periods=60).max() - 1
    ax.fill_between(d["Date"], dd * 100, 0, color=MPL["red"], alpha=0.55, lw=0)
    ax.plot(d["Date"], dd * 100, color=MPL["red"], lw=0.8)
    ax.axhline(-20, color="#9AA3AF", lw=0.8, ls=(0, (4, 3)))
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:.0f}%"))
    _dates(ax)
    return _png(fig)


def chart_volume(d, w=8.61, h=2.11):
    fig, ax = _fig(w, h)
    up = d["Close"].diff().fillna(0) >= 0
    ax.bar(d["Date"], d["Volume"] / 1e6, width=1.0, color=np.where(up, "#8FBF8F", "#D98B8B"), lw=0)
    ax.plot(d["Date"], d["Volume"].rolling(20).mean() / 1e6, color=MPL["navy"], lw=1.1, ls=(0, (4, 2)))
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:,.0f}M"))
    _dates(ax)
    return _png(fig)


def chart_bollinger(d, w=8.89, h=2.36):
    fig, ax = _fig(w, h)
    ma = d["Close"].rolling(20).mean()
    sd = d["Close"].rolling(20).std()
    ax.fill_between(d["Date"], ma - 2 * sd, ma + 2 * sd, color=MPL["navy"], alpha=0.08, lw=0)
    ax.plot(d["Date"], ma + 2 * sd, color=MPL["red"], lw=0.8, ls=(0, (4, 2)))
    ax.plot(d["Date"], ma - 2 * sd, color="#6BA368", lw=0.8, ls=(0, (4, 2)))
    ax.plot(d["Date"], d["Close"], color=MPL["navy"], lw=1.4)
    ax.yaxis.set_major_formatter(mticker.StrMethodFormatter("${x:,.0f}"))
    _dates(ax)
    return _png(fig)


def chart_rsi(d, w=8.89, h=1.57):
    fig, ax = _fig(w, h)
    rsi = d["RSI14"] if "RSI14" in d.columns else _rsi(d["Close"])
    ax.plot(d["Date"], rsi, color=MPL["purple"], lw=1.0)
    ax.axhline(70, color=MPL["red"], lw=0.8, ls=(0, (4, 2)))
    ax.axhline(30, color="#6BA368", lw=0.8, ls=(0, (4, 2)))
    ax.set_ylim(0, 100)
    _dates(ax)
    return _png(fig)


def chart_monte_carlo(paths, w=7.50, h=4.11):
    fig, ax = _fig(w, h)
    arr = np.asarray(paths, dtype=float)
    t = np.arange(arr.shape[0])
    p = {q: np.percentile(arr, q, axis=1) for q in (5, 25, 50, 75, 95)}
    ax.fill_between(t, p[5], p[95], color=MPL["navy"], alpha=0.10, lw=0, label="5th–95th percentile")
    ax.fill_between(t, p[25], p[75], color=MPL["navy"], alpha=0.22, lw=0, label="25th–75th percentile")
    ax.plot(t, p[50], color=MPL["navy"], lw=2.0, label="Median")
    ax.plot(t, p[5], color=MPL["red"], lw=0.9, ls=(0, (4, 2)))
    ax.plot(t, p[95], color="#2E7D32", lw=0.9, ls=(0, (4, 2)))
    ax.set_xlabel("Trading days forward", color="#5B6475", fontsize=8)
    ax.yaxis.set_major_formatter(mticker.StrMethodFormatter("${x:,.0f}"))
    ax.legend(loc="upper left", fontsize=8, frameon=False)
    return _png(fig)


def chart_football(rows, price, w=6.9, h=3.9):
    """Horizontal floating bars: (label, low, high)."""
    fig, ax = _fig(w, h)
    ax.grid(True, axis="x", color=MPL["grid"], lw=0.6)
    ax.grid(False, axis="y")
    ys = np.arange(len(rows))[::-1]
    for y, (lab, lo, hi) in zip(ys, rows):
        dcf = lab.startswith("DCF")
        ax.barh(y, hi - lo, left=lo, height=0.55, color=MPL["navy"] if dcf else MPL["light"])
        ax.text(hi, y, f"  {_usd(lo)}–{_usd(hi)}", va="center", fontsize=8, color="#1B2533")
    ax.axvline(price, color=MPL["red"], lw=1.4)
    ax.text(price, len(rows) - 0.35, f" price {_usd(price)}", color=MPL["red"], fontsize=8, va="bottom")
    ax.set_yticks(ys)
    ax.set_yticklabels([r[0] for r in rows], fontsize=8, color="#1B2533")
    ax.xaxis.set_major_formatter(mticker.StrMethodFormatter("${x:,.0f}"))
    lo = min(r[1] for r in rows + [("", price, price)])
    hi = max(r[2] for r in rows + [("", price, price)])
    ax.set_xlim(max(0, lo - (hi - lo) * 0.05), hi + (hi - lo) * 0.28)
    return _png(fig)


def chart_bars(labels, values, fmt="{:.0f}", w=6.1, h=3.4, color=MPL["light"], highlight=None, horizontal=False):
    fig, ax = _fig(w, h)
    ax.grid(False)
    cols = [MPL["navy"] if (highlight is not None and i == highlight) else color for i in range(len(values))]
    if horizontal:
        ax.barh(range(len(values))[::-1], values, color=cols)
        ax.set_yticks(range(len(values))[::-1])
        ax.set_yticklabels(labels, fontsize=8)
    else:
        ax.bar(range(len(values)), values, color=cols, width=0.6)
        ax.set_xticks(range(len(values)))
        ax.set_xticklabels(labels, fontsize=8)
        for i, v in enumerate(values):
            ax.text(i, v, fmt.format(v), ha="center", va="bottom", fontsize=8, color="#1B2533")
    return _png(fig)


def _rsi(close, n=14):
    d = close.diff()
    up, dn = d.clip(lower=0).rolling(n).mean(), (-d.clip(upper=0)).rolling(n).mean()
    return 100 - 100 / (1 + up / dn.replace(0, np.nan))


# ── native (editable) charts, styled as the reference's ───────────────────────
def _native_bar(slide, x, y, w, h, cats, vals, fmt, horizontal, highlight):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("", [float(v) for v in vals])
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 50
    plot.vary_by_categories = False
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = _rgb(BAR)
    ser.invert_if_negative = False
    if highlight is not None:
        pt = ser.points[highlight]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(BAR_HI)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size, dl.font.bold, dl.font.name = Pt(14), True, FONT
    dl.font.color.rgb = _rgb(TEXT)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(14), FONT
    ca.tick_labels.font.color.rgb = _rgb(TEXT)
    ca.format.line.color.rgb = _rgb(AXIS)
    ca.has_major_gridlines = False
    if horizontal:
        ca.reverse_order = True
    return gf


# ── the context every slide reads ─────────────────────────────────────────────
def _context(ticker, df, R, company_details, mc_summary, mc_sim_df, news, fundamentals,
             period_label, price_source, vhist, filings):
    cd = company_details or {}
    P = RT.profile(R)
    d = df.copy()
    d["Date"] = pd.to_datetime(d["Date"])
    d = d.sort_values("Date").reset_index(drop=True)
    close = d["Close"].astype(float)
    ret = close.pct_change()
    from analysis import window_stats
    ws = window_stats(d.assign(Daily_Return=ret), rf=R.get("rf"))
    last = float(close.iloc[-1])
    tail = d.tail(252)
    hi52 = float((tail["High"] if "High" in tail else tail["Close"]).max())
    lo52 = float((tail["Low"] if "Low" in tail else tail["Close"]).min())
    ma = {n: float(close.rolling(n).mean().iloc[-1]) if len(close) >= n else None for n in (20, 50, 200)}
    ma200_prev = (float(close.rolling(200).mean().iloc[-64]) if len(close) >= 264 else None)
    bb_ma, bb_sd = close.rolling(20).mean().iloc[-1], close.rolling(20).std().iloc[-1]
    pctb = float((last - (bb_ma - 2 * bb_sd)) / (4 * bb_sd)) if bb_sd else None
    rsi = (float(d["RSI14"].iloc[-1]) if "RSI14" in d.columns and _num(d["RSI14"].iloc[-1])
           else float(_rsi(close).iloc[-1]))
    dd60 = close / close.rolling(60, min_periods=60).max() - 1
    cum = close / close.cummax() - 1
    vol20 = float(ret.tail(20).std() * np.sqrt(252))
    idx = close / close.iloc[0] * 100
    spy_ret = (float((1 + d["SPY_Return"].fillna(0)).prod() - 1)
               if "SPY_Return" in d.columns and d["SPY_Return"].notna().sum() > 20 else None)
    mc = mc_summary or {}

    def mcv(k):
        v = mc.get(k)
        try:
            return float(str(v).rstrip("%")) / (100 if str(v).endswith("%") else 1)
        except Exception:
            return None
    now = pd.Timestamp(R["generated"])
    ctx = {"ticker": R["ticker"], "name": P["name"], "short": P["short"], "exchange": P["exchange"],
           "month": now.strftime("%B %Y"), "date": now, "d": d, "price": last, "R": R,
           "fb": R["fb"], "dcf_ok": R.get("dcf_ok"), "model": R.get("model"),
           "A": R.get("assumptions") or {}, "ws": ws, "hi52": hi52, "lo52": lo52, "ma": ma,
           "ma200_prev": ma200_prev, "pctb": pctb, "rsi": rsi, "dd60": dd60, "cum": cum, "vol20": vol20,
           "idx": idx, "spy_ret": spy_ret, "start": d["Date"].iloc[0], "end": d["Date"].iloc[-1],
           "period_label": period_label, "price_source": price_source or "market data feed",
           "mc": {"p5": mcv("Bear Case (P5)"), "p25": mcv("Low Case (P25)"), "p50": mcv("Median (P50)"),
                  "p75": mcv("Bull Case (P75)"), "p95": mcv("Best Case (P95)"),
                  "prob": mcv("Prob. of Gain"), "n": int(mc.get("Simulations") or 1000),
                  "days": int(mc.get("Forecast Horizon (days)") or 252), "start": mcv("Last Price"),
                  "vol": mcv("Ann. Volatility")},
           "mc_df": mc_sim_df, "news": news or [], "f": fundamentals or {}, "cd": cd, "vhist": vhist,
           "filings": filings or []}
    yrs = (ctx["end"] - ctx["start"]).days / 365.25
    ctx["window"] = (f"{yrs:.0f} years" if yrs >= 1.5 else "the past year" if yrs >= 0.9
                     else f"{yrs * 12:.0f} months")
    ctx["window_word"] = {1: "one year", 2: "two years", 3: "three years", 5: "five years",
                          10: "ten years"}.get(round(yrs), ctx["window"])
    ctx["window_adj"] = {1: "One-year", 2: "Two-year", 3: "Three-year", 5: "Five-year",
                         10: "Ten-year"}.get(round(yrs), f"{max(1, round(yrs))}-year")
    return ctx


# ── slides ────────────────────────────────────────────────────────────────────
def _verdict_words(c):
    R = c["R"]
    if not c["dcf_ok"]:
        return "no DCF anchor", None
    v = R["model"]["verdict"]
    return {"Priced above model value": "stretched valuation",
            "Priced below model value": "undemanding valuation",
            "Near model value": "valuation near model value"}.get(v, "valuation near model value"), v


def _momentum(c):
    ws, ma = c["ws"], c["ma"]
    r = (ws.get("period_ret") or 0) / 100
    above200 = ma.get(200) and c["price"] > ma[200]
    if r > 0.20 and above200:
        return "strong momentum"
    if r < -0.10 or (ma.get(200) and c["price"] < ma[200] * 0.95):
        return "weak momentum"
    return "steady trading"


def slide_cover(dk, c):
    s = dk.prs.slides.add_slide(dk.layout)
    dk.page += 1
    _rect(s, 0, 0, 13.33, 7.5, COVER_BG)
    _rect(s, 0, 0, 0.11, 7.5, COVER_ACCENT)
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, Inches(10.82), Inches(0.47), Inches(1.94), Inches(1.33))
    _box(s, 0.78, 1.67, 9.72, 0.25, f"EQUITY REVIEW  ·  {c['month'].upper()}", size=11, bold=True,
         color=COVER_ACCENT, spc=150)
    _box(s, 0.78, 2.03, 10.56, 0.78, f"{c['name']} ({c['ticker']})", size=40, bold=True, color="FFFFFF")
    val, _ = _verdict_words(c)
    _box(s, 0.78, 2.83, 11.11, 0.56, f"{_momentum(c).capitalize()}, {val}", size=24, color=COVER_TEXT)
    fy = c["fb"].get("fy") or "latest"
    _box(s, 0.78, 3.56, 11.11, 0.56,
         f"{c['window_adj']} price performance ({c['start']:%b %Y} – {c['end']:%b %Y}), "
         f"{fy} fundamentals, reverse DCF and 12-month Monte Carlo range", size=13, color=COVER_TEXT)
    _line(s, 0.78, 5.31, 12.77, 5.31, COVER_RULE, 9525)
    base = c["model"]["bridge"]["fair_value"] if c["dcf_ok"] else None
    kpis = [(_usd(c["price"], 2), "Last close"),
            (_pct((c["ws"].get("period_ret") or 0) / 100, 1, True), f"Total return since {c['start']:%b %Y}"),
            (_usd(base) if base else "n/a", "DCF base-case value per share"),
            (_pct(c["mc"]["prob"], 1) if _num(c["mc"]["prob"]) else "n/a",
             "Simulated probability of a 12-month gain")]
    for i, (v, lab) in enumerate(kpis):
        x = 0.78 + i * 3.0
        _box(s, x, 5.53, 2.77, 0.61, v, size=30, bold=True, color="FFFFFF")
        _box(s, x, 6.17, 2.77, 0.47, lab, size=12, color=COVER_TEXT)
        if i:
            _line(s, x - 0.14, 5.56, x - 0.14, 6.56, COVER_RULE, 9525)
    _box(s, 0.78, 6.94, 11.99, 0.25,
         f"Data: {c['price_source']}, SEC EDGAR. Prepared by QuantWizard, {c['date']:%d %b %Y}. "
         "For information only; not investment advice.", size=10, color=COVER_SRC)


def slide_summary(dk, c):
    fb, ws, A = c["fb"], c["ws"], c["A"]
    r = (ws.get("period_ret") or 0) / 100
    val, verdict = _verdict_words(c)
    short, T = c["short"], c["ticker"]
    fy = fb.get("fy_end")
    s = dk.slide("Summary", "Executive summary",
                 f"Source: QuantWizard analysis of {c['price_source']} prices and SEC EDGAR filings. "
                 f"Fundamentals: FY ended {fy:%d %b %Y}; TTM to {fb.get('flows_end'):%d %b %Y}. "
                 f"Returns: {c['start']:%b %Y} – {c['end']:%d %b %Y}." if fy is not None and fb.get('flows_end') is not None
                 else f"Source: QuantWizard analysis of {c['price_source']} prices. "
                      f"Returns: {c['start']:%b %Y} – {c['end']:%d %b %Y}.")
    _rect(s, M, 1.81, CW, 0.97, PANEL)
    _rect(s, M, 1.81, 0.07, 0.97, NAVY)
    if c["dcf_ok"]:
        rv = c["model"]["reverse"]
        lead = {"Priced above model value": f"{short} is priced above what its cash flows support on the model's assumptions.",
                "Priced below model value": f"{short} trades below the model's value of its cash flows.",
                "Near model value": f"{short} trades close to the model's value of its cash flows."}.get(verdict)
        detail = (f" The stock has returned {_pct(r, 1, True)} over {c['window_word']}"
                  + (f"; today's price needs {_pct(rv['growth'], 0)} year-one revenue growth against the "
                     f"model's {_pct(A['g1'], 0)} base case." if _num(rv.get("growth")) else "."))
    else:
        lead = f"{short} is valued here on multiples; the DCF does not apply."
        detail = f" The stock has returned {_pct(r, 1, True)} over {c['window_word']} ({c['R'].get('dcf_reason')})."
    _box(s, 0.83, 1.81, 11.71, 0.97, [(lead, {"bold": True, "color": NAVY}), (detail, {"color": TEXT})],
         size=16, anchor="ctr", lnspc=0.95)
    cols = _summary_columns(c)
    for i, (head, value, vcolor, caption, bullets) in enumerate(cols):
        x = M + i * 3.12
        _box(s, x, 3.06, 2.84, 0.22, head.upper(), size=10, bold=True, color=TAG, spc=100)
        _box(s, x, 3.31, 2.84, 0.56, value, size=30, bold=True, color=vcolor)
        _box(s, x, 3.89, 2.84, 0.47, caption, size=12, color=GREY, lnspc=0.9)
        _line(s, x, 4.42, x + 2.84, 4.42, RULE, 9525)
        _box(s, x, 4.50, 2.84, 1.39, bullets, size=14, bullets=True, lnspc=0.92, space_after=5)
    _line(s, M, 6.00, M + CW, 6.00, NAVY, 9525)
    _box(s, M, 6.06, CW, 0.56, [("Implication: ", {"bold": True, "color": NAVY}),
                               (_implication(c), {"color": TEXT})], size=14)


def _summary_columns(c):
    fb, ws, A, f = c["fb"], c["ws"], c["A"], c["f"]
    ttm = fb.get("ttm") or {}
    r = (ws.get("period_ret") or 0) / 100
    dd = (ws.get("dd_peak_trough") or 0) / 100
    below_hi = c["price"] / c["hi52"] - 1
    ma_state = [n for n in (50, 200) if c["ma"].get(n)]
    above = [n for n in ma_state if c["price"] > c["ma"][n]]
    perf_b = [f"Max peak-to-trough fall of {_pct(abs(dd), 0)} over the window"
              + (", since recovered" if c["cum"].iloc[-1] > -0.02 and dd < -0.1 else ""),
              f"{_pct(abs(below_hi), 1)} below 52-week high; "
              + ("above " if len(above) == len(ma_state) else "below " if not above else "mixed vs ")
              + "50- and 200-day averages"]
    q = f.get("quality") or {}
    rg, eg = ttm.get("rev_growth"), (f.get("growth") or {}).get("eps_yoy")
    fcf = (fb.get("C") or {}).get("cfo"), (fb.get("C") or {}).get("capex")
    fcf_v = (fcf[0] - fcf[1]) if (_num(fcf[0]) and _num(fcf[1])) else None
    fund_b = [(f"${fcf_v:,.1f}B free cash flow" if _num(fcf_v) else "Free cash flow n/a")
              + (f"; Piotroski score {q['f_score']} of 9" if q.get("f_score") is not None else ""),
              f"Revenue grew {_pct(rg, 1)}" + (f"; EPS grew {_pct(eg / 100, 1)}" if _num(eg) else "")]
    om = ttm.get("op_margin")
    cols = [("Performance", _pct(r, 1, True), NAVY if r >= 0 else RED, f"total return since {c['start']:%b %Y}", perf_b),
            ("Fundamentals", _pct(om, 1), NAVY, "operating margin, TTM", fund_b)]
    pe = (c["price"] / ttm["eps_clean"]) if (_num(ttm.get("eps_clean")) and ttm["eps_clean"] > 0) else None
    fcfy = (fcf_v / (c["price"] * (fb.get("shares_now") or 1))) if _num(fcf_v) and fb.get("shares_now") else None
    if c["dcf_ok"]:
        rv, br = c["model"]["reverse"], c["model"]["bridge"]
        ig = rv.get("growth")
        hot = _num(ig) and ig > 2 * max(A["g1"], 0.01)
        cols.append(("Valuation", _pct(ig, 1) if _num(ig) else "n/a", RED if hot else NAVY,
                     f"year-one revenue growth implied by the {_usd(c['price'])} price",
                     [f"Base-case DCF value {_usd(br['fair_value'])}, {_pct(abs(br['upside']), 0)} "
                      f"{'below' if br['upside'] < 0 else 'above'} the price",
                      f"{_mult(pe)} P/E, {_pct(fcfy, 1)} FCF yield"]))
    else:
        pb = (f.get("valuation") or {}).get("pb")
        cols.append(("Valuation", _mult(pe), NAVY, "P/E, trailing twelve months",
                     [(f"{_mult(pb)} price / book" if _num(pb) else f"{_pct(fcfy, 1)} FCF yield"),
                      "No DCF: judged on multiples (see the valuation slide)"]))
    if not _num(om):
        nm = (f.get("margins") or {}).get("net")
        cols[1] = ("Fundamentals", _pct(nm / 100, 1) if _num(nm) else "n/a", NAVY, "net margin, TTM", fund_b)
    mc = c["mc"]
    chg = (mc["p50"] / c["price"] - 1) if _num(mc["p50"]) else None
    cols.append(("12-month outlook", _usd(mc["p50"]), NAVY,
                 f"median simulated price, {_pct(chg, 0, True)} vs today",
                 [f"5th–95th percentile range {_usd(mc['p5'])}–{_usd(mc['p95'])}",
                  f"{_pct(mc['prob'], 1)} of {mc['n']:,} paths end higher"]))
    return cols


def _implication(c):
    A, T = c["A"], c["ticker"]
    if not c["dcf_ok"]:
        return f"without a DCF anchor, {T}'s price is best judged against peers and its own multiple history."
    v = c["model"]["verdict"]
    if v == "Priced above model value":
        return (f"the case for owning {T} at {_usd(c['price'])} rests on growth or margins well above the base "
                f"case ({_pct(A['g1'], 0)} year-one growth, {_pct(c['R']['inputs'].margin, 0)} margin).")
    if v == "Priced below model value":
        return (f"the price embeds less than the base case; the risk is that growth or margins fall short of "
                f"{_pct(A['g1'], 0)} and {_pct(c['R']['inputs'].margin, 0)}.")
    return "price and model agree; returns from here depend on delivering the base case."


def slide_performance(dk, c):
    ws, d = c["ws"], c["d"]
    r = (ws.get("period_ret") or 0) / 100
    dd = (ws.get("dd_peak_trough") or 0) / 100
    trough_i = int(np.argmin(c["cum"].values))
    trough_d = d["Date"].iloc[trough_i]
    recovered = c["cum"].iloc[-1] > -0.02
    title = f"{c['ticker']} returned {_pct(r, 1, True)} over {c['window_word']}"
    if dd < -0.15 and recovered:
        title += f", recovering fully from a {_pct(abs(dd), 0)} drawdown in {trough_d:%B %Y}"
    elif c["cum"].iloc[-1] < -0.10:
        title += f", and sits {_pct(abs(c['cum'].iloc[-1]), 0)} below its peak"
    s = dk.slide("Performance", title,
                 f"Source: {c['price_source']} daily prices, {c['start']:%b %Y} – {c['end']:%d %b %Y}; "
                 "QuantWizard calculations. Index = cumulative return, start = 100"
                 + ("; grey line = SPY." if c["spy_ret"] is not None else "."))
    _box(s, M, 1.78, 8.61, 0.25, f"{c['ticker']} cumulative return, index (start = 100)", size=12, bold=True)
    _image(s, chart_cumulative(d), M, 2.08, 8.61, 3.34)
    _panel_label(s, 9.50, 1.78, 3.26, "Return and risk, annualized")
    sh, so = ws.get("sharpe"), ws.get("sortino")
    _rows(s, 9.50, 2.11, 3.26, [
        ("Return", _pct((ws.get("ann_ret") or 0) / 100, 1, True), TEXT),
        ("Volatility", _pct((ws.get("ann_vol") or 0) / 100, 1), TEXT),
        ("Sharpe ratio", f"{sh:.2f}" if _num(sh) else "n/a", TEXT),
        ("Sortino ratio", f"{so:.2f}" if _num(so) else "n/a", TEXT),
        ("Max drawdown", _pct(dd, 1), RED)], label_w=1.96)
    note = ("Sortino above Sharpe: downside volatility ran below total volatility." if (_num(so) and _num(sh) and so > sh)
            else "Sortino at or below Sharpe: losses were as volatile as the stock overall.")
    _box(s, 9.50, 4.61, 3.26, 0.83, note, size=12, color=GREY, italic=True, lnspc=0.95)
    close = d["Close"].values
    peak_i = int(np.argmax(close[: trough_i + 1])) if trough_i > 0 else 0
    after = np.nonzero(close[trough_i:] >= close[peak_i])[0]
    if dd < -0.10:
        t1 = (f"The worst fall took the stock {_pct(abs(dd), 0)} below its {d['Date'].iloc[peak_i]:%B %Y} peak by "
              f"{trough_d:%B %Y}"
              + (f"; it was back at that peak by {d['Date'].iloc[trough_i + after[0]]:%B %Y}." if len(after)
                 else "; it has not regained that peak."))
        lead1 = "Drawdown was deep but contained." if (dd < -0.2 and len(after)) else (
            "Drawdown still open." if not len(after) else "Pullbacks stayed moderate.")
    else:
        t1 = f"The worst peak-to-trough fall in the window was only {_pct(abs(dd), 0)}."
        lead1 = "The path was steady."
    near = c["price"] / c["hi52"] - 1
    if near > -0.05:
        lead2, t2 = "The stock sits near its highs.", f"It closed within {_pct(abs(near), 1)} of its 52-week peak of {_usd(c['hi52'], 2)}."
    else:
        lead2, t2 = "Well off its highs.", f"The stock is {_pct(abs(near), 0)} below its 52-week peak of {_usd(c['hi52'], 2)}."
    if c["spy_ret"] is not None:
        t2 += f" SPY returned {_pct(c['spy_ret'], 0, True)} over the same window."
    _takeaways(s, [(lead1, t1), (lead2, t2)], h=0.67)


def slide_trend(dk, c):
    ma, p = c["ma"], c["price"]
    have = [n for n in (20, 50, 200) if ma.get(n)]
    above = [n for n in have if p > ma[n]]
    below_hi = p / c["hi52"] - 1
    if len(above) == len(have):
        pos = "above its " + ", ".join(f"{n}-" for n in have[:-1]) + (f" and {have[-1]}-day averages" if len(have) > 1 else f"{have[-1]}-day average")
    elif not above:
        pos = "below its " + ", ".join(f"{n}-" for n in have[:-1]) + (f" and {have[-1]}-day averages" if len(have) > 1 else f"{have[-1]}-day average")
    else:
        pos = f"above its {'/'.join(str(n) for n in above)}-day but below its {'/'.join(str(n) for n in have if n not in above)}-day averages"
    title = f"{c['ticker']} trades {pos} and {_pct(abs(below_hi), 1)} below its 52-week high"
    s = dk.slide("Trend", title,
                 f"Source: {c['price_source']} daily prices to {c['end']:%d %b %Y}; QuantWizard technical "
                 "calculations. MA = simple moving average; %B = position within 20-day, 2σ Bollinger Bands "
                 "(1.0 = upper band).")
    _box(s, M, 1.78, 8.61, 0.25, f"{c['ticker']} daily close and moving averages, $", size=12, bold=True)
    _image(s, chart_price_ma(c["d"]), M, 2.08, 8.61, 3.34)
    _panel_label(s, 9.50, 1.78, 3.26, "Price vs moving averages")
    for x, w, t, al in ((9.50, 1.18, "Average", "l"), (10.68, 1.11, "Level", "r"), (11.79, 0.97, "Price above", "r")):
        _box(s, x, 2.14, w, 0.22, t, size=10, color=SRC, align=al)
    y = 2.39
    for n in (20, 50, 200):
        if not ma.get(n):
            continue
        gap = p / ma[n] - 1
        _box(s, 9.50, y, 1.25, 0.39, f"{n}-day", size=14, color=GREY, anchor="ctr")
        _box(s, 10.68, y, 1.11, 0.39, _usd(ma[n], 2), size=14, color=TEXT, align="r", anchor="ctr")
        _box(s, 11.79, y, 0.97, 0.39, _pct(gap, 1, True), size=14, bold=True,
             color=(NAVY if n == 200 and gap > 0 else RED if gap < 0 else TEXT), align="r", anchor="ctr")
        _line(s, 9.50, y + 0.42, 12.76, y + 0.42)
        y += 0.44
    _box(s, 9.50, 3.83, 3.26, 0.25, "52-week range", size=12, bold=True)
    lo, hi = c["lo52"], c["hi52"]
    frac = min(1, max(0, (p - lo) / (hi - lo))) if hi > lo else 0.5
    _rect(s, 9.56, 4.25, 3.15, 0.08, RULE)
    _rect(s, 9.56, 4.25, max(0.02, 3.15 * frac), 0.08, NAVY)
    _rect(s, 9.56 + 3.15 * frac - 0.085, 4.21, 0.17, 0.17, NAVY, MSO_SHAPE.OVAL)
    _box(s, min(11.76, max(9.5, 9.56 + 3.15 * frac - 0.5)), 3.97, 1.0, 0.22, _usd(p, 2), size=11, bold=True,
         color=NAVY, align="c")
    _box(s, 9.50, 4.39, 1.53, 0.22, _usd(lo, 2), size=11, color=GREY)
    _box(s, 11.24, 4.39, 1.53, 0.22, _usd(hi, 2), size=11, color=GREY, align="r")
    _box(s, 9.50, 4.83, 3.26, 0.30, f"RSI (14) {c['rsi']:.1f}  ·  %B {c['pctb']:.2f}" if _num(c["pctb"]) else
         f"RSI (14) {c['rsi']:.1f}", size=14, bold=True)
    rsi = c["rsi"]
    note = ("Overbought: RSI above 70" if rsi >= 70 else "Oversold: RSI below 30" if rsi <= 30 else
            "Positive momentum; below 70 (overbought)" if rsi >= 50 else "Soft momentum; above 30 (oversold)")
    _box(s, 9.50, 5.12, 3.26, 0.33, note, size=12, color=GREY, italic=True)
    m200, prev = ma.get(200), c["ma200_prev"]
    if m200:
        rising = prev is not None and m200 > prev
        closes = c["d"]["Close"]
        m200s = closes.rolling(200).mean()
        side = closes > m200s
        streak = 0
        for v in side.values[::-1]:
            if bool(v) == bool(side.values[-1]):
                streak += 1
            else:
                break
        if p > m200 and rising:
            lead1, t1 = "The uptrend is intact.", (f"The 200-day average has risen over the last three months, "
                                                    f"and price has held above it for {streak} trading days.")
        elif p > m200:
            lead1, t1 = "Above trend, but the trend is flat.", "The 200-day average has not risen over the last three months."
        else:
            lead1, t1 = "The trend has turned.", (f"Price has been below the 200-day average for {streak} trading "
                                                  f"days, and the average is {'still rising' if rising else 'falling'}.")
        gap = p / m200 - 1
        if gap > 0:
            lead2 = "The gap to the long-term trend is wide." if gap > 0.10 else "Close to the long-term trend."
            t2 = (f"At {_pct(gap, 0)} above the 200-day average, the stock could fall about {_pct(1 - m200 / p, 0)} "
                  "and still sit on its long-term trend line.")
        else:
            lead2 = "Below the long-term trend."
            t2 = f"A close back above the {_usd(m200)} 200-day average would repair the trend."
        _takeaways(s, [(lead1, t1), (lead2, t2)])


def slide_risk(dk, c):
    ws, d = c["ws"], c["d"]
    dd60 = c["dd60"]
    # An episode starts below -20% and ends only once the stock is back above
    # -10%, so a drawdown hovering around the line counts once.
    episodes = []
    in_ep = False
    for dt, v in zip(d["Date"], dd60):
        if not _num(v):
            continue
        if not in_ep and v < -0.20:
            episodes.append([dt, v])
            in_ep = True
        elif in_ep and v < episodes[-1][1]:
            episodes[-1] = [dt, v]
        elif in_ep and v > -0.10:
            in_ep = False
    worst = float(dd60.min()) if dd60.notna().any() else None
    if len(episodes) == 1:
        title = f"The {episodes[0][0]:%B %Y} sell-off was the only drawdown in {c['window_word']} deeper than 20%"
    elif len(episodes) > 1:
        e = min(episodes, key=lambda x: x[1])
        title = f"{len(episodes)} sell-offs deeper than 20% in {c['window_word']}; the worst came in {e[0]:%B %Y}"
    else:
        title = f"No 60-day drawdown in {c['window_word']} went deeper than {_pct(abs(worst or 0), 0)}"
    s = dk.slide("Risk", title,
                 f"Source: {c['price_source']} daily prices and volume, {c['start']:%b %Y} – {c['end']:%d %b %Y}; "
                 "QuantWizard calculations. Max drawdown is peak-to-trough over the full window; the chart shows "
                 f"the rolling 60-day drawdown, which bottomed at {_pct(worst, 1)}.")
    _box(s, M, 1.75, 8.61, 0.25, f"{c['ticker']} rolling 60-day drawdown, %", size=12, bold=True)
    _image(s, chart_drawdown(d), M, 2.03, 8.61, 1.59)
    if "Volume" in d.columns and d["Volume"].fillna(0).sum() > 0:
        _box(s, M, 3.81, 8.61, 0.25, f"{c['ticker']} daily volume, shares, with 20-day average", size=12, bold=True)
        _image(s, chart_volume(d), M, 4.08, 8.61, 2.11)
    _panel_label(s, 9.50, 1.75, 3.26, "Downside and volatility")
    volf = (ws.get("ann_vol") or 0) / 100
    _rows(s, 9.50, 2.08, 3.26, [
        ("Max drawdown", _pct((ws.get("dd_peak_trough") or 0) / 100, 1), RED),
        ("Worst 60-day drawdown", _pct(worst, 1), RED),
        ("Volatility, 20-day", _pct(c["vol20"], 1), TEXT),
        ("Volatility, full period", _pct(volf, 1), TEXT)], label_w=2.22)
    n1 = (("One deep episode.", "Every other pullback in the window stayed shallower than −20%.") if len(episodes) == 1
          else ("Repeated stress.", f"{len(episodes)} separate falls went deeper than −20%.") if len(episodes) > 1
          else ("No deep episode.", "No 60-day fall went past −20%."))
    diff = (c["vol20"] - volf) * 100
    n2 = (("Calmer today.", f"Recent volatility runs {abs(diff):.0f} points below the period average.") if diff < -2
          else ("Rougher today.", f"Recent volatility runs {diff:.0f} points above the period average.") if diff > 2
          else ("Typical conditions.", "Recent volatility is close to the period average."))
    notes = [n1, n2]
    if "Volume" in d.columns and d["Volume"].fillna(0).sum() > 0 and dd60.notna().any():
        v20 = d["Volume"].rolling(20).mean()
        med = float(v20.median())
        ti = int(np.nanargmin(dd60.values.astype(float))) if dd60.notna().any() else None
        if ti is not None and med:
            near = v20.iloc[max(0, ti - 30): ti + 30]
            if len(near) and near.max() > 1.3 * med:
                notes.append(("Volume confirms stress.",
                              f"Around the {d['Date'].iloc[ti]:%B %Y} low the 20-day average reached about "
                              f"{near.max() / 1e6:,.0f}M shares, against a usual {med / 1e6:,.0f}M."))
            elif len(near):
                notes.append(("Orderly selling.", "Volume around the worst drawdown stayed close to normal."))
    _notes(s, 9.50, 4.11, 3.26, 2.22, notes)


def slide_fundamentals(dk, c):
    fb, f = c["fb"], c["f"]
    g = f.get("growth") or {}
    basis = "TTM" if g.get("yoy_basis") in ("ttm", "ttm_eps_fy") else (fb.get("fy") or "FY")
    rg, eg = g.get("revenue_yoy"), g.get("eps_yoy")
    rc, ec = g.get("revenue_cagr"), g.get("eps_cagr")
    title = (f"EPS grew {_pct(eg / 100, 1)} over the last twelve months on revenue growth of {_pct(rg / 100, 1)}"
             if basis == "TTM" and _num(eg) and _num(rg) else
             f"EPS grew {_pct(eg / 100, 1)} in {basis} on revenue growth of {_pct(rg / 100, 1)}"
             if _num(eg) and _num(rg) else f"{c['short']}: growth, margins and balance-sheet quality")
    pb = (f.get("valuation") or {}).get("pb")
    s = dk.slide("Fundamentals", title,
                 f"Source: SEC EDGAR filings (fiscal year ended {fb['fy_end']:%d %b %Y}"
                 + (f"; TTM to {fb['flows_end']:%d %b %Y}" if fb.get("has_ttm") else "")
                 + f"); QuantWizard calculations. CAGR over the filed years ({(fb.get('fy_labels') or ['?'])[0]}–{fb.get('fy')})."
                 + (f" *ROE is inflated by a small book-equity base (P/B {pb:.0f}x) and is not comparable across "
                    "companies." if _num(pb) and pb > 20 else "")
                 if fb.get("fy_end") is not None else "Source: SEC EDGAR filings.")
    _box(s, M, 1.75, 6.11, 0.25, f"Growth rates, {basis}", size=12, bold=True)
    cats, vals = [], []
    for lab, v in ((f"Revenue, {basis} YoY", rg), ("Revenue, CAGR", rc), (f"EPS, {basis} YoY", eg), ("EPS, CAGR", ec)):
        if _num(v):
            cats.append(lab)
            vals.append(v)
    if vals:
        hi = int(np.argmax(vals))
        _native_bar(s, M, 2.03, 6.11, 3.33, cats, vals, '0.0"%"', True, hi)
    _panel_label(s, 7.22, 1.75, 5.54, f"Quality scorecard, {basis}")
    m, q, lv, rr = f.get("margins") or {}, f.get("quality") or {}, f.get("leverage") or {}, f.get("returns") or {}
    fcf = (f.get("fcf") or {}).get("fcf")
    groups = [("Profitability", [("Gross margin", _pct((m.get("gross") or 0) / 100, 1) if _num(m.get("gross")) else "n/a"),
                                 ("Operating margin", _pct(m["operating"] / 100, 1) if _num(m.get("operating")) else "n/a"),
                                 ("Net margin", _pct(m["net"] / 100, 1) if _num(m.get("net")) else "n/a")]),
              ("Cash and quality", [("Free cash flow", f"${fcf / 1e9:,.1f}B" if _num(fcf) else "n/a"),
                                    ("Piotroski F-score", f"{q['f_score']} of 9" if q.get("f_score") is not None else "n/a"),
                                    ("Altman Z-score", f"{q['z_score']:.2f} ({q.get('z_zone')})" if _num(q.get("z_score")) else "n/a")]),
              ("Balance sheet", [("Current ratio", f"{lv['current_ratio']:.2f}" if _num(lv.get("current_ratio")) else "n/a"),
                                 ("Debt / equity", f"{lv['debt_to_equity']:.2f}" if _num(lv.get("debt_to_equity")) else "n/a"),
                                 ("Return on equity" + ("*" if _num(pb) and pb > 20 else ""),
                                  _pct(rr["roe"] / 100, 1) if _num(rr.get("roe")) else "n/a")])]
    y = 2.14
    for gname, items in groups:
        _box(s, 7.22, y, 5.54, 0.22, gname.upper(), size=10, bold=True, color=TAG, spc=100)
        y += 0.25
        for lab, v in items:
            _box(s, 7.22, y, 3.33, 0.33, lab, size=14, color=GREY)
            _box(s, 9.99, y, 2.77, 0.33, v, size=14, bold=True, color=TEXT, align="r")
            _line(s, 7.22, y + 0.33, 12.76, y + 0.33)
            y += 0.36
        y += 0.1
    if _num(eg) and _num(rg) and rg > 0 and eg > rg * 1.5:
        lead, text = (f"EPS outgrew revenue by {eg / rg:.1f}x.",
                      "Since EPS = sales × net margin ÷ shares, the gap must come from wider margins or fewer "
                      "shares; it narrows if either stalls.")
    elif _num(eg) and _num(rg) and eg < rg:
        lead, text = ("Revenue outgrew EPS.", "Margins or the share count moved against earnings; growth in sales "
                                              "did not fully reach the bottom line.")
    else:
        lead, text = ("Earnings tracked sales.", "EPS and revenue grew at similar rates, so margins and the share "
                                                 "count held roughly steady.")
    _box(s, M, 5.44, 6.11, 0.89, [(lead + " ", {"bold": True, "color": NAVY}), (text, {"color": TEXT})], size=14,
         lnspc=0.95)


def slide_valuation(dk, c):
    A, p, fb = c["A"], c["price"], c["fb"]
    ttm = fb.get("ttm") or {}
    pe = (p / ttm["eps_clean"]) if (_num(ttm.get("eps_clean")) and ttm["eps_clean"] > 0) else None
    f = c["f"]
    fcf = (f.get("fcf") or {})
    fcfy = (fcf.get("fcf_yield") / 100) if _num(fcf.get("fcf_yield")) else None
    if c["dcf_ok"]:
        sc, rv, W_ = c["model"]["scenarios"], c["model"]["reverse"], c["R"]["wacc"]
        ig = rv.get("growth")
        if _num(ig):
            # "Over 3 times" a 0.8% base case is true and says little: below 3%
            # the gap reads better in points.
            ratio = ig / A["g1"] if A["g1"] >= 0.03 else None
            if ratio and ratio >= 2:
                rel = f"over {int(ratio)} times"
            elif A["g1"] < 0.03 and abs(ig - A["g1"]) >= 0.001:
                rel = f"{abs(ig - A['g1']) * 100:.1f} points {'above' if ig > A['g1'] else 'below'}"
            else:
                rel = "above" if ig > A["g1"] else "below"
            # Small rates get a decimal, so "2.8%, 2.0 points above 0.8%" adds up.
            title = (f"Today's {_usd(p)} price requires {_pct(ig, 1 if A['g1'] < 0.03 else 0)} year-one "
                     f"revenue growth, {rel} the model's {_pct(A['g1'], 1)} base case")
        else:
            title = f"Today's {_usd(p)} price sits outside what the model can reach on growth alone"
        src = (f"Source: QuantWizard revenue-driven DCF: {_pct(W_['wacc'], 1)} discount rate, 10-year explicit "
               f"period, {_pct(A['tg'], 1)} terminal growth, long-run operating margin {_pct(c['R']['inputs'].margin, 0)} "
               f"({A['anchor']} anchor), long-run capex {_pct(A['lr_capex'], 1)} of revenue. Multiples: SEC EDGAR "
               f"and the {c['end']:%d %b %Y} price. A model output, not a price target; full assumptions and "
               "sensitivity grids in the workbook's DCF and Scenarios tabs.")
    else:
        pb = (f.get("valuation") or {}).get("pb")
        parts = [f"{_mult(pe)} earnings"] if _num(pe) else []
        if _num(pb):
            parts.append(f"{_mult(pb)} book value")
        elif _num(f.get("ev_ebitda")):
            parts.append(f"{_mult(f['ev_ebitda'])} EBITDA")
        title = (f"Without a DCF, {c['ticker']} is judged on multiples: " + " and ".join(parts)
                 if parts else f"Without a DCF, {c['ticker']} is judged on its multiples")
        src = f"Source: SEC EDGAR filings and the {c['end']:%d %b %Y} price. No DCF: {c['R'].get('dcf_reason')}."
    s = dk.slide("Valuation", title, src)
    if c["dcf_ok"]:
        vals = [sc["bear"]["fair_value"], sc["base"]["fair_value"], sc["bull"]["fair_value"], p]
        base = vals[1]
        rel = (f"  ·  price is {p / base:.1f}x the base case" if base and p / base >= 1.15 else
               f"  ·  price is {1 - p / base:.0%} below the base case" if base and p / base <= 0.87 else "")
        _box(s, M, 1.75, 6.53, 0.25, [("DCF value per share vs current price, $", {"bold": True}),
                                      (rel, {"bold": False, "color": GREY})], size=12)
        _native_bar(s, M, 2.03, 6.53, 3.47, ["DCF bear", "DCF base", "DCF bull", "Share price"], vals,
                    '"$"#,##0', False, 3)
        _panel_label(s, 7.50, 1.75, 5.26, "What the price implies")
        hot = _num(ig) and ig > 2 * max(A["g1"], 0.01)
        _box(s, 7.50, 2.17, 2.50, 0.61, _pct(ig, 1) if _num(ig) else "n/a", size=32, bold=True,
             color=RED if hot else NAVY)
        _box(s, 7.50, 2.78, 2.50, 0.50, f"year-one revenue growth needed to justify {_usd(p)}", size=12, color=GREY,
             lnspc=0.95)
        _line(s, 10.14, 2.22, 10.14, 3.22, GREY, 12700)
        _box(s, 10.35, 2.17, 2.42, 0.61, _pct(A["g1"], 1), size=32, bold=True, color=NAVY)
        _box(s, 10.35, 2.78, 2.42, 0.50, "year-one revenue growth in the model's base case", size=12, color=GREY,
             lnspc=0.95)
        y0 = 3.47
    else:
        y0 = 1.75
        roe = (f.get("returns") or {}).get("roe")
        pb = (f.get("valuation") or {}).get("pb")
        fin_co = "financial" in str(c["R"].get("dcf_reason") or "")
        _rect(s, M, 1.75, 6.53, 3.9, SIDE)
        _box(s, M + 0.22, 1.95, 6.1, 0.25, "WHY THERE IS NO DCF", size=10, bold=True, color=TAG, spc=100)
        why = str(c["R"].get("dcf_reason") or "the model does not apply")
        body = [[("The model does not fit. ", {"bold": True, "color": NAVY}),
                 (why[:1].upper() + why[1:] + ".", {"color": TEXT})]]
        if fin_co and _num(roe) and _num(pb):
            body.append([("The better lens. ", {"bold": True, "color": NAVY}),
                         (f"For a bank or insurer, price-to-book against return on equity: {c['short']} "
                          f"earns {_pct(roe / 100, 1)} on equity and trades at {_mult(pb)} book.",
                          {"color": TEXT})])
            body.append([("Rule of thumb. ", {"bold": True, "color": NAVY}),
                         ("A bank earning its cost of equity (~10%) is worth about 1x book; higher ROE "
                          "justifies a premium.", {"color": TEXT})])
        else:
            body.append([("What to use instead. ", {"bold": True, "color": NAVY}),
                         ("The trading multiples at right, the peer comparison and the company's own "
                          "multiple history.", {"color": TEXT})])
        _box(s, M + 0.22, 2.3, 6.1, 3.2, body, size=14, bullets=True, lnspc=0.95, space_after=10)
    _box(s, 7.50, y0, 5.26, 0.22, "TRADING MULTIPLES", size=10, bold=True, color=TAG, spc=100)
    mrows = [("Price / earnings" + (" (excl. gains)" if fb.get("gains_material") else ""), _mult(pe), TEXT)]
    if _num(f.get("ev_ebitda")):
        mrows.append(("EV / EBITDA", _mult(f.get("ev_ebitda")), TEXT))
    mrows.append(("Price / sales", _mult((f.get("valuation") or {}).get("ps")), TEXT))
    if not c["dcf_ok"]:
        mrows.append(("Price / book", _mult((f.get("valuation") or {}).get("pb")), TEXT))
        dy = (f.get("capital_return") or {}).get("dividend_yield")
        mrows.append(("Dividend yield", _pct(dy / 100, 1) if _num(dy) else "n/a", TEXT))
    if _num(fcfy):
        mrows.append(("Free cash flow yield", _pct(fcfy, 1), TEXT))
    y = _rows(s, 7.50, y0 + 0.25, 5.26, mrows, label_w=3.16, row_h=0.39, value_size=14)
    if _num(fcfy) and fcfy > 0:
        _box(s, 7.50, y + 0.1, 5.26, 0.69, f"A {_pct(fcfy, 1)} FCF yield means investors pay about "
                                           f"${1 / fcfy:,.0f} for each $1 of current free cash flow.",
             size=12, color=GREY, italic=True)
    if c["dcf_ok"]:
        bear, bull = sc["bear"]["fair_value"], sc["bull"]["fair_value"]
        cagr = c["fb"].get("rev_cagr5")
        if bull < p:
            lead, text = (f"Even the bull case sits {_pct(1 - bull / p, 0)} below the price.",
                          f"Closing the gap needs growth far above {c['short']}'s "
                          + (f"{_pct(cagr, 1)} five-year revenue CAGR." if _num(cagr) else "recent revenue growth."))
        elif bear > p:
            lead, text = ("Even the bear case is above the price.",
                          "The market is pricing in less than the model's most cautious scenario.")
        else:
            lead, text = ("The price sits inside the scenario range.",
                          f"Bear {_usd(bear)} to bull {_usd(bull)}; probability-weighted value "
                          f"{_usd(sc['prob_weighted'])}.")
        _box(s, M, 5.61, 6.53, 0.75, [(lead + " ", {"bold": True, "color": NAVY}), (text, {"color": TEXT})],
             size=14, lnspc=0.95)


def slide_valuation_range(dk, c):
    """Extra: the workbook's football field and the WACC × growth grid."""
    if not c["dcf_ok"]:
        return
    m, p, fb, R = c["model"], c["price"], c["fb"], c["R"]
    sc, sens = m["scenarios"], m["sensitivity"]
    rows = [("52-week range", c["lo52"], c["hi52"]),
            ("DCF bear–bull", sc["bear"]["fair_value"], sc["bull"]["fair_value"])]
    g = [v for row in sens["wacc_tg"]["grid"] for v in row if _num(v)]
    if g:
        rows.append(("DCF: WACC × terminal g", min(g), max(g)))
    g2 = [v for row in sens["growth_margin"]["grid"] for v in row if _num(v)]
    if g2:
        rows.append(("DCF: growth × margin", min(g2), max(g2)))
    eps = (fb.get("ttm") or {}).get("eps_clean")
    pes = [x["pe"] for x in R["peers"] if _num(x.get("pe")) and x["pe"] > 0]
    if _num(eps) and eps > 0 and pes:
        rows.append(("Peer P/E × clean EPS", min(pes) * eps, max(pes) * eps))
    vh = [r["pe_avg"] for r in ((c["vhist"] or {}).get("rows") or []) if _num(r.get("pe_avg"))]
    if _num(eps) and eps > 0 and vh:
        rows.append(("Own P/E history × EPS", min(vh) * eps, max(vh) * eps))
    mids = sorted((lo + hi) / 2 for _, lo, hi in rows[1:])
    lo_all = mids[0] if mids else None
    hi_all = mids[-1] if mids else None
    title = (f"Methods centre between {_usd(lo_all)} and {_usd(hi_all)} a share; the price is {_usd(p)}"
             if mids else f"Valuation range against the {_usd(p)} price")
    s = dk.slide("Valuation range", title,
                 "Source: QuantWizard model (DCF scenarios and sensitivity grids), peer multiples from each "
                 "peer's filings, P/E history from filed EPS and yearly prices. Bars show ranges, not targets.")
    _box(s, M, 1.75, 6.9, 0.25, "Value per share by method, $", size=12, bold=True)
    _image(s, chart_football(rows, p), M, 2.03, 6.9, 3.9)
    _panel_label(s, 7.75, 1.75, 5.01, "Base value: WACC × terminal growth")
    st = sens["wacc_tg"]
    cw = 5.01 / 6
    _box(s, 7.75, 2.14, cw, 0.3, "WACC ↓  g →", size=9, color=SRC)
    for j, tg in enumerate(st["cols"]):
        _box(s, 7.75 + cw * (j + 1), 2.14, cw, 0.3, _pct(tg, 1), size=11, bold=True, color=GREY, align="r")
    for i, w in enumerate(st["rows"]):
        y = 2.48 + i * 0.42
        _box(s, 7.75, y, cw, 0.38, _pct(w, 1), size=11, bold=True, color=GREY, anchor="ctr")
        for j, v in enumerate(st["grid"][i]):
            centre = (i == 2 and j == 2)
            if centre:
                _rect(s, 7.75 + cw * (j + 1), y, cw, 0.38, PANEL)
            _box(s, 7.75 + cw * (j + 1), y, cw - 0.05, 0.38, _usd(v), size=12, bold=centre,
                 color=(NAVY if centre else (TEXT if (_num(v) and v < p) else GREEN_T)), align="r", anchor="ctr")
        _line(s, 7.75, y + 0.40, 12.76, y + 0.40)
    _box(s, 7.75, 4.66, 5.01, 0.6, f"Green = at or above today's {_usd(p)} price. Centre = base case. "
                                   "Each row adds 0.5 point to the discount rate.", size=11, color=GREY, italic=True)


def slide_outlook(dk, c):
    mc, p = c["mc"], c["price"]
    if not _num(mc["p50"]):
        return
    chg = mc["p50"] / p - 1
    title = (f"Simulations center on a {_pct(abs(chg), 0)} {'gain' if chg >= 0 else 'loss'} over 12 months, "
             f"within a wide {_usd(mc['p5'])}–{_usd(mc['p95'])} range")
    s = dk.slide("12-month outlook", title,
                 f"Source: QuantWizard Monte Carlo, {mc['n']:,} paths over {mc['days']} trading days from a "
                 f"{_usd(mc['start'] or p, 2)} starting price, using {_pct(mc['vol'], 1)} annualized volatility "
                 "from the price history and a CAPM drift. A distribution of outcomes, not a forecast or price target.")
    _box(s, M, 1.75, 7.50, 0.25, f"Simulated {c['ticker']} price paths, $ (percentile bands)", size=12, bold=True)
    if c["mc_df"] is not None:
        _image(s, chart_monte_carlo(c["mc_df"]), M, 2.03, 7.50, 4.11)
    _panel_label(s, 8.47, 1.75, 4.29, "Price after 12 months, by percentile")
    base = mc["start"] or p
    y = 2.11
    for lab, key in (("95th percentile", "p95"), ("75th percentile", "p75"), ("Median", "p50"),
                     ("5th percentile", "p5")):
        v = mc[key]
        ch = v / base - 1 if _num(v) else None
        if key == "p50":
            _rect(s, 8.47, y, 4.29, 0.47, PANEL)
        _box(s, 8.58, y, 2.08, 0.47, lab, size=14, bold=(key == "p50"), color=NAVY if key == "p50" else GREY, anchor="ctr")
        _box(s, 10.56, y, 1.11, 0.47, _usd(v), size=16, bold=True, color=TEXT, align="r", anchor="ctr")
        _box(s, 11.67, y, 1.00, 0.47, _pct(ch, 0, True), size=14, bold=True,
             color=(RED if (_num(ch) and ch < 0) else NAVY if key == "p50" else TEXT), align="r", anchor="ctr")
        _line(s, 8.47, y + 0.49, 12.76, y + 0.49)
        y += 0.50
    _box(s, 8.47, 4.19, 1.67, 0.61, _pct(mc["prob"], 1), size=30, bold=True, color=NAVY)
    _box(s, 10.17, 4.22, 2.60, 0.56, "of paths end above the starting price", size=13, color=GREY)
    prob = mc["prob"] or 0
    n1 = (("Barely better than even odds.", "The median gain is small relative to the spread.") if 0.5 <= prob < 0.6
          else ("Odds favour a gain.", "More than 60% of paths finish higher.") if prob >= 0.6
          else ("Odds lean to a loss.", "Fewer than half of the paths finish higher."))
    up = (mc["p95"] / base - 1) if _num(mc["p95"]) else None
    dn = (1 - mc["p5"] / base) if _num(mc["p5"]) else None
    n2 = ("Skew is built in.", f"The upside tail ({_pct(up, 0, True)}) is larger than the downside tail "
                               f"({_pct(-dn, 0, True)}), a property of compounding returns, not a view on {c['short']}.") \
        if (_num(up) and _num(dn)) else ("Wide range.", "Outcomes spread widely around the median.")
    _notes(s, 8.47, 4.94, 4.29, 1.33, [n1, n2])


def slide_signposts(dk, c):
    fb, A, ma = c["fb"], c["A"], c["ma"]
    ttm = fb.get("ttm") or {}
    s = dk.slide("Signposts", "Four signposts would change the view; growth and margins matter most",
                 "Source: figures from the preceding slides (SEC EDGAR; prices to "
                 f"{c['end']:%d %b %Y}; QuantWizard DCF and risk calculations). Signposts and gaps are rule-based "
                 "readings of the data, not model outputs.")
    _box(s, M, 1.75, 8.06, 0.25, "What to watch", size=12, bold=True)
    rv = (c["model"] or {}).get("reverse") or {}
    cagr = fb.get("rev_cagr5")
    om = ttm.get("op_margin")
    g = c["f"].get("growth") or {}
    m200 = ma.get(200)
    gap = (c["price"] / m200 - 1) if m200 else None
    worst = float(c["dd60"].min()) if c["dd60"].notna().any() else None
    rows = [
        ("Revenue growth",
         (f"{_pct(A.get('g1'), 1)} base case vs {_pct(rv.get('growth'), 1)} implied" if c["dcf_ok"]
          else f"{_pct(ttm.get('rev_growth'), 1)} over the last twelve months"),
         f"Growth runs well above the {_pct(cagr, 1)} five-year pace" if _num(cagr) else "Growth accelerates",
         "Growth tracks the base case; the gap to the price persists" if c["dcf_ok"] else "Growth slows"),
        ("Margins",
         f"{_pct(om, 1)} operating margin; EPS {_pct((g.get('eps_yoy') or 0) / 100, 1, True)} vs revenue "
         f"{_pct((g.get('revenue_yoy') or 0) / 100, 1, True)}",
         "Margins keep widening", "Margins plateau and EPS growth converges toward revenue growth"),
        ("Price trend",
         f"{_pct(abs(gap), 1)} {'above' if gap >= 0 else 'below'} the {_usd(m200)} 200-day average" if _num(gap) else "n/a",
         "Price holds above the 200-day average", "A close below the 200-day average"),
        ("Volatility",
         f"{_pct(c['vol20'], 1)} 20-day vs {_pct((c['ws'].get('ann_vol') or 0) / 100, 1)} annualized",
         "Volatility stays below its long-run level",
         f"Volatility returns to stress levels ({_pct(worst, 0)} 60-day drawdown)" if _num(worst) else
         "Volatility rises above its long-run level")]
    tbl = s.shapes.add_table(5, 4, Inches(M), Inches(2.06), Inches(8.06), Inches(3.81)).table
    widths = [1.55, 1.97, 2.27, 2.27]
    for j, wv in enumerate(widths):
        tbl.columns[j].width = Inches(wv)
    heads = ["Signpost", "Today", "Supports upside if", "Supports downside if"]
    _style_table(tbl, [heads] + [list(r) for r in rows],
                 col_colors=[TEXT, GREY, GREEN_T, RED], first_bold=True)
    tbl.rows[0].height = Inches(0.36)
    for i in range(1, 5):
        tbl.rows[i].height = Inches(0.86)
    gaps = []
    if not c["R"]["peers"]:
        gaps.append(("Peer multiples.", "No peer data was available for comparison."))
    if not (c["R"].get("segments") or {}).get("rows"):
        gaps.append(("Segment mix.", "The filings carry no segment breakdown in XBRL."))
    if not (c["vhist"] or {}).get("rows"):
        gaps.append(("Multiple history.", "Not enough years of prices and EPS."))
    if not c["R"].get("consensus"):
        gaps.append(("Street estimates.", "No analyst revenue consensus was available."))
    _rect(s, 9.03, 2.06, 3.74, 3.81, SIDE)
    if gaps:
        _box(s, 9.22, 2.22, 3.35, 0.25, "DATA GAPS TO CLOSE", size=10, bold=True, color=TAG, spc=100)
        _box(s, 9.22, 2.50, 3.35, 0.33, "Before acting on this review, add:", size=13, color=GREY, italic=True)
        _box(s, 9.22, 2.89, 3.35, 2.92, [[(a + " ", {"bold": True, "color": NAVY}), (b, {"color": TEXT})]
                                          for a, b in gaps], size=14, bullets=True, lnspc=0.95, space_after=6)
    else:
        _box(s, 9.22, 2.22, 3.35, 0.25, "MODEL CHECKS", size=10, bold=True, color=TAG, spc=100)
        checks = (c["model"] or {}).get("checks", {})
        items = checks.get("items") or []
        _box(s, 9.22, 2.50, 3.35, 0.33, checks.get("integrity", "No DCF"), size=13, color=GREY, italic=True)
        short = {"Return on new capital between WACC and 50%": "Return on new capital in range",
                 "Closed-form recompute = table": "Model recompute matches the table",
                 "Net cash = cash & securities − debt": "Net cash ties to the balance sheet",
                 "Latest price is the last Data row": "Price is the latest close"}
        items = [(short.get(lab, lab), st) for lab, st in items]
        y = 2.89
        for lab, st in items[:7]:
            _box(s, 9.22, y, 3.35, 0.36, [(st[:1] + "  ", {"bold": True, "color": NAVY if st.startswith("✓") else RED}),
                                          (lab, {"color": TEXT})], size=12)
            y += 0.4


def _style_table(tbl, data, col_colors=None, first_bold=False, header_color=NAVY, size=14):
    tbl.first_row = True
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.background()
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            r = tf.paragraphs[0].add_run()
            r.text = str(val if val is not None else "")
            r.font.name = FONT
            r.font.size = Pt(11 if i == 0 else size)
            r.font.bold = (i == 0) or (first_bold and j == 0)
            color = header_color if i == 0 else ((col_colors or [TEXT] * 8)[j] if col_colors else TEXT)
            r.font.color.rgb = _rgb(color)
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.find(qn(tag))
                if ln is not None:
                    tcPr.remove(ln)
            for tag, show, col, wv in (("a:lnL", False, None, 0), ("a:lnR", False, None, 0),
                                       ("a:lnT", i == 0, NAVY, 12700), ("a:lnB", True, NAVY if i == 0 else RULE,
                                                                        12700 if i == 0 else 6350)):
                ln = tcPr.makeelement(qn(tag), {"w": str(wv)})
                if show:
                    sf = ln.makeelement(qn("a:solidFill"), {})
                    sf.append(sf.makeelement(qn("a:srgbClr"), {"val": col}))
                    ln.append(sf)
                else:
                    ln.append(ln.makeelement(qn("a:noFill"), {}))
                tcPr.append(ln)
    # plain table style (no banding)
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "1")


def slide_technical(dk, c):
    s = dk.slide("Appendix", "Technical detail: Bollinger Bands and RSI",
                 f"Source: {c['price_source']} daily prices, {c['start']:%b %Y} – {c['end']:%d %b %Y}; QuantWizard "
                 "calculations. Bollinger Bands: 20-day average ± 2 standard deviations. RSI: 14-day relative "
                 "strength index; 70 = overbought, 30 = oversold.")
    _box(s, M, 1.72, 8.89, 0.25, f"{c['ticker']} daily close with 20-day Bollinger Bands, $", size=12, bold=True)
    _image(s, chart_bollinger(c["d"]), M, 1.99, 8.89, 2.36)
    _box(s, M, 4.39, 8.89, 0.25, "14-day RSI", size=12, bold=True)
    _image(s, chart_rsi(c["d"]), M, 4.65, 8.89, 1.57)
    _panel_label(s, 9.86, 1.72, 2.90, "Current readings")
    pb, rsi = c["pctb"], c["rsi"]
    _box(s, 9.86, 2.11, 2.90, 0.56, f"{pb:.2f}" if _num(pb) else "n/a", size=28, bold=True, color=NAVY)
    band = ("above the upper band" if _num(pb) and pb > 1 else "below the lower band" if _num(pb) and pb < 0 else
            "in the upper fifth of its band" if _num(pb) and pb >= 0.8 else "in the lower fifth of its band"
            if _num(pb) and pb <= 0.2 else "inside its band")
    _box(s, 9.86, 2.64, 2.90, 0.50, f"%B: price {band}", size=13, color=GREY)
    _line(s, 9.86, 3.22, 12.76, 3.22)
    _box(s, 9.86, 3.31, 2.90, 0.56, f"{rsi:.1f}", size=28, bold=True, color=NAVY)
    rtxt = ("overbought, above 70" if rsi >= 70 else "oversold, below 30" if rsi <= 30 else
            "positive momentum, below 70" if rsi >= 50 else "soft momentum, above 30")
    _box(s, 9.86, 3.83, 2.90, 0.50, f"RSI: {rtxt}", size=13, color=GREY)
    _line(s, 9.86, 4.42, 12.76, 4.42)
    reading = ("momentum is firm but not stretched." if 50 <= rsi < 70 else "momentum is stretched." if rsi >= 70
               else "momentum is weak." if rsi <= 30 else "momentum is soft.")
    _box(s, 9.86, 4.53, 2.90, 1.67, [("Reading: ", {"bold": True, "color": NAVY}),
                                     (reading + " These indicators describe the recent trend; they carry no view "
                                                "on value.", {"color": TEXT})], size=13, lnspc=0.95)


def slide_profile(dk, c):
    cd, fb, R = c["cd"], c["fb"], c["R"]
    s = dk.slide("Appendix", f"Company profile: {c['name']}",
                 f"Source: company filings (SEC EDGAR XBRL) and the Yahoo Finance profile; market "
                 f"capitalization at {c['end']:%d %b %Y}. Segment figures from the latest filing.")
    _box(s, M, 1.75, 7.22, 0.25, "How the business works", size=12, bold=True)
    _line(s, M, 2.06, M + 7.22, 2.06, NAVY, 12700)
    bullets = []
    desc = str(cd.get("Description") or "")
    if desc and desc != "N/A":
        import re as _re
        # Sentence ends, but not after the abbreviations company names carry.
        parts = _re.split(r"(?<!\bInc)(?<!\bCorp)(?<!\bCo)(?<!\bLtd)(?<!\bNo)(?<!\bU\.S)(?<!\bSt)\.\s+(?=[A-Z])",
                          desc)
        first = parts[0].strip().rstrip(".") + "."
        if len(first) > 170:
            first = first[:168].rsplit(" ", 1)[0] + "…"
        bullets.append([("What it does. ", {"bold": True, "color": NAVY}), (first, {"color": TEXT})])
    seg = R.get("segments") or {}
    srows = [r for r in (seg.get("rows") or []) if _num(r.get("rev"))]
    tot = sum(r["rev"] for r in srows) or 0
    per = "quarter" if seg.get("basis") == "quarter" else "year"
    for r in srows[:3]:
        share = r["rev"] / tot if tot else None
        gr = (r["rev"] / r["rev_prev"] - 1) if _num(r.get("rev_prev")) and r["rev_prev"] > 0 else None
        mg = (r["oi"] / r["rev"]) if _num(r.get("oi")) and r["rev"] else None
        txt = (f"{_pct(share, 0)} of revenue in the latest {per}"
               + (f", {_pct(gr, 0, True)} y/y" if _num(gr) else "")
               + (f"; {_pct(mg, 0)} operating margin" if _num(mg) else "") + ".")
        bullets.append([(f"{r['name']}. ", {"bold": True, "color": NAVY}), (txt, {"color": TEXT})])
    fcf = [x for x in (fb.get("fy_fcf") or []) if _num(x)]
    ret = [(b or 0) + (d or 0) for b, d in zip(fb.get("fy_buybacks") or [], fb.get("fy_dividends") or [])]
    if fcf and sum(fcf) > 0 and any(ret):
        bullets.append([("Where the cash goes. ", {"bold": True, "color": NAVY}),
                        (f"{sum(ret) / sum(fcf):.0%} of free cash flow returned to shareholders over "
                         f"{len(fcf)} fiscal years.", {"color": TEXT})])
    if bullets:
        _box(s, M, 2.19, 7.22, 4.17, bullets[:5], size=16, bullets=True, lnspc=0.95, space_after=10)
    _rect(s, 8.61, 1.75, 4.15, 4.58, SIDE)
    _box(s, 8.83, 1.92, 3.71, 0.25, "KEY FACTS", size=10, bold=True, color=TAG, spc=100)
    mcap = c["price"] * (fb.get("shares_now") or 0) * 1e9
    emp = cd.get("Employees")
    facts = [("Ticker", f"{c['ticker']} ({c['exchange'] or 'n/a'})"),
             ("Market cap", (f"${mcap / 1e12:,.2f}T" if mcap >= 1e12 else f"${mcap / 1e9:,.1f}B") if mcap else "n/a"),
             ("Employees", f"{int(emp):,}" if _num(emp) else "n/a"),
             ("Headquarters", str(cd.get("Country") or "n/a")),
             ("Sector", str(cd.get("Industry") or cd.get("Sector") or "n/a")[:26]),
             ("Fiscal year end", (lambda t: t[:1].upper() + t[1:])(RT.fy_end_desc(fb.get("fy_ends")))),
             ("Website", str(cd.get("Website") or "n/a").replace("https://", "").replace("http://", "").strip("/"))]
    y = 2.25
    for lab, v in facts:
        _box(s, 8.83, y, 1.53, 0.50, lab, size=13, color=GREY, anchor="ctr")
        _box(s, 10.28, y, 2.26, 0.50, v, size=14, bold=True, color=TEXT, align="r", anchor="ctr")
        _line(s, 8.83, y + 0.53, 12.54, y + 0.53)
        y += 0.556


def slide_news(dk, c):
    items = [n for n in c["news"] if n.get("Headline")][:7]
    s = dk.slide("Appendix", "Recent news and filings",
                 "Source: QuantWizard news feed (aggregators and SEC EDGAR 8-Ks). Headlines shown as published "
                 "and not verified by QuantWizard. Quote pages, headlines about other companies and off-topic "
                 "press releases are filtered out.")
    if not items:
        _box(s, M, 1.9, CW, 0.4, "No recent headlines passed the relevance filter.", size=14, color=GREY)
        return
    rows = [["Date", "Source", "Headline"]]
    for n in items:
        try:
            dt = pd.Timestamp(str(n.get("Date"))[:10]).strftime("%d %b %Y")
        except Exception:
            dt = str(n.get("Date") or "")
        rows.append([dt, str(n.get("Publisher") or ""), str(n.get("Headline") or "")[:110]])
    tbl = s.shapes.add_table(len(rows), 3, Inches(M), Inches(1.81), Inches(CW), Inches(0.51 * len(rows))).table
    for j, wv in enumerate((1.25, 1.95, 9.01)):
        tbl.columns[j].width = Inches(wv)
    _style_table(tbl, rows, col_colors=[GREY, TEXT, TEXT], size=13)
    for i in range(1, len(rows)):
        tbl.cell(i, 1).text_frame.paragraphs[0].runs[0].font.bold = True


def slide_disclosures(dk, c):
    s = dk.slide("Disclosures", "Important disclaimer", None)
    paras = ["This report has been generated by QuantWizard for informational and educational purposes only. "
             "It does not constitute financial, investment, legal, or tax advice. The information presented is "
             "derived from company filings (SEC EDGAR) and third-party market data and is believed to be accurate "
             "but is not guaranteed.",
             "Past performance is not indicative of future results. All investments involve risk, including the "
             "possible loss of principal. You should not make any investment decision based solely on the "
             "information in this report.",
             "QuantWizard is not a registered investment adviser, broker-dealer, or financial planner. Always "
             "consult a qualified financial professional before making investment decisions."]
    tb = _box(s, M, 1.94, 8.89, 2.36, paras[0], size=14, color=TEXT, lnspc=1.0)
    for ptxt in paras[1:]:
        p = tb.text_frame.add_paragraph()
        p.space_before = Pt(10)
        r = p.add_run()
        r.text = ptxt
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(14), _rgb(TEXT)
    _line(s, M, 4.42, M + CW, 4.42)
    _box(s, M, 4.53, 8.89, 0.28, f"© {c['date']:%Y} QuantWizard  ·  quantwizard.co", size=11, color=SRC)


def slide_business_mix(dk, c):
    seg = c["R"].get("segments") or {}
    srows = [r for r in (seg.get("rows") or []) if _num(r.get("rev"))]
    if len(srows) < 2:
        return
    tot = sum(r["rev"] for r in srows)
    top = max(srows, key=lambda r: r["rev"])
    per = "quarter" if seg.get("basis") == "quarter" else "fiscal year"
    fast = max((r for r in srows if _num(r.get("rev_prev")) and r["rev_prev"] > 0 and r["rev"] / tot > 0.05),
               key=lambda r: r["rev"] / r["rev_prev"], default=None)
    title = (f"{top['name']} is {_pct(top['rev'] / tot, 0)} of revenue"
             + (f"; {fast['name']} is growing fastest at {_pct(fast['rev'] / fast['rev_prev'] - 1, 0, True)}"
                if fast and fast is not top else ""))
    s = dk.slide("Business mix", title,
                 f"Source: segment revenue and operating income from the latest filing's XBRL (the {per} to "
                 f"{pd.Timestamp(seg.get('cur_end')):%d %b %Y}, against the same period a year earlier).")
    _box(s, M, 1.75, 6.5, 0.25, f"Revenue by segment, latest {per}, $B", size=12, bold=True)
    _native_bar(s, M, 2.03, 6.5, 3.6, [r["name"][:28] for r in srows[:6]],
                [r["rev"] / 1e9 for r in srows[:6]], '"$"#,##0.0', True, 0)
    _panel_label(s, 7.4, 1.75, 5.36, "Growth and margin by segment")
    rows = []
    for r in srows[:6]:
        gr = (r["rev"] / r["rev_prev"] - 1) if _num(r.get("rev_prev")) and r["rev_prev"] > 0 else None
        mg = (r["oi"] / r["rev"]) if _num(r.get("oi")) and r["rev"] else None
        rows.append((r["name"][:24], _pct(gr, 0, True), RED if (_num(gr) and gr < 0) else TEXT,
                     (_pct(mg, 0) if _num(mg) else "–", RED if (_num(mg) and mg < 0) else GREY)))
    _box(s, 7.4, 2.12, 5.36, 0.22, "Segment · growth y/y · op. margin", size=10, color=SRC)
    _rows(s, 7.4, 2.36, 5.36, rows, label_w=3.0, row_h=0.45, value_size=14, extra_w=1.0)


def slide_capital(dk, c):
    fb = c["fb"]
    labels = fb.get("fy_labels") or []
    fcf = fb.get("fy_fcf") or []
    bb = fb.get("fy_buybacks") or []
    dv = fb.get("fy_dividends") or []
    sbc = fb.get("fy_sbc") or []
    if not labels or not any(_num(x) for x in bb + dv):
        return
    tf_ = sum(x for x in fcf if _num(x))
    tr = sum((b or 0) + (d or 0) for b, d in zip(bb, dv))
    ts = sum(x for x in sbc if _num(x))
    title = (f"{c['short']} returned {_pct(tr / tf_, 0)} of free cash flow to shareholders over "
             f"{len(labels)} years" if tf_ > 0 else f"{c['short']}: buybacks and dividends by year")
    s = dk.slide("Capital returns", title,
                 "Source: SEC EDGAR cash-flow statements (10-K): operating cash flow less capex, repurchases of "
                 "common stock, dividends paid and share-based compensation, per fiscal year as filed.")
    fig, ax = _fig(8.0, 3.9)
    x = np.arange(len(labels))
    ax.bar(x - 0.2, [v or 0 for v in fcf], width=0.4, color=MPL["light"], label="Free cash flow")
    ax.bar(x + 0.2, [b or 0 for b in bb], width=0.4, color=MPL["navy"], label="Buybacks")
    ax.bar(x + 0.2, [d or 0 for d in dv], width=0.4, bottom=[b or 0 for b in bb], color="#6C8EB5", label="Dividends")
    ax.plot(x, [v or 0 for v in sbc], color=MPL["red"], lw=1.2, marker="o", ms=3, label="Stock-based pay")
    ax.set_xticks(x)
    ax.set_xticklabels([l.replace("FY20", "FY") for l in labels], fontsize=8)
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}B"))
    ax.legend(loc="upper left", fontsize=8, frameon=False, ncol=2)
    _box(s, M, 1.75, 8.0, 0.25, "Free cash flow and capital returned, $B", size=12, bold=True)
    _image(s, _png(fig), M, 2.03, 8.0, 3.9)
    _panel_label(s, 8.95, 1.75, 3.81, f"{labels[0]}–{labels[-1]} totals")
    _rows(s, 8.95, 2.11, 3.81, [("Free cash flow", f"${tf_:,.0f}B", TEXT),
                               ("Buybacks", f"${sum(b or 0 for b in bb):,.0f}B", TEXT),
                               ("Dividends", f"${sum(d or 0 for d in dv):,.0f}B", TEXT),
                               ("Stock-based pay", f"${ts:,.0f}B", RED)], label_w=2.1)
    if ts and sum(b or 0 for b in bb):
        _box(s, 8.95, 4.1, 3.81, 1.2, [("Dilution offset. ", {"bold": True, "color": NAVY}),
                                       (f"Stock-based pay equalled {ts / sum(b or 0 for b in bb):.0%} of buybacks, so "
                                        "part of the buyback only offsets new shares issued to employees.",
                                        {"color": TEXT})], size=13, lnspc=0.95)


def slide_peers(dk, c):
    R, fb = c["R"], c["fb"]
    peers = R["peers"]
    vh = (c["vhist"] or {}).get("rows") or []
    if not peers and not vh:
        return
    ttm = fb.get("ttm") or {}
    pe = (c["price"] / ttm["eps_clean"]) if (_num(ttm.get("eps_clean")) and ttm["eps_clean"] > 0) else None
    pes = [p["pe"] for p in peers if _num(p.get("pe"))]
    med = float(np.median(pes)) if pes else None
    hist = [r["pe_avg"] for r in vh if _num(r.get("pe_avg"))]
    havg = float(np.mean(hist)) if hist else None
    bits = []
    if _num(pe) and _num(med):
        bits.append(f"{_mult(pe)} earnings against a {_mult(med)} peer median")
    if _num(pe) and _num(havg):
        bits.append(f"{'above' if pe > havg else 'below'} its own {_mult(havg)} average")
    title = (f"{c['ticker']} trades at " + ", ".join(bits)) if bits else f"{c['ticker']} against peers and its history"
    s = dk.slide("Peers & history", title,
                 "Source: each company's own latest 10-Q / 10-K (TTM where newer than the 10-K) and market cap; "
                 "P/E history from filed EPS (restated to today's share basis) and yearly price ranges. "
                 + (f"{c['ticker']} P/E excludes equity-security gains." if fb.get("gains_material") else ""))
    if peers:
        _box(s, M, 1.75, 7.4, 0.25, "Peer comparison", size=12, bold=True)
        rows = [["Ticker", "P/E", "EV/EBITDA", "Rev. growth", "Op. margin", "FCF yield"]]
        me = [c["ticker"], _mult(pe), _mult(c["f"].get("ev_ebitda")), _pct(ttm.get("rev_growth"), 0),
              _pct(ttm.get("op_margin"), 0),
              _pct(((c["f"].get("fcf") or {}).get("fcf_yield") or 0) / 100, 1)]
        rows.append(me)
        for p in peers:
            rows.append([p["ticker"], _mult(p["pe"]), _mult(p["ev_ebitda"]), _pct(p["rev_growth"], 0),
                         _pct(p["op_margin"], 0), _pct(p["fcf_yield"], 1)])
        if pes:
            rows.append(["Peer median", _mult(med), _mult(_md([p["ev_ebitda"] for p in peers])),
                         _pct(_md([p["rev_growth"] for p in peers]), 0), _pct(_md([p["op_margin"] for p in peers]), 0),
                         _pct(_md([p["fcf_yield"] for p in peers]), 1)])
        tbl = s.shapes.add_table(len(rows), 6, Inches(M), Inches(2.06), Inches(7.4), Inches(0.45 * len(rows))).table
        for j, wv in enumerate((1.5, 1.1, 1.3, 1.25, 1.2, 1.05)):
            tbl.columns[j].width = Inches(wv)
        _style_table(tbl, rows, size=13, first_bold=True)
        for j in range(6):
            run = tbl.cell(1, j).text_frame.paragraphs[0].runs[0]
            run.font.bold, run.font.color.rgb = True, _rgb(NAVY)
    if vh:
        x0 = 8.3
        _box(s, x0, 1.75, 4.46, 0.25, "Average P/E by fiscal year", size=12, bold=True)
        ys = [str(r["year"]) for r in vh if _num(r.get("pe_avg"))]
        vs = [r["pe_avg"] for r in vh if _num(r.get("pe_avg"))]
        if vs:
            _native_bar(s, x0, 2.03, 4.46, 3.4, ys[-8:], vs[-8:], '0.0"x"', False, None)


def _md(xs):
    xs = [x for x in xs if _num(x)]
    return float(np.median(xs)) if xs else None


# ── entry point ───────────────────────────────────────────────────────────────
def build_deck(ticker, df, R, company_details=None, mc_summary=None, mc_sim_df=None, news_rows=None,
               fundamentals=None, period_label="", price_source=None, vhist=None, filings=None):
    """The equity review deck as BytesIO. `df` is the report window's prices
    (with SPY_Return when available); `R` is report_inputs.build()."""
    c = _context(ticker, df, R, company_details, mc_summary, mc_sim_df, news_rows, fundamentals,
                 period_label, price_source, vhist, filings)
    dk = Deck(c)
    slide_cover(dk, c)
    slide_summary(dk, c)
    slide_performance(dk, c)
    slide_trend(dk, c)
    slide_risk(dk, c)
    if c["fb"].get("ok"):
        slide_fundamentals(dk, c)
        slide_business_mix(dk, c)
        slide_capital(dk, c)
    slide_valuation(dk, c)
    slide_valuation_range(dk, c)
    slide_peers(dk, c)
    slide_outlook(dk, c)
    slide_signposts(dk, c)
    slide_technical(dk, c)
    slide_profile(dk, c)
    slide_news(dk, c)
    slide_disclosures(dk, c)
    buf = io.BytesIO()
    __import__("doc_props").stamp(dk.prs)
    dk.prs.save(buf)
    buf.seek(0)
    return buf
